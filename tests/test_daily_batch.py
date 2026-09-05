from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter

from gns_app.database import utc_now
from gns_app.domain import (
    AbsCheckResult,
    AbsStatus,
    CaseStatus,
    ExtractedFields,
    ExtractedTaxpayer,
    PageStatus,
    PageType,
    QrStatus,
)
from gns_app.services.workflow import WorkflowValidationError
from gns_app.services.scanner_service import ScannerError
from gns_app.services.outlook_service import (
    OutlookConnectionError,
    OutlookDraftResult,
    OutlookIntegrationError,
    OutlookOutgoingService,
    OutlookSentMatch,
    OutlookSentScan,
    OutlookService,
)


def _insert_ready_case(
    workflow,
    inn: str,
    name: str,
    *,
    source_kind: str = "qr_official",
    district_place: str = "по Ленинскому району города Бишкек",
    recipient_full_name: str = "Телтаев Рахатбек Замирбекович",
    recipient_display_name: str = "Телтаеву Р. З.",
) -> str:
    upload_id = uuid4().hex
    case_id = uuid4().hex
    now = utc_now()
    workflow.db.execute(
        """
        INSERT INTO uploads(
            id, original_filename, stored_path, sha256,
            page_count, status, created_at
        ) VALUES (?, ?, ?, ?, 1, 'ready', ?)
        """,
        (upload_id, f"{case_id}.pdf", f"{case_id}.pdf", case_id, now),
    )
    workflow.db.execute(
        """
        INSERT INTO cases(
            id, upload_id, status, source_kind, district_place,
            recipient_position, recipient_full_name,
            recipient_display_name, period_start, period_end,
            employee_name, fields_confirmed, abs_status,
            created_at, updated_at
        ) VALUES (?, ?, 'ready_for_abs', ?, ?,
                  'Зам. начальника управления', ?, ?,
                  '2020-01-01', '2026-01-01',
                  'Гапарова Э.', 1, 'not_checked', ?, ?)
        """,
        (
            case_id,
            upload_id,
            source_kind,
            district_place,
            recipient_full_name,
            recipient_display_name,
            now,
            now,
        ),
    )
    workflow.db.execute(
        """
        INSERT INTO taxpayers(
            id, case_id, display_order, name, inn,
            name_source, inn_source, manually_confirmed,
            created_at, updated_at
        ) VALUES (?, ?, 1, ?, ?, ?, ?, 1, ?, ?)
        """,
        (
            uuid4().hex,
            case_id,
            name,
            inn,
            source_kind,
            source_kind,
            now,
            now,
        ),
    )
    return case_id


def test_today_batch_checks_and_groups_by_recipient(workflow):
    first = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )
    second = _insert_ready_case(
        workflow, "23456789012345", 'ОсОО "Второй"'
    )

    summary = workflow.check_abs_today("batch-user", "one-time-secret")
    overview = workflow.today_overview()

    assert summary["case_count"] == 2
    assert workflow.abs_session_active()
    assert len(overview["not_found_groups"]) == 1
    group = overview["not_found_groups"][0]
    assert group["case_count"] == 2
    assert group["taxpayer_count"] == 2
    assert group["can_generate"]

    group_id, output = workflow.generate_daily_response(group["group_key"])
    document_text = "\n".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )

    assert group_id
    assert 'ОсОО "Первый"' in document_text
    assert 'ОсОО "Второй"' in document_text
    assert "12345678901234" in document_text
    assert "23456789012345" in document_text
    taxpayer_lines = [
        line.strip()
        for line in document_text.splitlines()
        if "ИНН:" in line
    ]
    assert taxpayer_lines[0].startswith("1. ")
    assert taxpayer_lines[1].startswith("2. ")
    assert workflow.get_case(first)["status"] == CaseStatus.RESPONSE_CREATED
    assert workflow.get_case(second)["status"] == CaseStatus.RESPONSE_CREATED
    assert not workflow.today_overview()["not_found_groups"]
    audit_text = "\n".join(
        event["payload_json"] for event in workflow.get_audit()
    )
    assert "batch-user" not in audit_text
    assert "one-time-secret" not in audit_text


def test_same_inn_from_qr_and_manual_case_appears_once_in_response(workflow):
    inn = "12345678901234"
    name = 'ОсОО "Один налогоплательщик"'
    _insert_ready_case(workflow, inn, name, source_kind="qr_official")
    _insert_ready_case(workflow, inn, name, source_kind="manual")

    workflow.check_abs_today("batch-user", "one-time-secret")
    groups = workflow.today_overview()["not_found_groups"]

    assert len(groups) == 1
    assert groups[0]["case_count"] == 2
    assert groups[0]["taxpayer_count"] == 1
    assert groups[0]["can_generate"]

    _, output = workflow.generate_daily_response(groups[0]["group_key"])
    document_text = "\n".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )
    assert document_text.count(inn) == 1
    assert document_text.count(name) == 1


def test_personal_ip_prefix_difference_is_not_sent_to_manual_review(workflow):
    inn = "12909198201284"
    full_name = "Абдылдаева Анипахан Акматалиевна"
    qr_case_id = _insert_ready_case(
        workflow,
        inn,
        f"ИП {full_name}",
        source_kind="qr_official",
    )
    manual_case_id = _insert_ready_case(
        workflow,
        inn,
        full_name,
        source_kind="manual",
    )

    workflow.check_abs_today("batch-user", "one-time-secret")
    overview = workflow.today_overview()

    assert not workflow.list_case_match_reviews()
    assert overview["unresolved_matches"] == 0
    assert len(overview["not_found_groups"]) == 1
    group = overview["not_found_groups"][0]
    assert group["case_count"] == 2
    assert group["taxpayer_count"] == 1
    assert group["taxpayers"][0]["name"] == f"ИП {full_name}"
    assert group["can_generate"]
    assert workflow.get_taxpayers(qr_case_id)[0]["name"] == f"ИП {full_name}"
    assert workflow.get_taxpayers(manual_case_id)[0]["name"] == full_name


def test_same_inn_with_different_names_blocks_response(workflow):
    inn = "12345678901234"
    _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Название из QR"',
        source_kind="qr_official",
    )
    _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Название после ручной проверки"',
        source_kind="manual",
    )

    workflow.check_abs_today("batch-user", "one-time-secret")
    overview = workflow.today_overview()
    reviews = workflow.list_case_match_reviews()

    assert not overview["not_found_groups"]
    assert overview["unresolved_matches"] == 1
    assert len(reviews) == 1
    assert reviews[0]["differing_field"] == "name"
    assert reviews[0]["left"]["taxpayer"]["inn"] == inn
    assert reviews[0]["right"]["taxpayer"]["inn"] == inn
    assert workflow.get_case(reviews[0]["left_case_id"])["status"] == (
        CaseStatus.NEEDS_REVIEW
    )
    assert workflow.get_case(reviews[0]["right_case_id"])["status"] == (
        CaseStatus.NEEDS_REVIEW
    )


def test_manual_name_correction_resolves_group_conflict(workflow):
    inn = "12345678901234"
    expected_name = 'ОсОО "Название из письма"'
    _insert_ready_case(workflow, inn, expected_name)
    corrected_case_id = _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Ошибочное название"',
        source_kind="manual",
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    taxpayer = workflow.get_taxpayers(corrected_case_id)[0]

    result = workflow.correct_taxpayer(
        corrected_case_id,
        taxpayer["id"],
        name=expected_name,
        inn=inn,
        actor="Гапарова Э.",
    )

    assert result == {
        "changed_fields": ["name"],
        "requires_abs_recheck": False,
    }
    updated = workflow.get_taxpayers(corrected_case_id)[0]
    assert updated["name"] == expected_name
    assert updated["name_source"] == "manual"
    assert updated["abs_result"] == AbsStatus.NOT_FOUND
    group = workflow.today_overview()["not_found_groups"][0]
    assert group["can_generate"]
    assert not group["issues"]
    assert any(
        event["event_type"] == "taxpayer_manually_corrected"
        for event in workflow.get_audit()
    )


def test_manual_inn_correction_requires_new_abs_check(workflow):
    case_id = _insert_ready_case(
        workflow,
        "12345678901234",
        'ОсОО "Налогоплательщик"',
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    taxpayer = workflow.get_taxpayers(case_id)[0]

    result = workflow.correct_taxpayer(
        case_id,
        taxpayer["id"],
        name=taxpayer["name"],
        inn="99999999999999",
    )

    assert result["requires_abs_recheck"]
    assert workflow.get_case(case_id)["status"] == CaseStatus.READY_FOR_ABS
    assert workflow.get_case(case_id)["abs_status"] == AbsStatus.NOT_CHECKED
    updated = workflow.get_taxpayers(case_id)[0]
    assert updated["inn"] == "99999999999999"
    assert updated["inn_source"] == "manual"
    assert updated["abs_result"] is None


def test_structured_conflict_resolution_uses_explicit_field_sources(workflow):
    inn = "12345678901234"
    expected_name = 'ОсОО "Название из QR"'
    first_case_id = _insert_ready_case(
        workflow,
        inn,
        expected_name,
        source_kind="qr_official",
    )
    second_case_id = _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Другое название"',
        source_kind="manual",
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    workflow.today_overview()
    review = workflow.list_case_match_reviews()[0]
    chosen = (
        review["left"]
        if review["left_case_id"] == first_case_id
        else review["right"]
    )["taxpayer"]

    result = workflow.resolve_case_match_review(
        review["id"],
        decision="same",
        name_choice=chosen["id"],
        inn_choice=chosen["id"],
        actor="Гапарова Э.",
    )

    assert result["decision"] == "same"
    assert not result["requires_abs_recheck"]
    for case_id in (first_case_id, second_case_id):
        taxpayer = workflow.get_taxpayers(case_id)[0]
        assert taxpayer["name"] == expected_name
        assert taxpayer["inn"] == inn
        assert taxpayer["name_source"] == "qr_official"
        assert taxpayer["name_source_reference"] == (
            f"taxpayer:{chosen['id']}"
        )
        assert taxpayer["abs_result"] == AbsStatus.NOT_FOUND
    updated_group = workflow.today_overview()["not_found_groups"][0]
    assert updated_group["can_generate"]
    assert not updated_group["conflicts"]


def test_structured_conflict_manual_inn_resets_all_affected_cases(workflow):
    original_inn = "12345678901234"
    new_inn = "99999999999999"
    first_case_id = _insert_ready_case(
        workflow,
        original_inn,
        'ОсОО "Первый вариант"',
    )
    second_case_id = _insert_ready_case(
        workflow,
        original_inn,
        'ОсОО "Второй вариант"',
        source_kind="manual",
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    workflow.today_overview()
    review = workflow.list_case_match_reviews()[0]
    chosen_name = review["left"]["taxpayer"]

    result = workflow.resolve_case_match_review(
        review["id"],
        decision="same",
        name_choice=chosen_name["id"],
        inn_choice="manual",
        manual_inn=new_inn,
    )

    assert result["requires_abs_recheck"]
    assert result["recheck_case_ids"] == sorted(
        [first_case_id, second_case_id]
    )
    for case_id in (first_case_id, second_case_id):
        case = workflow.get_case(case_id)
        taxpayer = workflow.get_taxpayers(case_id)[0]
        # В тестовом режиме повторная АБС запускается сразу после явного
        # подтверждения сотрудника.
        assert case["status"] == CaseStatus.READY_FOR_RESPONSE
        assert case["abs_status"] == AbsStatus.NOT_FOUND
        assert taxpayer["inn"] == new_inn
        assert taxpayer["inn_source"] == "manual"
        assert taxpayer["inn_source_reference"] is None
        assert taxpayer["abs_result"] == AbsStatus.NOT_FOUND


def test_late_qr_variance_creates_a_new_manual_confirmation(workflow):
    inn = "12345678901234"
    first_case_id = _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Подтверждённое название"',
        source_kind="manual",
    )
    second_case_id = _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Ручной вариант"',
        source_kind="manual",
    )
    workflow.reconcile_case_match_reviews()
    first_review = workflow.list_case_match_reviews()[0]
    chosen = first_review["left"]["taxpayer"]
    workflow.resolve_case_match_review(
        first_review["id"],
        decision="same",
        name_choice=chosen["id"],
        inn_choice=chosen["id"],
    )
    assert not workflow.list_case_match_reviews()

    # Официальная версия пришла позже и снова изменила подтверждённое поле.
    qr_case_id = second_case_id
    qr_taxpayer = workflow.get_taxpayers(qr_case_id)[0]
    workflow.db.execute(
        """
        UPDATE cases
        SET source_kind = 'qr_official', official_parse_version = 2,
            status = 'ready_for_response', updated_at = ?
        WHERE id = ?
        """,
        (utc_now(), qr_case_id),
    )
    workflow.db.execute(
        """
        UPDATE taxpayers
        SET name = 'ОсОО "Название из позднего QR"',
            name_source = 'qr_official', updated_at = ?
        WHERE id = ?
        """,
        (utc_now(), qr_taxpayer["id"]),
    )

    assert workflow.reconcile_case_match_reviews() == 1
    second_review = workflow.list_case_match_reviews()[0]
    assert second_review["id"] != first_review["id"]
    assert second_review["differing_field"] == "name"
    assert workflow.get_case(first_case_id)["status"] == CaseStatus.NEEDS_REVIEW
    assert workflow.get_case(second_case_id)["status"] == CaseStatus.NEEDS_REVIEW
    assert not workflow.today_overview()["not_found_groups"]


def test_same_name_with_different_inn_is_blocked_until_confirmation(workflow):
    name = 'ОсОО "Одинаковое название"'
    _insert_ready_case(workflow, "12345678901234", name)
    _insert_ready_case(workflow, "99999999999999", name, source_kind="manual")

    workflow.check_abs_today("batch-user", "one-time-secret")
    assert not workflow.today_overview()["not_found_groups"]
    review = workflow.list_case_match_reviews()[0]
    assert review["differing_field"] == "inn"
    assert "отличается ИНН" in review["issue_message"]


@pytest.mark.parametrize(
    ("inn", "warning"),
    [
        ("223323", "6 цифр вместо 14"),
        ("123456789012345", "15 цифр вместо 14"),
    ],
)
def test_nonstandard_inn_warns_but_does_not_block_response(
    workflow, inn, warning
):
    case_id = _insert_ready_case(
        workflow,
        inn,
        'ОсОО "ИНН из письма"',
    )
    workflow.db.execute(
        "UPDATE cases SET period_start = NULL, period_end = NULL, "
        "period_route = 'no_odb' WHERE id = ?",
        (case_id,),
    )

    assert workflow.auto_check_abs(case_id)

    case = workflow.get_case(case_id)
    taxpayer = workflow.get_taxpayers(case_id)[0]
    group = workflow.today_overview()["not_found_groups"][0]
    assert case["status"] == CaseStatus.READY_FOR_RESPONSE
    assert case["abs_status"] == AbsStatus.INVALID_INN
    assert taxpayer["abs_result"] == AbsStatus.INVALID_INN
    assert group["can_generate"]
    assert warning in group["warnings"][0]

    _group_id, output = workflow.generate_daily_response(group["group_key"])
    assert output.exists()


def test_manual_fields_accept_nonstandard_numeric_inn(workflow):
    taxpayers = workflow._validate_manual_fields(
        "г.Ош",
        "Начальнику",
        "Получатель",
        "",
        "",
        "no_odb",
        "Сотрудник",
        [{"name": 'ОсОО "ИНН из письма"', "inn": "223323"}],
    )

    assert taxpayers == [
        {"name": 'ОсОО "ИНН из письма"', "inn": "223323"}
    ]


def test_abs_checks_valid_inn_and_skips_nonstandard_in_same_letter(workflow):
    case_id = _insert_ready_case(
        workflow,
        "12345678901234",
        'ОсОО "Обычный ИНН"',
    )
    workflow._replace_taxpayers(
        case_id,
        [
            {"name": 'ОсОО "Обычный ИНН"', "inn": "12345678901234"},
            {"name": 'ОсОО "ИНН из письма"', "inn": "223323"},
        ],
    )
    workflow.db.execute(
        "UPDATE cases SET period_start = NULL, period_end = NULL, "
        "period_route = 'no_odb' WHERE id = ?",
        (case_id,),
    )

    result = workflow.check_abs(case_id, "batch-user", "one-time-secret")

    taxpayers = {item["inn"]: item for item in workflow.get_taxpayers(case_id)}
    assert result.status == AbsStatus.NOT_FOUND
    assert taxpayers["12345678901234"]["abs_result"] == AbsStatus.NOT_FOUND
    assert taxpayers["223323"]["abs_result"] == AbsStatus.INVALID_INN
    assert workflow.get_case(case_id)["status"] == CaseStatus.READY_FOR_RESPONSE


def test_distinct_decision_restores_both_versions_without_requeue(workflow):
    inn = "12345678901234"
    _insert_ready_case(workflow, inn, 'ОсОО "Первое лицо"')
    _insert_ready_case(
        workflow,
        inn,
        'ОсОО "Второе лицо"',
        source_kind="manual",
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    workflow.today_overview()
    review = workflow.list_case_match_reviews()[0]

    result = workflow.resolve_case_match_review(
        review["id"], decision="distinct"
    )

    assert result["decision"] == "distinct"
    assert not workflow.list_case_match_reviews()
    groups = workflow.today_overview()["not_found_groups"]
    assert len(groups) == 1
    assert groups[0]["can_generate"]
    assert groups[0]["taxpayer_count"] == 2


def test_same_inn_for_another_recipient_stays_in_separate_response(workflow):
    inn = "12345678901234"
    name = 'ОсОО "Один налогоплательщик"'
    _insert_ready_case(workflow, inn, name, source_kind="qr_official")
    _insert_ready_case(
        workflow,
        inn,
        name,
        source_kind="manual",
        recipient_full_name="Асанова Айгуль Токтогуловна",
        recipient_display_name="Асановой А. Т.",
    )

    workflow.check_abs_today("batch-user", "one-time-secret")
    groups = workflow.today_overview()["not_found_groups"]

    assert len(groups) == 2
    assert all(group["taxpayer_count"] == 1 for group in groups)
    assert all(group["can_generate"] for group in groups)


def test_abs_session_reuses_credentials_only_in_memory(workflow):
    first = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )
    second = _insert_ready_case(
        workflow, "23456789012345", 'ОсОО "Второй"'
    )

    workflow.check_abs(first, "session-user", "session-secret")
    assert workflow.abs_session_active()
    workflow.check_abs(second)

    assert workflow.get_case(second)["status"] == CaseStatus.READY_FOR_RESPONSE
    audit_text = "\n".join(event["payload_json"] for event in workflow.get_audit())
    assert "session-user" not in audit_text
    assert "session-secret" not in audit_text


@pytest.mark.parametrize(
    ("username", "password", "expected_status"),
    [
        ("session-user", "invalid", AbsStatus.AUTH_ERROR),
        ("offline", "session-secret", AbsStatus.UNAVAILABLE),
    ],
)
def test_abs_error_clears_session_and_requires_login(
    workflow,
    username,
    password,
    expected_status,
):
    case_id = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Клиент"'
    )

    result = workflow.check_abs(case_id, username, password)

    assert result.status == expected_status
    assert not workflow.abs_session_active()
    assert workflow.abs_login_required()


def test_real_abs_credentials_are_not_reused_between_checks(workflow):
    class RealReadOnlyGateway:
        is_fake = False
        supports_session = False

        @staticmethod
        def check(username, password, taxpayers):
            return AbsCheckResult(
                status=AbsStatus.NOT_FOUND,
                taxpayers=[
                    {
                        "inn": item["inn"],
                        "name": item["name"],
                        "result": AbsStatus.NOT_FOUND,
                    }
                    for item in taxpayers
                ],
                message="Счета не найдены.",
                is_fake=False,
            )

    workflow.abs = RealReadOnlyGateway()
    first = _insert_ready_case(workflow, "12345678901234", "Первый")
    second = _insert_ready_case(workflow, "23456789012345", "Второй")

    workflow.check_abs(first, "employee", "one-time-secret")

    assert not workflow.abs_session_active()
    audit = workflow.get_audit()
    event = next(item for item in audit if item["event_type"] == "abs_checked")
    assert "employee" not in event["payload_json"]
    assert "one-time-secret" not in event["payload_json"]
    with pytest.raises(WorkflowValidationError, match="логин и пароль"):
        workflow.check_abs(second)


def test_found_taxpayer_is_separate_and_never_uses_absence_template(workflow):
    _insert_ready_case(
        workflow, "11111111111111", 'ОсОО "Найденный"'
    )

    workflow.check_abs_today("batch-user", "one-time-secret")
    overview = workflow.today_overview()

    assert not overview["not_found_groups"]
    assert len(overview["found_groups"]) == 1
    assert not overview["found_groups"][0]["can_generate"]
    assert not overview["found_groups"][0]["issues"]


def test_ocr_only_decision_stays_in_manual_review(workflow):
    status, issue_code, _ = workflow._page_outcome(
        PageType.DECISION,
        0.99,
        QrStatus.NOT_FOUND,
        False,
        None,
    )

    assert status == PageStatus.NEEDS_REVIEW
    assert issue_code == "decision_type_not_confident"


def test_confident_ocr_decision_can_finish_without_manual_fields(workflow):
    status, issue_code, issue_message = workflow._page_outcome(
        PageType.DECISION,
        0.78,
        QrStatus.NOT_FOUND,
        False,
        None,
        True,
    )

    assert status == PageStatus.COMPLETED
    assert issue_code is None
    assert issue_message is None


def test_ocr_disagreement_does_not_prefill_critical_fields(workflow):
    upload_id = uuid4().hex
    now = utc_now()
    workflow.db.execute(
        """
        INSERT INTO uploads(
            id, original_filename, stored_path, sha256,
            page_count, status, created_at
        ) VALUES (?, 'scan.pdf', 'scan.pdf', 'hash', 1, 'ready', ?)
        """,
        (upload_id, now),
    )
    case_id = workflow._create_scan_case(upload_id, "page-for-audit")
    fields = ExtractedFields(
        period_start="2020-01-01",
        period_end="2026-01-01",
        taxpayers=[
            ExtractedTaxpayer(
                name='ОсОО "Сомнительный"',
                inn="12345678901234",
                confidence=0.5,
            )
        ],
    )

    workflow._prefill_scan_case(
        case_id,
        fields,
        critical_fields_agree=False,
    )

    case = workflow.get_case(case_id)
    assert case["period_start"] is None
    assert case["period_end"] is None
    assert workflow.get_taxpayers(case_id) == []


def test_generate_all_ready_daily_responses_creates_one_file_per_recipient(
    workflow,
):
    first = _insert_ready_case(
        workflow,
        "12345678901234",
        'ОсОО "Первый"',
        district_place="по Ленинскому району города Бишкек",
        recipient_full_name="Телтаев Рахатбек Замирбекович",
        recipient_display_name="Телтаеву Р. З.",
    )
    second = _insert_ready_case(
        workflow,
        "23456789012345",
        'ОсОО "Второй"',
        district_place="по Свердловскому району города Бишкек",
        recipient_full_name="Асанова Айгуль Токтогуловна",
        recipient_display_name="Асановой А. Т.",
    )

    workflow.check_abs_today("batch-user", "one-time-secret")
    overview = workflow.today_overview()
    assert len(overview["not_found_groups"]) == 2

    summary = workflow.generate_all_ready_daily_responses()

    assert len(summary["created"]) == 2
    assert summary["errors"] == []
    assert workflow.get_case(first)["status"] == CaseStatus.RESPONSE_CREATED
    assert workflow.get_case(second)["status"] == CaseStatus.RESPONSE_CREATED
    assert not workflow.today_overview()["not_found_groups"]

    # Повторный вызов не должен падать и не находит новых готовых групп.
    empty_summary = workflow.generate_all_ready_daily_responses()
    assert empty_summary == {"created": [], "errors": []}


def test_split_response_registers_one_letter_per_word_section(workflow):
    for index in range(1, 5):
        _insert_ready_case(
            workflow,
            f"{index:014d}",
            f'ОсОО "Тест {index}"',
        )

    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, _ = workflow.generate_daily_response(
        group["group_key"], taxpayers_per_page=2
    )

    letters = workflow.db.fetch_all(
        "SELECT * FROM response_letters WHERE response_group_id = ? "
        "ORDER BY letter_order",
        (group_id,),
    )
    assert [letter["letter_order"] for letter in letters] == [1, 2]
    assert [letter["taxpayer_start_order"] for letter in letters] == [1, 3]
    assert [letter["taxpayer_count"] for letter in letters] == [2, 2]
    assert all(letter["outgoing_number"] is None for letter in letters)


def test_batch_outgoing_numbers_are_previewed_assigned_and_written_to_word(
    workflow,
):
    for index in range(1, 5):
        _insert_ready_case(
            workflow,
            f"{index:014d}",
            f'ОсОО "Тест {index}"',
        )

    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, output = workflow.generate_daily_response(
        group["group_key"], taxpayers_per_page=2
    )
    preview = workflow.preview_outgoing_numbers("9544")

    assert preview["count"] == 2
    assert preview["first_number"] == "9544"
    assert preview["last_number"] == "9545"
    assert [
        item["proposed_number"] for item in preview["letters"]
    ] == ["9544", "9545"]

    summary = workflow.assign_outgoing_numbers(
        "9544",
        [item["id"] for item in preview["letters"]],
        actor="Тестовый сотрудник",
    )

    assert summary["letter_count"] == 2
    letters = workflow.db.fetch_all(
        "SELECT * FROM response_letters WHERE response_group_id = ? "
        "ORDER BY letter_order",
        (group_id,),
    )
    assert [letter["outgoing_number"] for letter in letters] == [
        "9544",
        "9545",
    ]
    document_text = "\n".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )
    assert "04-1/9544" in document_text
    assert "04-1/9545" in document_text
    events = workflow.db.fetch_all(
        "SELECT * FROM audit_events "
        "WHERE event_type = 'outgoing_number_assigned' ORDER BY id"
    )
    assert len(events) == 2
    assert all(event["actor"] == "Тестовый сотрудник" for event in events)


def test_batch_outgoing_numbers_can_start_from_an_inline_letter_suffix(
    workflow,
):
    for index in range(1, 9):
        _insert_ready_case(
            workflow,
            f"{index:014d}",
            f'ОсОО "Тест {index}"',
        )

    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, output = workflow.generate_daily_response(
        group["group_key"], taxpayers_per_page=2
    )
    letters = workflow.db.fetch_all(
        "SELECT * FROM response_letters WHERE response_group_id = ? "
        "ORDER BY letter_order",
        (group_id,),
    )
    workflow.set_outgoing_number(letters[2]["id"], "7000")

    summary = workflow.assign_outgoing_numbers(
        "9544",
        [letters[1]["id"], letters[3]["id"]],
        actor="Тестовый сотрудник",
    )

    assert summary["letter_count"] == 2
    refreshed = workflow.db.fetch_all(
        "SELECT outgoing_number FROM response_letters "
        "WHERE response_group_id = ? ORDER BY letter_order",
        (group_id,),
    )
    assert [letter["outgoing_number"] for letter in refreshed] == [
        None,
        "9544",
        "7000",
        "9545",
    ]
    document_text = "\n".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )
    assert "04-1/9544" in document_text
    assert "04-1/7000" in document_text
    assert "04-1/9545" in document_text


def test_batch_outgoing_numbers_reject_non_contiguous_inline_selection(
    workflow,
):
    for index in range(1, 7):
        _insert_ready_case(
            workflow,
            f"{index:014d}",
            f'ОсОО "Тест {index}"',
        )

    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, _ = workflow.generate_daily_response(
        group["group_key"], taxpayers_per_page=2
    )
    letters = workflow.db.fetch_all(
        "SELECT * FROM response_letters WHERE response_group_id = ? "
        "ORDER BY letter_order",
        (group_id,),
    )

    with pytest.raises(WorkflowValidationError):
        workflow.assign_outgoing_numbers(
            "9544",
            [letters[0]["id"], letters[2]["id"]],
        )

    refreshed = workflow.db.fetch_all(
        "SELECT outgoing_number FROM response_letters "
        "WHERE response_group_id = ? ORDER BY letter_order",
        (group_id,),
    )
    assert [letter["outgoing_number"] for letter in refreshed] == [
        None,
        None,
        None,
    ]


def test_outgoing_number_correction_is_unique_and_regenerates_word(workflow):
    _insert_ready_case(workflow, "12345678901234", 'ОсОО "Первый"')
    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, output = workflow.generate_daily_response(group["group_key"])
    letter = workflow.db.fetch_one(
        "SELECT * FROM response_letters WHERE response_group_id = ?",
        (group_id,),
    )
    workflow.mark_response_group_opened_for_print(group_id)
    assert workflow.get_response_group(group_id)["opened_for_print_at"]

    workflow.set_outgoing_number(letter["id"], "9544")
    workflow.set_outgoing_number(letter["id"], "9600")

    refreshed = workflow.db.fetch_one(
        "SELECT * FROM response_letters WHERE id = ?", (letter["id"],)
    )
    assert refreshed["outgoing_number"] == "9600"
    assert workflow.get_response_group(group_id)["opened_for_print_at"] is None
    document_text = "\n".join(
        paragraph.text for paragraph in Document(output).paragraphs
    )
    assert "04-1/9600" in document_text
    assert "04-1/9544" not in document_text
    changed = workflow.db.fetch_one(
        "SELECT * FROM audit_events "
        "WHERE entity_id = ? AND event_type = 'outgoing_number_changed'",
        (letter["id"],),
    )
    assert changed is not None


def test_outgoing_number_cannot_be_reused_for_another_letter(workflow):
    _insert_ready_case(
        workflow,
        "12345678901234",
        'ОсОО "Первый"',
        recipient_full_name="Телтаев Рахатбек Замирбекович",
        recipient_display_name="Телтаеву Р. З.",
    )
    _insert_ready_case(
        workflow,
        "23456789012345",
        'ОсОО "Второй"',
        recipient_full_name="Асанова Айгуль Токтогуловна",
        recipient_display_name="Асановой А. Т.",
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    workflow.generate_all_ready_daily_responses()
    letters = workflow.list_outgoing_letters()

    workflow.set_outgoing_number(letters[0]["id"], "9544")
    with pytest.raises(WorkflowValidationError, match="уже используется"):
        workflow.set_outgoing_number(letters[1]["id"], "9544")


def test_created_responses_assign_numbers_from_the_inline_letter_field(
    workflow, monkeypatch
):
    from starlette.testclient import TestClient

    from gns_app import main

    _insert_ready_case(workflow, "12345678901234", 'ОсОО "Первый"')
    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    workflow.generate_daily_response(group["group_key"])
    letter = workflow.list_outgoing_letters(only_unnumbered=True)[0]
    monkeypatch.setattr(main, "workflow", workflow)
    client = TestClient(main.app)

    page = client.get("/?tab=responses&response_view=created")

    assert page.status_code == 200
    assert 'class="number-sequence-bar"' not in page.text
    assert 'data-number-sequence-start' not in page.text
    assert (
        f'data-inline-number-form data-letter-id="{letter["id"]}"'
        in page.text
    )
    assert 'data-inline-number-input' in page.text
    assert "Первый свободный исх. №" not in page.text

    assigned = client.post(
        "/today/outgoing-numbers/assign",
        data={"first_number": "9544", "letter_ids": [letter["id"]]},
        follow_redirects=False,
    )

    assert assigned.status_code == 303
    refreshed = workflow.db.fetch_one(
        "SELECT outgoing_number FROM response_letters WHERE id = ?",
        (letter["id"],),
    )
    assert refreshed["outgoing_number"] == "9544"


def test_created_responses_number_old_unfinished_backlog(
    workflow, monkeypatch
):
    from starlette.testclient import TestClient

    from gns_app import main

    _insert_ready_case(workflow, "12345678901234", 'ОсОО "Старое письмо"')
    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, _ = workflow.generate_daily_response(group["group_key"])
    workflow.db.execute(
        "UPDATE response_groups SET business_date = ? WHERE id = ?",
        ("2000-01-01", group_id),
    )
    letter = workflow.list_outgoing_letters(only_unnumbered=True)[0]
    monkeypatch.setattr(main, "workflow", workflow)

    page = TestClient(main.app).get(
        "/?tab=responses&response_view=created"
    )

    assert page.status_code == 200
    assert 'class="number-sequence-bar"' not in page.text
    assert (
        f'data-inline-number-form data-letter-id="{letter["id"]}"'
        in page.text
    )


def test_created_responses_do_not_open_scan_drawer_without_explicit_focus(
    workflow,
    monkeypatch,
):
    from starlette.testclient import TestClient

    from gns_app import main

    _, letter = _ready_response_letter(workflow)
    workflow.start_signed_scan_session(letter["id"])
    monkeypatch.setattr(main, "workflow", workflow)
    client = TestClient(main.app)

    closed = client.get("/?tab=responses&response_view=created")
    opened = client.get(
        "/?tab=responses&response_view=created"
        f"&focus_letter={letter['id']}"
    )
    unknown = client.get(
        "/?tab=responses&response_view=created&focus_letter=missing"
    )

    assert closed.status_code == 200
    assert 'id="letter-workflow"' not in closed.text
    assert opened.status_code == 200
    assert 'id="letter-workflow"' in opened.text
    assert unknown.status_code == 200
    assert 'id="letter-workflow"' not in unknown.text


def test_outgoing_number_preview_rejects_range_overflow(
    workflow, monkeypatch
):
    monkeypatch.setattr(
        workflow,
        "list_outgoing_letters",
        lambda *args, **kwargs: [{"id": "one"}, {"id": "two"}],
    )

    with pytest.raises(WorkflowValidationError, match="превышает"):
        workflow.preview_outgoing_numbers("999999999")


def test_today_page_opens_generated_response_in_word(workflow, monkeypatch):
    from starlette.testclient import TestClient

    from gns_app import main

    _insert_ready_case(workflow, "12345678901234", 'ОсОО "Первый"')
    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, output = workflow.generate_daily_response(group["group_key"])
    opened: list[Path] = []
    focus_requests: list[dict] = []
    monkeypatch.setattr(main, "workflow", workflow)
    monkeypatch.setattr(main, "open_word_document", opened.append)
    monkeypatch.setattr(
        main,
        "start_foreground_watcher",
        lambda **kwargs: focus_requests.append(kwargs),
    )
    client = TestClient(main.app)

    before = client.get("/?tab=responses&response_view=created")
    assert re.search(r">\s*Открыть письмо\s*<", before.text)
    assert not re.search(r">\s*Сканировать\s*<", before.text)

    response = client.post(
        f"/response-groups/{group_id}/open-for-print",
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith(
        "/?tab=responses&response_view=created"
    )
    assert opened == [output]
    assert focus_requests[0]["title_parts"] == (output.stem,)
    assert focus_requests[0]["class_parts"] == ("opusapp",)
    assert focus_requests[0]["keep_foreground_seconds"] == 2.0
    assert workflow.get_response_group(group_id)["opened_for_print_at"]

    after = client.get("/?tab=responses&response_view=created")
    assert re.search(r">\s*Сканировать\s*<", after.text)
    assert "Можно сканировать" in after.text


def _ready_response_letter(workflow):
    _insert_ready_case(workflow, "12345678901234", 'ОсОО "Первый"')
    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, _ = workflow.generate_daily_response(group["group_key"])
    letter = workflow.db.fetch_one(
        "SELECT * FROM response_letters WHERE response_group_id = ?",
        (group_id,),
    )
    return group_id, letter


def _png_scan() -> BytesIO:
    stream = BytesIO()
    Image.new("RGB", (1240, 1754), "white").save(stream, "PNG")
    stream.seek(0)
    return stream


def test_signed_a4_scan_is_registered_and_ready_without_second_step(workflow):
    _, letter = _ready_response_letter(workflow)

    scan = workflow.register_signed_response_scan(
        letter["id"],
        "Подписанный ответ.png",
        _png_scan(),
        actor="Тестовый сотрудник",
    )

    assert scan["source"] == "upload"
    assert scan["status"] == "confirmed"
    assert len(PdfReader(scan["pdf_path"]).pages) == 1
    overview_letter = workflow.today_overview()["generated_groups"][0][
        "letters"
    ][0]
    assert overview_letter["signed_scan"]["id"] == scan["id"]
    with pytest.raises(WorkflowValidationError, match="уже проверен"):
        workflow.confirm_signed_response_scan(
            scan["id"],
            correct_letter=True,
            signature_present=True,
            bank_seal_present=True,
            actor="Тестовый сотрудник",
        )

    confirmed = workflow.get_signed_response_scan(scan["id"])
    assert confirmed["status"] == "confirmed"
    assert confirmed["confirmed_by"] == "Тестовый сотрудник"
    assert confirmed["correct_letter_confirmed"] == 1
    assert confirmed["signature_confirmed"] == 1
    assert confirmed["bank_seal_confirmed"] == 1


def test_new_signed_scan_supersedes_previous_without_deleting_it(workflow):
    _, letter = _ready_response_letter(workflow)
    first = workflow.register_signed_response_scan(
        letter["id"], "first.png", _png_scan()
    )
    second = workflow.register_signed_response_scan(
        letter["id"], "second.png", _png_scan()
    )

    assert workflow.get_signed_response_scan(first["id"])["status"] == "superseded"
    assert workflow.get_signed_response_scan(second["id"])["status"] == "confirmed"
    assert Path(first["original_path"]).is_file()


def test_restart_manual_review_supersedes_only_unsent_response_group(workflow):
    case_id = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )
    case = workflow.get_case(case_id)
    assert case is not None
    page_id = uuid4().hex
    now = utc_now()
    workflow.db.execute(
        """
        INSERT INTO pages(
            id, upload_id, case_id, page_number, status,
            created_at, updated_at
        ) VALUES (?, ?, ?, 1, 'completed', ?, ?)
        """,
        (page_id, case["upload_id"], case_id, now, now),
    )
    workflow.check_abs_today("batch-user", "one-time-secret")
    group = workflow.today_overview()["not_found_groups"][0]
    group_id, _ = workflow.generate_daily_response(group["group_key"])
    letter = workflow.db.fetch_one(
        "SELECT * FROM response_letters WHERE response_group_id = ?",
        (group_id,),
    )
    assert letter is not None
    scan = workflow.register_signed_response_scan(
        letter["id"], "signed.png", _png_scan()
    )

    page = workflow.restart_case_manual_review(case_id)

    assert page["id"] == page_id
    assert workflow.get_case(case_id)["status"] == "needs_review"
    assert workflow.get_response_group(group_id)["status"] == "superseded"
    assert workflow.get_signed_response_scan(scan["id"])["status"] == "superseded"


def test_signed_scan_preserves_multi_page_pdf(workflow):
    _, letter = _ready_response_letter(workflow)
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    stream = BytesIO()
    writer.write(stream)
    stream.seek(0)

    scan = workflow.register_signed_response_scan(
        letter["id"], "two-pages.pdf", stream
    )

    assert scan["page_count"] == 2
    assert len(PdfReader(scan["pdf_path"]).pages) == 2
    assert scan["status"] == "confirmed"


def test_device_scan_is_registered_through_wia_gateway(workflow, monkeypatch):
    _, letter = _ready_response_letter(workflow)

    scan_profiles = []

    def fake_acquire(destination, timeout_seconds=900, **scanner_profile):
        scan_profiles.append(scanner_profile)
        Image.new("RGB", (1240, 1754), "white").save(destination, "PNG")
        return destination

    monkeypatch.setattr(workflow.scanner, "acquire_a4", fake_acquire)

    scan = workflow.acquire_signed_response_scan(letter["id"])

    assert scan["source"] == "wia"
    assert scan["status"] == "confirmed"
    assert scan_profiles == [{"dpi": 150, "color_mode": "grayscale"}]


def test_multi_page_wia_session_builds_one_pdf(workflow, monkeypatch):
    _, letter = _ready_response_letter(workflow)
    acquired = 0

    def fake_acquire(destination, timeout_seconds=900, **scanner_profile):
        nonlocal acquired
        acquired += 1
        color = "white" if acquired == 1 else "lightgray"
        Image.new("RGB", (620, 877), color).save(destination, "PNG")
        return destination

    monkeypatch.setattr(workflow.scanner, "acquire_a4", fake_acquire)
    session = workflow.start_signed_scan_session(
        letter["id"], actor="Тестовый сотрудник"
    )
    first = workflow.acquire_signed_scan_session_page(session["id"])
    second = workflow.acquire_signed_scan_session_page(session["id"])

    assert first["page_count"] == 1
    assert second["page_count"] == 2
    assert all(
        Path(page["original_path"]).is_file() for page in second["pages"]
    )

    scan = workflow.finalize_signed_scan_session(
        session["id"], actor="Тестовый сотрудник"
    )

    assert scan["source"] == "wia"
    assert scan["page_count"] == 2
    assert len(PdfReader(scan["pdf_path"]).pages) == 2
    completed = workflow.get_signed_scan_session(session["id"])
    assert completed["status"] == "completed"
    assert completed["result_scan_id"] == scan["id"]
    assert Path(scan["original_path"]).is_file()


def test_wia_session_error_keeps_previous_pages_and_can_resume(
    workflow,
    monkeypatch,
):
    _, letter = _ready_response_letter(workflow)
    calls = 0

    def flaky_acquire(destination, timeout_seconds=900, **scanner_profile):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ScannerError("Тестовая ошибка сканера")
        Image.new("RGB", (620, 877), "white").save(destination, "PNG")
        return destination

    monkeypatch.setattr(workflow.scanner, "acquire_a4", flaky_acquire)
    session = workflow.start_signed_scan_session(letter["id"])
    workflow.acquire_signed_scan_session_page(session["id"])

    with pytest.raises(ScannerError, match="Тестовая ошибка"):
        workflow.acquire_signed_scan_session_page(session["id"])

    failed = workflow.get_signed_scan_session(session["id"])
    assert failed["status"] == "technical_error"
    assert failed["page_count"] == 1
    assert len(failed["pages"]) == 1

    resumed = workflow.acquire_signed_scan_session_page(session["id"])
    assert resumed["status"] == "collecting"
    assert resumed["error_message"] is None
    assert resumed["page_count"] == 2


def test_wia_session_reorders_replaces_and_removes_pages(workflow, monkeypatch):
    _, letter = _ready_response_letter(workflow)
    acquired = 0

    def fake_acquire(destination, timeout_seconds=900, **scanner_profile):
        nonlocal acquired
        acquired += 1
        Image.new("RGB", (620, 877), (acquired, acquired, acquired)).save(
            destination,
            "PNG",
        )
        return destination

    monkeypatch.setattr(workflow.scanner, "acquire_a4", fake_acquire)
    session = workflow.start_signed_scan_session(letter["id"])
    workflow.acquire_signed_scan_session_page(session["id"])
    current = workflow.acquire_signed_scan_session_page(session["id"])
    first_id, second_id = [page["id"] for page in current["pages"]]

    moved = workflow.move_signed_scan_session_page(
        session["id"], second_id, "up"
    )
    assert [page["id"] for page in moved["pages"]] == [second_id, first_id]

    replaced = workflow.acquire_signed_scan_session_page(
        session["id"], replace_page_id=second_id
    )
    replacement_id = replaced["pages"][0]["id"]
    assert replacement_id != second_id
    removed_raw = workflow.db.fetch_one(
        "SELECT * FROM signed_scan_session_pages WHERE id = ?",
        (second_id,),
    )
    assert removed_raw["status"] == "removed"
    assert Path(removed_raw["original_path"]).is_file()

    remaining = workflow.remove_signed_scan_session_page(
        session["id"], replacement_id
    )
    assert remaining["page_count"] == 1
    assert [page["page_order"] for page in remaining["pages"]] == [1]


def test_cancelled_wia_session_does_not_create_or_supersede_scan(workflow):
    _, letter = _ready_response_letter(workflow)
    existing = workflow.register_signed_response_scan(
        letter["id"], "existing.png", _png_scan()
    )
    session = workflow.start_signed_scan_session(letter["id"])

    workflow.cancel_signed_scan_session(session["id"])

    assert workflow.get_signed_scan_session(session["id"])["status"] == "cancelled"
    assert workflow.get_signed_response_scan(existing["id"])["status"] == (
        "confirmed"
    )


def test_ready_file_upload_closes_unfinished_wia_session(workflow):
    _, letter = _ready_response_letter(workflow)
    session = workflow.start_signed_scan_session(letter["id"])

    scan = workflow.register_signed_response_scan(
        letter["id"], "ready-from-mfu.png", _png_scan()
    )

    assert scan["status"] == "confirmed"
    assert workflow.get_signed_scan_session(session["id"])["status"] == (
        "cancelled"
    )
    assert workflow.get_active_signed_scan_session(letter["id"]) is None


def test_signed_scan_upload_route_marks_scan_ready(workflow, monkeypatch):
    from starlette.testclient import TestClient

    from gns_app import main

    group_id, letter = _ready_response_letter(workflow)
    png = _png_scan().getvalue()
    monkeypatch.setattr(main, "workflow", workflow)
    client = TestClient(main.app)

    uploaded = client.post(
        f"/response-letters/{letter['id']}/scan-upload",
        data={"group_id": group_id},
        files={"scan_file": ("signed.png", png, "image/png")},
        follow_redirects=False,
    )

    assert uploaded.status_code == 303
    scan = workflow.db.fetch_one(
        "SELECT * FROM signed_response_scans WHERE response_letter_id = ?",
        (letter["id"],),
    )
    preview = client.get(f"/signed-response-scans/{scan['id']}")
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "application/pdf"
    assert "filename*=utf-8''" in preview.headers["content-disposition"]

    drawer = client.get(
        "/?tab=responses&response_view=created"
        f"&focus_letter={letter['id']}"
    )
    assert drawer.status_code == 200
    assert (
        f'src="/signed-response-scans/{scan["id"]}'
        '#page=1&amp;zoom=page-fit"' in drawer.text
    )
    assert "Открыть PDF" not in drawer.text
    assert 'type="checkbox"' not in drawer.text

    assert workflow.get_signed_response_scan(scan["id"])["status"] == "confirmed"


def test_confirmed_multi_page_scan_creates_one_idempotent_outlook_draft(
    workflow,
):
    workflow.initialize_gns_offices()
    workflow.initialize_gns_office_emails()
    _, letter = _ready_response_letter(workflow)
    workflow.set_outgoing_number(letter["id"], "9544")

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    stream = BytesIO()
    writer.write(stream)
    stream.seek(0)
    scan = workflow.register_signed_response_scan(
        letter["id"], "Подписанный ответ.pdf", stream
    )
    class DraftGateway:
        def __init__(self):
            self.requests = []
            self.missing_existing = False

        def create_draft(self, **request):
            assert request["attachment_path"].is_file()
            self.requests.append(request)
            if self.missing_existing and not request["create_if_missing"]:
                raise OutlookConnectionError("Черновик не найден")
            return OutlookDraftResult(
                entry_id="outlook-entry-1",
                recipient_email=request["recipient_email"],
                subject=request["subject"],
                attachment_name=request["attachment_name"],
                existing=len(self.requests) > 1,
            )

    gateway = DraftGateway()
    outgoing = OutlookOutgoingService(
        workflow.db,
        workflow.settings,
        OutlookService(gateway),
    )
    outgoing.update_subject_template(
        "Ответ № {outgoing_number} — {office_name}"
    )

    created = outgoing.create_draft(workflow, letter["id"])
    reopened = outgoing.create_draft(workflow, letter["id"])
    gateway.missing_existing = True
    with pytest.raises(OutlookConnectionError, match="не найден"):
        outgoing.create_draft(workflow, letter["id"])

    assert created["status"] == "draft_created"
    assert created["recipient_email"] == "002lenin@sti.gov.kg"
    assert created["subject"] == (
        "Ответ № 9544 — УГНС по Ленинскому району"
    )
    assert created["attachment_name"] == "Ответ_исх_9544.pdf"
    assert not created["existing_outlook_draft"]
    assert reopened["existing_outlook_draft"]
    assert len(gateway.requests) == 3
    assert gateway.requests[0]["draft_key"] == f"gns-scan-{scan['id']}"
    assert gateway.requests[0]["create_if_missing"] is True
    assert gateway.requests[1]["create_if_missing"] is False
    assert gateway.requests[2]["create_if_missing"] is False
    assert workflow.db.fetch_one(
        "SELECT status FROM outlook_outgoing_messages WHERE id = ?",
        (created["id"],),
    ) == {"status": "draft_created"}
    assert workflow.db.fetch_one(
        "SELECT COUNT(*) AS count FROM outlook_outgoing_messages"
    )["count"] == 1


def test_outlook_sent_reconciliation_updates_status_once(workflow, monkeypatch):
    workflow.initialize_gns_offices()
    workflow.initialize_gns_office_emails()
    _, letter = _ready_response_letter(workflow)
    workflow.set_outgoing_number(letter["id"], "9546")
    scan = workflow.register_signed_response_scan(
        letter["id"], "scan.png", _png_scan()
    )

    class Gateway:
        def __init__(self):
            self.draft_calls = 0
            self.sent_scans = 0
            self.fail_scan = True
            self.simulate_competing_update = False

        def create_draft(self, **request):
            self.draft_calls += 1
            return OutlookDraftResult(
                entry_id="draft-entry",
                recipient_email=request["recipient_email"],
                subject=request["subject"],
                attachment_name=request["attachment_name"],
            )

        def scan_sent_items(self, **request):
            self.sent_scans += 1
            if self.fail_scan:
                raise OutlookConnectionError("Outlook временно недоступен")
            candidate = request["candidates"][0]
            if self.simulate_competing_update:
                workflow.db.execute(
                    "UPDATE outlook_outgoing_messages SET status = 'sent' "
                    "WHERE draft_key = ?",
                    (candidate.draft_key,),
                )
            return OutlookSentScan(
                inspected_mail_count=1,
                candidate_mail_count=1,
                rejected_mail_count=0,
                ambiguous_candidate_count=0,
                matches=(
                    OutlookSentMatch(
                        draft_key=candidate.draft_key,
                        entry_id="sent-entry",
                        sent_at="2026-09-04T10:30:00+06:00",
                    ),
                ),
            )

    gateway = Gateway()
    outgoing = OutlookOutgoingService(
        workflow.db,
        workflow.settings,
        OutlookService(gateway),
    )
    created = outgoing.create_draft(workflow, letter["id"])
    with pytest.raises(OutlookConnectionError, match="временно недоступен"):
        outgoing.reconcile_sent_messages()
    unchanged = workflow.db.fetch_one(
        "SELECT status FROM outlook_outgoing_messages WHERE id = ?",
        (created["id"],),
    )
    gateway.fail_scan = False
    first = outgoing.reconcile_sent_messages()
    second = outgoing.reconcile_sent_messages()

    stored = workflow.db.fetch_one(
        "SELECT status, outlook_entry_id, sent_at, draft_key "
        "FROM outlook_outgoing_messages WHERE id = ?",
        (created["id"],),
    )
    workflow.db.execute(
        "UPDATE outlook_outgoing_messages SET status = 'draft_created' "
        "WHERE id = ?",
        (created["id"],),
    )
    gateway.simulate_competing_update = True
    raced = outgoing.reconcile_sent_messages()
    audits = workflow.db.fetch_one(
        "SELECT COUNT(*) AS count FROM audit_events "
        "WHERE entity_id = ? AND event_type = 'outlook_sent_confirmed'",
        (created["id"],),
    )
    assert first["confirmed"] == 1
    assert unchanged == {"status": "draft_created"}
    assert second["pending"] == 0
    assert raced["confirmed"] == 0
    assert gateway.sent_scans == 3
    assert stored == {
        "status": "sent",
        "outlook_entry_id": "sent-entry",
        "sent_at": "2026-09-04T10:30:00+06:00",
        "draft_key": f"gns-scan-{scan['id']}",
    }
    assert audits["count"] == 1
    with pytest.raises(OutlookIntegrationError, match="уже отправлено"):
        outgoing.create_draft(workflow, letter["id"])
    assert gateway.draft_calls == 1

    from starlette.testclient import TestClient

    from gns_app import main

    monkeypatch.setattr(main, "workflow", workflow)
    client = TestClient(main.app)
    page = client.get(
        "/?tab=responses&response_view=created"
    )
    assert page.status_code == 200
    assert "Отправлено" in page.text
    assert "2026-09-04 10:30" in page.text
    assert "Подтверждено Outlook" in page.text
    assert "Подготовить письмо в Outlook" not in page.text
    history = client.get("/history")
    assert history.status_code == 200
    assert "Отправлено" in history.text
    assert "2026-09-04 10:30" in history.text


def test_repeat_draft_reuses_exact_confirmed_scan_and_is_confirmed(workflow):
    workflow.initialize_gns_offices()
    workflow.initialize_gns_office_emails()
    _, letter = _ready_response_letter(workflow)
    workflow.set_outgoing_number(letter["id"], "9550")
    scan = workflow.register_signed_response_scan(
        letter["id"], "scan.png", _png_scan()
    )
    expected_pdf = Path(scan["pdf_path"]).read_bytes()

    class Gateway:
        def __init__(self):
            self.requests = []
            self.attachments = []

        def create_draft(self, **request):
            self.requests.append(request)
            self.attachments.append(request["attachment_path"].read_bytes())
            return OutlookDraftResult(
                entry_id=f"draft-{len(self.requests)}",
                recipient_email=request["recipient_email"],
                subject=request["subject"],
                attachment_name=request["attachment_name"],
                existing=not request["create_if_missing"],
            )

        def scan_sent_items(self, **request):
            candidate = request["candidates"][0]
            return OutlookSentScan(
                inspected_mail_count=1,
                candidate_mail_count=1,
                rejected_mail_count=0,
                ambiguous_candidate_count=0,
                matches=(OutlookSentMatch(
                    draft_key=candidate.draft_key,
                    entry_id="resent-entry",
                    sent_at="2026-09-04T11:00:00+06:00",
                ),),
            )

    gateway = Gateway()
    outgoing = OutlookOutgoingService(
        workflow.db, workflow.settings, OutlookService(gateway)
    )
    original = outgoing.create_draft(workflow, letter["id"])
    workflow.db.execute(
        "UPDATE outlook_outgoing_messages SET status = 'sent', sent_at = ? "
        "WHERE id = ?",
        ("2026-09-04T10:30:00+06:00", original["id"]),
    )

    first = outgoing.create_resend_draft(workflow, original["id"])
    reopened = outgoing.create_resend_draft(workflow, original["id"])
    confirmed = outgoing.reconcile_sent_messages()
    stored = workflow.db.fetch_one(
        "SELECT resend_sequence, resend_status, resend_draft_key, "
        "resend_outlook_entry_id, resent_at FROM outlook_outgoing_messages "
        "WHERE id = ?",
        (original["id"],),
    )
    second = outgoing.create_resend_draft(workflow, original["id"])

    assert first["resend_sequence"] == 1
    assert not first["existing_outlook_draft"]
    assert reopened["existing_outlook_draft"]
    assert confirmed["confirmed"] == 1
    assert stored["resend_status"] == "sent"
    assert stored["resend_outlook_entry_id"] == "resent-entry"
    assert stored["resent_at"] == "2026-09-04T11:00:00+06:00"
    assert second["resend_sequence"] == 2
    assert gateway.requests[1]["create_if_missing"] is True
    assert gateway.requests[2]["create_if_missing"] is False
    assert gateway.requests[3]["create_if_missing"] is True
    assert gateway.attachments == [expected_pdf] * 4


def test_outlook_draft_requires_confirmed_scan(workflow):
    workflow.initialize_gns_offices()
    workflow.initialize_gns_office_emails()
    _, letter = _ready_response_letter(workflow)
    workflow.set_outgoing_number(letter["id"], "9545")
    workflow.register_signed_response_scan(
        letter["id"], "scan.png", _png_scan()
    )
    outgoing = OutlookOutgoingService(
        workflow.db,
        workflow.settings,
        OutlookService(object()),
    )

    with pytest.raises(OutlookIntegrationError, match="подтвердите"):
        outgoing.create_draft(workflow, letter["id"])


def test_found_abs_questionnaire_waits_for_manual_account_result(workflow):
    case_id = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )
    workflow.db.execute(
        "UPDATE taxpayers SET abs_result = 'found' WHERE case_id = ?",
        (case_id,),
    )
    workflow.db.execute(
        "UPDATE cases SET status = 'needs_review', abs_status = 'found' "
        "WHERE id = ?",
        (case_id,),
    )

    result = workflow.confirm_abs_account_taxpayer(
        case_id, "12345678901234", "not_found"
    )

    assert result["next_status"] == "ready_for_response"
    taxpayer = workflow.get_taxpayers(case_id)[0]
    assert taxpayer["abs_result"] == "found"
    assert taxpayer["abs_account_result"] == "not_found"
    assert workflow.get_case(case_id)["status"] == "ready_for_response"


def test_abs_account_task_moves_to_manual_response_without_duplicate(workflow):
    case_id = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )
    workflow.db.execute(
        "UPDATE taxpayers SET abs_result = 'found' WHERE case_id = ?",
        (case_id,),
    )
    workflow.db.execute(
        "UPDATE cases SET status = 'needs_review', abs_status = 'found' "
        "WHERE id = ?",
        (case_id,),
    )

    before = workflow.manual_review_overview()
    assert len(before["account_groups"]) == 1
    assert not workflow.manual_response_groups()

    workflow.confirm_abs_account_taxpayer(
        case_id,
        "12345678901234",
        "found",
    )

    after = workflow.manual_review_overview()
    manual = workflow.manual_response_groups()
    assert not after["account_groups"]
    assert len(manual) == 1
    assert manual[0]["taxpayers"][0]["source_case_id"] == case_id


def test_real_abs_account_counts_are_persisted_without_auto_confirmation(
    workflow, monkeypatch
):
    from bs4 import BeautifulSoup
    from starlette.testclient import TestClient

    from gns_app import main

    class AccountSummaryGateway:
        is_fake = False
        supports_session = False

        @staticmethod
        def check(username, password, taxpayers):
            assert username == "employee"
            assert password == "one-time-secret"
            return AbsCheckResult(
                status=AbsStatus.FOUND,
                taxpayers=[
                    {
                        "inn": taxpayers[0]["inn"],
                        "name": taxpayers[0]["name"],
                        "result": AbsStatus.FOUND,
                        "active_account_count": 2,
                        "closed_account_count": 1,
                    }
                ],
                message="Анкета найдена.",
                is_fake=False,
            )

    workflow.abs = AccountSummaryGateway()
    case_id = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )

    workflow.check_abs(case_id, "employee", "one-time-secret")

    taxpayer = workflow.get_taxpayers(case_id)[0]
    assert taxpayer["abs_active_account_count"] == 2
    assert taxpayer["abs_closed_account_count"] == 1
    assert taxpayer["abs_account_result"] is None
    assert workflow.get_case(case_id)["status"] == "needs_review"
    overview_taxpayer = workflow.today_overview()["found_groups"][0][
        "taxpayers"
    ][0]
    assert overview_taxpayer["abs_active_account_count"] == 2
    assert overview_taxpayer["abs_closed_account_count"] == 1
    monkeypatch.setattr(main, "workflow", workflow)
    client = TestClient(main.app)
    response = client.get("/?tab=review")
    assert response.status_code == 200
    visible_text = " ".join(
        BeautifulSoup(response.text, "html.parser")
        .get_text(" ", strip=True)
        .split()
    )
    assert "Активных счетов: 2 · закрытых: 1" in visible_text
    assert "Расчётный счёт" in response.text
    assert "task-row-abs-found" in response.text
    assert "АБС · НАЙДЕН" in response.text
    assert 'name="account_result" value="found"' in response.text
    assert 'name="account_result" value="not_found"' in response.text

    confirmed = client.post(
        f"/cases/{case_id}/abs-account/taxpayer",
        data={
            "taxpayer_inn": "12345678901234",
            "account_result": "found",
            "return_to": "review",
        },
        follow_redirects=False,
    )
    assert confirmed.status_code == 303
    assert not workflow.manual_review_overview()["account_groups"]
    manual = workflow.manual_response_groups()
    assert len(manual) == 1
    assert manual[0]["taxpayers"][0]["inn"] == "12345678901234"
    manual_page = client.get("/?tab=responses&response_view=manual")
    assert manual_page.status_code == 200
    assert "12345678901234" in manual_page.text
    assert "response-row-abs-found" in manual_page.text
    assert "АБС: счёт найден" in manual_page.text


def test_found_abs_questionnaire_with_account_stays_manual(workflow):
    case_id = _insert_ready_case(
        workflow, "12345678901234", 'ОсОО "Первый"'
    )
    workflow.db.execute(
        "UPDATE taxpayers SET abs_result = 'found' WHERE case_id = ?",
        (case_id,),
    )
    workflow.db.execute(
        "UPDATE cases SET status = 'needs_review' WHERE id = ?",
        (case_id,),
    )

    result = workflow.confirm_abs_account_taxpayer(
        case_id, "12345678901234", "found"
    )

    assert result["next_status"] == "needs_review"


def test_office_spelling_aliases_share_one_response_group(workflow):
    workflow.initialize_gns_offices()
    first = _insert_ready_case(
        workflow,
        "12345678901234",
        'ОсОО "Первый"',
        district_place="по г. Балыкчы Иссык-Кульской области",
    )
    second = _insert_ready_case(
        workflow,
        "23456789012345",
        'ОсОО "Второй"',
        district_place="по городу Балыкчы Ысык-Кульской области",
    )
    for case_id in (first, second):
        workflow.db.execute(
            "UPDATE cases SET status = 'ready_for_response' WHERE id = ?",
            (case_id,),
        )
        workflow.db.execute(
            "UPDATE taxpayers SET abs_result = 'not_found' WHERE case_id = ?",
            (case_id,),
        )

    groups = workflow.today_overview()["not_found_groups"]

    assert len(groups) == 1
    assert groups[0]["taxpayer_count"] == 2
    assert groups[0]["gns_office_key"]
