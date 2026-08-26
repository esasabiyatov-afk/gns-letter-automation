from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfWriter

from gns_app.config import Settings
from gns_app.database import Database
from gns_app.services.workflow import WorkflowService


@pytest.fixture
def project_root() -> Path:
    return Path(__file__).resolve().parents[1]


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """Create a harmless two-page PDF used by workflow tests."""
    path = tmp_path / "sample-letter.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    with path.open("wb") as stream:
        writer.write(stream)
    return path


@pytest.fixture
def templates_dir(tmp_path: Path) -> Path:
    """Generate minimal Word templates without private bank documents."""
    directory = tmp_path / "templates"
    directory.mkdir()
    for filename in (
        "шаблон ответа одиночный.docx",
        "шаблон ответа много.docx",
    ):
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        normal = document.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        document.add_paragraph("04-1/______")
        document.add_paragraph("[Дата.Сегодня]")
        document.add_paragraph("[Район.Место]")
        document.add_paragraph("[Должность.Отправитель]")
        document.add_paragraph("[ФИО. Отправитель]")
        document.add_paragraph(
            "Настоящим ЗАО АКБ сообщает сведения по вашему запросу."
        )
        document.add_paragraph("[Перечисления.Субьект] [ИНН.Субьект]")
        for _ in range(7):
            document.add_paragraph(
                "Ответ подготовлен на основании доступных банковских данных."
            )
        document.add_paragraph("Исполнитель: [ФИО.Исп]")
        document.save(directory / filename)
    return directory


@pytest.fixture
def test_settings(
    tmp_path: Path,
    project_root: Path,
    templates_dir: Path,
) -> Settings:
    runtime = tmp_path / "runtime"
    settings = Settings(
        project_root=project_root,
        runtime_dir=runtime,
        database_path=runtime / "test.sqlite3",
        uploads_dir=runtime / "uploads",
        previews_dir=runtime / "previews",
        official_dir=runtime / "official",
        responses_dir=runtime / "responses",
        inbox_dir=runtime / "inbox",
        source_templates_dir=templates_dir,
        period_threshold=date(2019, 1, 1),
        max_upload_bytes=150 * 1024 * 1024,
        allowed_qr_hosts=frozenset({"qr.salyk.kg"}),
        allowed_qr_paths=frozenset({"/getsti010decission"}),
        auto_download_official=False,
        employee_name="Гапарова Э.",
        auto_registry_check=False,
    )
    settings.ensure_directories()
    return settings


@pytest.fixture
def workflow(test_settings: Settings) -> WorkflowService:
    database = Database(test_settings.database_path)
    database.initialize()
    return WorkflowService(database, test_settings)
