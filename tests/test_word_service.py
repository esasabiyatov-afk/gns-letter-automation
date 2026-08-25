from __future__ import annotations

import shutil
import zipfile

import pytest
from docx import Document

from gns_app.services.word_service import WordTemplateError, WordTemplateService


def _all_word_text(path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    with zipfile.ZipFile(path) as archive:
        parts.extend(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )
    return "\n".join(parts)


def test_renders_single_response_without_placeholders(
    tmp_path, project_root
):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "single.docx"
    service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        [{"name": 'ОсОО "Тест"', "inn": "02312201410117"}],
    )
    text = _all_word_text(output)
    assert output.exists()
    assert "02312201410117" in text
    assert 'ОсОО "Тест"' in text
    assert "Заместителю начальника управления" in text
    assert "Зам. начальника управления" not in text
    assert "[Дата.Сегодня]" not in text
    assert "[ФИО.Исп]" not in text


def test_render_returns_overflow_flag_alongside_path(tmp_path, project_root):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "single.docx"

    result_path, likely_overflow = service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        [{"name": 'ОсОО "Тест"', "inn": "02312201410117"}],
    )

    assert result_path == output
    assert likely_overflow is False


def test_outgoing_number_replaces_legacy_template_placeholder(
    tmp_path, project_root
):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "numbered.docx"

    service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
            "outgoing_number": "9544",
        },
        [{"name": 'ОсОО "Тест"', "inn": "02312201410117"}],
    )

    text = _all_word_text(output)
    assert "04-1/9544" in text
    assert "[Исх.Номер]" not in text
    assert "______" not in text


def test_explicit_outgoing_number_tag_is_replaced(tmp_path, project_root):
    templates = tmp_path / "templates"
    templates.mkdir()
    source = project_root / "УГНС" / "шаблон ответа одиночный.docx"
    derived = templates / source.name
    shutil.copy2(source, derived)
    document = Document(str(derived))
    marker = next(
        paragraph for paragraph in document.paragraphs if "______" in paragraph.text
    )
    WordTemplateService._replace_across_runs(
        marker,
        "______",
        WordTemplateService.TOKENS["outgoing_number"],
    )
    document.save(derived)

    output = tmp_path / "tag-numbered.docx"
    WordTemplateService(templates).render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
            "outgoing_number": "9544",
        },
        [{"name": 'ОсОО "Тест"', "inn": "02312201410117"}],
    )

    text = _all_word_text(output)
    assert "04-1/9544" in text
    assert "[Исх.Номер]" not in text


def test_short_letter_is_not_flagged_as_overflowing(tmp_path, project_root):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "short.docx"

    _, likely_overflow = service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        [{"name": 'ОсОО "Тест"', "inn": "02312201410117"}],
    )

    assert likely_overflow is False


def test_long_taxpayer_list_is_flagged_as_likely_overflowing(
    tmp_path, project_root
):
    # Откалибровано по реальным замерам: при рендере через LibreOffice
    # именно на 14-м налогоплательщике письмо реально перестаёт помещаться
    # на одну страницу (13 - ещё влезает, 14 - уже нет). Берём заведомо
    # длинный список, чтобы не зависеть от точной границы в один
    # налогоплательщик.
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "long.docx"
    taxpayers = [
        {"name": f'ОсОО "Тестовая Компания Номер {i}"', "inn": f"{i:014d}"}
        for i in range(1, 31)
    ]

    _, likely_overflow = service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        taxpayers,
    )

    assert likely_overflow is True


def test_render_pages_reports_overflow_when_a_chunk_itself_overflows(
    tmp_path, project_root
):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "chunked.docx"
    taxpayers = [
        {"name": f'ОсОО "Тестовая Компания Номер {i}"', "inn": f"{i:014d}"}
        for i in range(1, 31)
    ]

    # Каждый кусок по 30 налогоплательщиков сам по себе не влезет на
    # одну страницу, даже несмотря на явный разрыв страницы между кусками.
    _, likely_overflow = service.render_pages(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        taxpayers,
        taxpayers_per_page=30,
    )

    assert likely_overflow is True


def test_render_pages_does_not_flag_overflow_when_split_keeps_each_chunk_short(
    tmp_path, project_root
):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "chunked_short.docx"
    taxpayers = [
        {"name": f'ОсОО "Тестовая Компания Номер {i}"', "inn": f"{i:014d}"}
        for i in range(1, 31)
    ]

    # Тот же список, но разбит на короткие куски по 5 - каждый кусок
    # заведомо помещается на одну страницу.
    _, likely_overflow = service.render_pages(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        taxpayers,
        taxpayers_per_page=5,
    )

    assert likely_overflow is False


def test_renders_multiple_taxpayers_on_separate_lines(
    tmp_path, project_root
):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "multi.docx"
    service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Зам. начальника управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        [
            {"name": 'ОсОО "Первый"', "inn": "02312201410117"},
            {"name": 'ОсОО "Второй"', "inn": "02801202210185"},
        ],
    )
    text = _all_word_text(output)
    assert output.exists()
    assert "02312201410117" in text
    assert "02801202210185" in text
    assert "Заместителю начальника управления" in text
    assert "[Перечисления.Субьект]" not in text
    assert "[ИНН.Субьект]" not in text


def test_response_adds_ip_prefix_for_person(tmp_path, project_root):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "person.docx"

    service.render(
        output,
        {
            "district_place": "по Ленинскому району города Бишкек",
            "recipient_position": "Начальник управления",
            "recipient_display_name": "Телтаеву Р. З.",
            "employee_name": "Гапарова Э.",
        },
        [{"name": "Ишен кызы Саида", "inn": "10207200101109"}],
    )

    text = _all_word_text(output)
    assert "ИП Ишен кызы Саида" in text
    assert "ИП ИП" not in text


def test_split_response_repeats_full_letter_on_new_page(tmp_path, project_root):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "split.docx"
    case = {
        "district_place": "по Ленинскому району города Бишкек",
        "recipient_position": "Зам. начальника управления",
        "recipient_display_name": "Телтаеву Р. З.",
        "employee_name": "Гапарова Э.",
    }
    taxpayers = [
        {"name": f'ОсОО "Тест {index}"', "inn": f"{index:014d}"}
        for index in range(1, 5)
    ]

    service.render_pages(output, case, taxpayers, taxpayers_per_page=2)

    text = _all_word_text(output)
    assert text.count("по Ленинскому району города Бишкек") >= 2
    assert text.count("Настоящим ЗАО АКБ") >= 2
    assert all(item["inn"] in text for item in taxpayers)


def test_split_response_uses_distinct_outgoing_numbers(tmp_path, project_root):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "split-numbered.docx"
    case = {
        "district_place": "по Ленинскому району города Бишкек",
        "recipient_position": "Зам. начальника управления",
        "recipient_display_name": "Телтаеву Р. З.",
        "employee_name": "Гапарова Э.",
    }
    taxpayers = [
        {"name": f'ОсОО "Тест {index}"', "inn": f"{index:014d}"}
        for index in range(1, 5)
    ]

    service.render_pages(
        output,
        case,
        taxpayers,
        taxpayers_per_page=2,
        outgoing_numbers=["9544", "9545"],
    )

    text = _all_word_text(output)
    assert "04-1/9544" in text
    assert "04-1/9545" in text


def test_split_response_rejects_one_number_for_multiple_letters(
    tmp_path, project_root
):
    service = WordTemplateService(project_root / "УГНС")
    output = tmp_path / "split-invalid-number.docx"
    case = {
        "district_place": "по Ленинскому району города Бишкек",
        "recipient_position": "Зам. начальника управления",
        "recipient_display_name": "Телтаеву Р. З.",
        "employee_name": "Гапарова Э.",
        "outgoing_number": "9544",
    }
    taxpayers = [
        {"name": f'ОсОО "Тест {index}"', "inn": f"{index:014d}"}
        for index in range(1, 5)
    ]

    with pytest.raises(WordTemplateError, match="отдельный исходящий номер"):
        service.render_pages(
            output,
            case,
            taxpayers,
            taxpayers_per_page=2,
        )
