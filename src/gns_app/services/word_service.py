from __future__ import annotations

import math
import os
import re
import zipfile
from copy import deepcopy
from datetime import date
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Length
from docx.text.paragraph import Paragraph
from lxml import etree

from gns_app.services.name_service import NameService
from gns_app.services.taxpayer_service import response_taxpayer_name


class WordTemplateError(ValueError):
    pass


# --- Оценка переполнения на вторую страницу ---------------------------
#
# python-docx не умеет верстать документ и не знает, сколько страниц
# получится - это знает только сам Word при открытии файла. Здесь -
# приближённая (эвристическая) оценка высоты содержимого письма в пунктах,
# калиброванная по РЕАЛЬНЫМ замерам: сгенерированные этим сервисом письма
# были превращены в PDF через LibreOffice, и координаты каждой строки
# текста были извлечены и сопоставлены с ожидаемой высотой строки.
# Такая калибровка учитывает реальный шрифт и вёрстку именно этих двух
# шаблонов ("шаблон ответа одиночный.docx" и "шаблон ответа много.docx"),
# а не абстрактную типографику.
#
# Оценка не может быть идеально точной (это не полноценный движок вёрстки),
# поэтому предупреждение показывается с запасом: даже "пограничные" случаи,
# где реального переполнения может и не быть, помечаются как "могут не
# поместиться" - лучше лишний раз предложить проверить глазами, чем
# промолчать про реальное переполнение.
DEFAULT_FONT_PT = 12.0
DEFAULT_LINE_SPACING_MULTIPLE = 1.0
AVG_CHAR_WIDTH_FACTOR = 0.62
NATURAL_LINE_FACTOR = 16.10 / (12.0 * 1.15)
OVERFLOW_WARNING_RATIO = 0.92


def _resolve_font_size_pt(paragraph: Paragraph) -> float:
    for run in paragraph.runs:
        if run.font.size:
            return run.font.size.pt
    style = paragraph.style
    while style is not None:
        if style.font.size:
            return style.font.size.pt
        style = style.base_style
    return DEFAULT_FONT_PT


def _resolve_line_spacing(paragraph: Paragraph):
    pf = paragraph.paragraph_format
    spacing = pf.line_spacing
    if spacing is not None:
        return spacing
    style = paragraph.style
    while style is not None:
        spacing = style.paragraph_format.line_spacing
        if spacing is not None:
            return spacing
        style = style.base_style
    return DEFAULT_LINE_SPACING_MULTIPLE


def _resolve_space_pt(getter, paragraph: Paragraph, font_size_pt: float) -> float:
    value = getter(paragraph.paragraph_format)
    if value is not None:
        return value.pt if isinstance(value, Length) else float(value) * font_size_pt
    style = paragraph.style
    while style is not None:
        value = getter(style.paragraph_format)
        if value is not None:
            return (
                value.pt
                if isinstance(value, Length)
                else float(value) * font_size_pt
            )
        style = style.base_style
    return 0.0


def _resolve_indent_pt(getter, paragraph: Paragraph) -> float:
    value = getter(paragraph.paragraph_format)
    if value is not None:
        return value.pt
    style = paragraph.style
    while style is not None:
        value = getter(style.paragraph_format)
        if value is not None:
            return value.pt
        style = style.base_style
    return 0.0


def _line_height_pt(paragraph: Paragraph) -> float:
    font_size_pt = _resolve_font_size_pt(paragraph)
    spacing = _resolve_line_spacing(paragraph)
    if isinstance(spacing, Length):
        return spacing.pt
    return font_size_pt * NATURAL_LINE_FACTOR * float(spacing)


def _wrap_line_count(text: str, chars_per_line: float) -> int:
    """Перенос по словам (как в реальном текстовом процессоре), а не
    деление длины текста на ширину строки - иначе для узких колонок
    (например, адресного блока с большим отступом) число строк сильно
    занижается: слово целиком не влезает - переносится целиком, а не
    "дозаполняет" строку по числу символов."""
    if not text:
        return 1
    words = text.split(" ")
    lines = 1
    current_len = 0
    for word in words:
        word_len = len(word)
        if current_len == 0:
            current_len = word_len
        elif current_len + 1 + word_len <= chars_per_line:
            current_len += 1 + word_len
        else:
            lines += 1
            current_len = word_len
        while current_len > chars_per_line:
            current_len -= chars_per_line
            lines += 1
    return lines


def _estimate_paragraph_height_pt(
    paragraph: Paragraph, section_usable_width_pt: float
) -> float:
    font_size_pt = _resolve_font_size_pt(paragraph)
    line_height_pt = _line_height_pt(paragraph)
    space_before = _resolve_space_pt(
        lambda pf: pf.space_before, paragraph, font_size_pt
    )
    space_after = _resolve_space_pt(
        lambda pf: pf.space_after, paragraph, font_size_pt
    )

    left_indent = _resolve_indent_pt(lambda pf: pf.left_indent, paragraph)
    right_indent = _resolve_indent_pt(lambda pf: pf.right_indent, paragraph)
    usable_width_pt = max(
        20.0, section_usable_width_pt - left_indent - right_indent
    )
    chars_per_line = max(
        1.0, usable_width_pt / (AVG_CHAR_WIDTH_FACTOR * font_size_pt)
    )

    text = paragraph.text
    sub_lines = text.split("\n") if text else [""]
    total_lines = sum(_wrap_line_count(s, chars_per_line) for s in sub_lines)
    return space_before + space_after + total_lines * line_height_pt


def estimate_single_page_overflow(document: Document) -> bool:
    """True, если письмо, вероятнее всего, не уместится на одной странице.

    Это приближённая оценка (см. пояснение к константам выше), а не точный
    подсчёт разрывов страниц - при сомнении она скорее предупредит лишний
    раз, чем промолчит о реальном переполнении.
    """
    section = document.sections[0]
    usable_width_pt = (
        section.page_width.pt - section.left_margin.pt - section.right_margin.pt
    )
    usable_height_pt = (
        section.page_height.pt - section.top_margin.pt - section.bottom_margin.pt
    )
    total_height_pt = sum(
        _estimate_paragraph_height_pt(paragraph, usable_width_pt)
        for paragraph in document.paragraphs
    )
    if usable_height_pt <= 0:
        return False
    return (total_height_pt / usable_height_pt) > OVERFLOW_WARNING_RATIO


MONTHS_RU = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)


class WordTemplateService:
    TOKENS = {
        "today": "[Дата.Сегодня]",
        "district": "[Район.Место]",
        "position": "[Должность.Отправитель]",
        "recipient": "[ФИО. Отправитель]",
        "subjects": "[Перечисления.Субьект]",
        "inns": "[ИНН.Субьект]",
        "employee": "[ФИО.Исп]",
    }

    def __init__(
        self,
        templates_dir: Path,
        name_service: NameService | None = None,
    ):
        self.templates_dir = templates_dir
        self.names = name_service or NameService()

    def render(
        self,
        output_path: Path,
        case: dict[str, str],
        taxpayers: list[dict[str, str]],
    ) -> tuple[Path, bool]:
        self._validate(case, taxpayers)
        template_name = (
            "шаблон ответа одиночный.docx"
            if len(taxpayers) == 1
            else "шаблон ответа много.docx"
        )
        template_path = self.templates_dir / template_name
        if not template_path.exists():
            raise WordTemplateError(f"Не найден шаблон: {template_name}")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        working_path = output_path.with_suffix(".working.docx")
        document = Document(str(template_path))

        replacements = {
            self.TOKENS["today"]: self._format_date(date.today()),
            self.TOKENS["district"]: case["district_place"].strip(),
            self.TOKENS["position"]: self.names.position_display(
                case["recipient_position"]
            ),
            self.TOKENS["recipient"]: case["recipient_display_name"].strip(),
            self.TOKENS["employee"]: case["employee_name"].strip(),
        }

        if len(taxpayers) == 1:
            taxpayer = taxpayers[0]
            replacements[self.TOKENS["subjects"]] = (
                f" {response_taxpayer_name(taxpayer['name'], taxpayer['inn'])}"
            )
            replacements[self.TOKENS["inns"]] = (
                f"ИНН: {taxpayer['inn'].strip()} "
            )
        else:
            self._replace_multi_taxpayer_paragraph(document, taxpayers)
            replacements[self.TOKENS["subjects"]] = self._plain_taxpayer_list(
                taxpayers
            )
            replacements[self.TOKENS["inns"]] = ""

        for paragraph in self._iter_visible_paragraphs(document):
            for token, value in replacements.items():
                self._replace_across_runs(paragraph, token, value)

        likely_overflow = estimate_single_page_overflow(document)

        document.save(working_path)
        self._patch_all_xml(working_path, replacements)
        self._assert_no_tokens(working_path)
        os.replace(working_path, output_path)
        return output_path, likely_overflow

    def render_pages(
        self,
        output_path: Path,
        case: dict[str, str],
        taxpayers: list[dict[str, str]],
        taxpayers_per_page: int,
    ) -> tuple[Path, bool]:
        if taxpayers_per_page <= 0 or taxpayers_per_page >= len(taxpayers):
            return self.render(output_path, case, taxpayers)
        chunks = [
            taxpayers[index : index + taxpayers_per_page]
            for index in range(0, len(taxpayers), taxpayers_per_page)
        ]
        temporary_paths: list[Path] = []
        any_chunk_overflows = False
        output_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            for index, chunk in enumerate(chunks, 1):
                temporary = output_path.with_name(
                    f".{output_path.stem}-page-{index}-{uuid4().hex}.docx"
                )
                _, chunk_overflow = self.render(temporary, case, chunk)
                any_chunk_overflows = any_chunk_overflows or chunk_overflow
                temporary_paths.append(temporary)

            combined = Document(str(temporary_paths[0]))
            for temporary in temporary_paths[1:]:
                combined.add_page_break()
                source = Document(str(temporary))
                for element in source.element.body:
                    if element.tag.endswith("}sectPr"):
                        continue
                    combined.element.body.insert(-1, deepcopy(element))

            working_path = output_path.with_suffix(".working.docx")
            combined.save(working_path)
            os.replace(working_path, output_path)
            return output_path, any_chunk_overflows
        finally:
            for temporary in temporary_paths:
                temporary.unlink(missing_ok=True)

    @staticmethod
    def _validate(
        case: dict[str, str], taxpayers: list[dict[str, str]]
    ) -> None:
        required = (
            "district_place",
            "recipient_position",
            "recipient_display_name",
            "employee_name",
        )
        missing = [field for field in required if not case.get(field)]
        if missing:
            raise WordTemplateError(
                "Не заполнены обязательные поля: " + ", ".join(missing)
            )
        if not taxpayers:
            raise WordTemplateError("Нет подтверждённых налогоплательщиков")
        for taxpayer in taxpayers:
            inn = re.sub(r"\D", "", taxpayer.get("inn", ""))
            if len(inn) != 14 or not taxpayer.get("name", "").strip():
                raise WordTemplateError(
                    "У каждого налогоплательщика нужны наименование и 14-значный ИНН"
                )

    @staticmethod
    def _format_date(value: date) -> str:
        return f"«{value.day:02d}» {MONTHS_RU[value.month]} {value.year} г."

    def _replace_multi_taxpayer_paragraph(
        self,
        document: Document,
        taxpayers: list[dict[str, str]],
    ) -> None:
        target = next(
            (
                paragraph
                for paragraph in document.paragraphs
                if self.TOKENS["subjects"] in paragraph.text
                and self.TOKENS["inns"] in paragraph.text
            ),
            None,
        )
        if target is None:
            raise WordTemplateError(
                "В множественном шаблоне не найдена строка налогоплательщиков"
            )

        reference_properties = None
        if target.runs and target.runs[0]._r.rPr is not None:
            reference_properties = deepcopy(target.runs[0]._r.rPr)

        target.clear()
        for index, taxpayer in enumerate(taxpayers, 1):
            if index > 1:
                target.add_run().add_break()
            run = target.add_run(
                f"{index}. {response_taxpayer_name(taxpayer['name'], taxpayer['inn'])} "
                f"ИНН: {taxpayer['inn'].strip()};"
            )
            if reference_properties is not None:
                run._r.insert(0, deepcopy(reference_properties))

    @staticmethod
    def _plain_taxpayer_list(taxpayers: list[dict[str, str]]) -> str:
        return "; ".join(
            f"{index}. {response_taxpayer_name(item['name'], item['inn'])} "
            f"ИНН: {item['inn'].strip()}"
            for index, item in enumerate(taxpayers, 1)
        )

    @staticmethod
    def _iter_visible_paragraphs(document: Document):
        yield from document.paragraphs
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for section in document.sections:
            for part in (section.header, section.footer):
                yield from part.paragraphs
                for table in part.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            yield from cell.paragraphs

    @staticmethod
    def _replace_across_runs(
        paragraph: Paragraph,
        token: str,
        replacement: str,
    ) -> None:
        while token in "".join(run.text for run in paragraph.runs):
            full_text = "".join(run.text for run in paragraph.runs)
            start = full_text.index(token)
            end = start + len(token)
            cursor = 0
            inserted = False
            for run in paragraph.runs:
                run_start = cursor
                run_end = cursor + len(run.text)
                cursor = run_end
                if run_end <= start or run_start >= end:
                    continue

                local_start = max(0, start - run_start)
                local_end = min(len(run.text), end - run_start)
                prefix = run.text[:local_start]
                suffix = run.text[local_end:]
                if not inserted:
                    run.text = prefix + replacement + suffix
                    inserted = True
                else:
                    run.text = prefix + suffix

    def _patch_all_xml(
        self, docx_path: Path, replacements: dict[str, str]
    ) -> None:
        temporary = docx_path.with_name(
            f"{docx_path.stem}.{uuid4().hex}.zip"
        )
        namespace = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        }

        with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(
            temporary, "w", zipfile.ZIP_DEFLATED
        ) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename.startswith("word/") and item.filename.endswith(
                    ".xml"
                ):
                    try:
                        root = etree.fromstring(data)
                        for paragraph in root.xpath(
                            ".//w:p", namespaces=namespace
                        ):
                            text_nodes = paragraph.xpath(
                                ".//w:t", namespaces=namespace
                            )
                            for token, value in replacements.items():
                                self._replace_xml_nodes(
                                    text_nodes, token, value
                                )
                        data = etree.tostring(
                            root,
                            xml_declaration=True,
                            encoding="UTF-8",
                            standalone=True,
                        )
                    except etree.XMLSyntaxError:
                        pass
                target.writestr(item, data)

        os.replace(temporary, docx_path)

    @staticmethod
    def _replace_xml_nodes(
        nodes: list[etree._Element],
        token: str,
        replacement: str,
    ) -> None:
        while token in "".join(node.text or "" for node in nodes):
            full = "".join(node.text or "" for node in nodes)
            start = full.index(token)
            end = start + len(token)
            cursor = 0
            inserted = False
            for node in nodes:
                value = node.text or ""
                node_start = cursor
                node_end = cursor + len(value)
                cursor = node_end
                if node_end <= start or node_start >= end:
                    continue
                local_start = max(0, start - node_start)
                local_end = min(len(value), end - node_start)
                prefix = value[:local_start]
                suffix = value[local_end:]
                if not inserted:
                    node.text = prefix + replacement + suffix
                    inserted = True
                else:
                    node.text = prefix + suffix

    def _assert_no_tokens(self, docx_path: Path) -> None:
        with zipfile.ZipFile(docx_path) as archive:
            remaining: set[str] = set()
            for name in archive.namelist():
                if not (name.startswith("word/") and name.endswith(".xml")):
                    continue
                text = archive.read(name).decode("utf-8", errors="ignore")
                for token in self.TOKENS.values():
                    if token in text:
                        remaining.add(token)
        if remaining:
            raise WordTemplateError(
                "В ответе остались незаполненные коды: "
                + ", ".join(sorted(remaining))
            )
