from __future__ import annotations

import hashlib
import csv
import json
import re
import shutil
import sqlite3
import time as monotonic_time
from concurrent.futures import ThreadPoolExecutor, as_completed
from difflib import SequenceMatcher
from datetime import UTC, date, datetime, time, timedelta, timezone
from pathlib import Path
from threading import RLock
from typing import Any, BinaryIO, Callable
from uuid import uuid4

from docx import Document
from PIL import Image, ImageSequence, UnidentifiedImageError
from pypdf import PdfReader

from gns_app.config import Settings
from gns_app.database import Database, utc_now
from gns_app.diagnostics import record_event, record_exception
from gns_app.domain import (
    AbsCheckResult,
    AbsStatus,
    CaseStatus,
    OdbStatus,
    OcrResult,
    OcrStatus,
    PageStatus,
    PageType,
    QrStatus,
    UploadStatus,
    ValueSource,
)
from gns_app.services.abs_service import create_abs_gateway
from gns_app.services.classifier import PageClassifier
from gns_app.services.extractor import FieldExtractor
from gns_app.services.name_service import NameService
from gns_app.services.ocr_service import OcrService
from gns_app.services.official_document import (
    OfficialDocumentClient,
    OfficialDocumentError,
)
from gns_app.services.pdf_service import PdfProcessingError, PdfService
from gns_app.services.qr_service import QrService
from gns_app.services.registry_service import (
    CompositeRegistryClient,
    OsooRegistryClient,
    ReestrKgClient,
    RegistryLookupError,
    RegistryLookupResult,
)
from gns_app.services.scanner_service import (
    DEFAULT_SCANNER_COLOR_MODE,
    DEFAULT_SCANNER_DPI,
    SCANNER_COLOR_INTENTS,
    SCANNER_DPI_CHOICES,
    ScannerCancelled,
    ScannerError,
    ScannerService,
)
from gns_app.services.visual_classifier import PageVisualAnalyzer
from gns_app.services.storage import (
    IMAGE_INPUT_EXTENSIONS,
    SUPPORTED_INPUT_EXTENSIONS,
    ensure_within,
    sanitize_filename,
    save_input_stream,
    save_pdf_stream,
)
from gns_app.services.taxpayer_service import (
    TaxpayerKind,
    classify_taxpayer,
    response_taxpayer_name,
)
from gns_app.services.word_service import WordTemplateError, WordTemplateService
from gns_app.text_cleanup import clean_location, clean_taxpayer_name


class WorkflowValidationError(ValueError):
    pass


class WorkflowService:
    ACTIVE_EMPLOYEE_SETTING = "active_employee_key"
    GNS_OFFICES_SOURCE_SETTING = "gns_offices_source_hash"
    INTAKE_SOURCES = frozenset(
        {"outlook", "manual_upload", "inbox_folder", "legacy"}
    )
    GNS_EMAILS_SOURCE_SETTING = "gns_emails_source_hash"
    EMAIL_RE = re.compile(r"[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}")
    GNS_EMAILS_SOURCE_URL = (
        "https://sti.gov.kg/section/0/electronic_appeals_of_citizens"
    )
    PERIOD_THRESHOLD_SETTING = "period_threshold"
    INBOX_DIR_SETTING = "inbox_dir"
    REGISTRY_PRIORITY_SETTING = "registry_priority"
    PROCESSING_WORKERS_SETTING = "processing_workers"
    SCANNER_DPI_SETTING = "scanner_dpi"
    SCANNER_COLOR_MODE_SETTING = "scanner_color_mode"
    UI_SETTING_DEFAULTS = {
        "allow_multiple_taxpayers": False,
        "show_recipient_salutation": True,
        "require_review_checkbox": False,
    }

    @staticmethod
    def _file_sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                digest.update(chunk)
        return digest.hexdigest()
    OFFICIAL_PARSER_VERSION = 4
    BUSINESS_TIMEZONE = timezone(timedelta(hours=6), "Asia/Bishkek")
    DEFAULT_GNS_OFFICES_PATH = (
        Path(__file__).resolve().parents[1]
        / "data"
        / "ugns_addresses.csv"
    )
    DEFAULT_GNS_EMAILS_PATH = (
        Path(__file__).resolve().parents[1]
        / "data"
        / "ugns_emails.csv"
    )

    def __init__(self, db: Database, settings: Settings):
        self.db = db
        self.settings = settings
        self.pdf = PdfService()
        self.qr = QrService(
            settings.allowed_qr_hosts,
            settings.allowed_qr_paths,
        )
        self.ocr = OcrService(
            self.pdf,
            settings.ocr_fast_data_dir,
            settings.ocr_best_data_dir,
        )
        self.classifier = PageClassifier()
        self.visual = PageVisualAnalyzer()
        self.extractor = FieldExtractor()
        self.names = NameService()
        self.official = OfficialDocumentClient(
            settings.allowed_qr_hosts,
            settings.allowed_qr_paths,
        )
        self.abs = create_abs_gateway(settings)
        # Только оперативная память процесса: в БД и журнал не попадает.
        self._abs_session_credentials: tuple[str, str] | None = None
        # Одна проверка АБС за раз. Дополнительно каждая карточка захватывается
        # атомарным UPDATE по статусу, поэтому повторный клик не запустит её
        # параллельно второй раз.
        self._abs_check_lock = RLock()
        self.registry = CompositeRegistryClient(
            OsooRegistryClient(),
            ReestrKgClient(),
            priority_getter=self.get_registry_priority,
        )
        self.word = WordTemplateService(
            settings.source_templates_dir,
            self.names,
        )
        self.scanner = ScannerService()
        self._gns_offices_cache: list[dict[str, Any]] | None = None
        self._gns_office_terms_cache: list[
            tuple[dict[str, Any], frozenset[str]]
        ] | None = None
        self._gns_office_exact_cache: dict[
            str, tuple[dict[str, Any], ...]
        ] | None = None
        # OCR и рендер независимых страниц выполняются параллельно, а создание
        # и связывание карточек остаётся последовательным. Это исключает две
        # карточки для страниц с одинаковым QR.
        self._case_lock = RLock()

    def abs_session_active(self) -> bool:
        return self._abs_session_credentials is not None

    def _clear_abs_session(self) -> None:
        self._abs_session_credentials = None
        reset_session = getattr(self.abs, "reset_session", None)
        if callable(reset_session):
            reset_session()

    def recover_interrupted_abs_checks(self) -> int:
        """Return checks interrupted by process termination to the ABS queue."""

        now = utc_now()
        with self.db.connect() as connection:
            rows = connection.execute(
                "SELECT id FROM cases WHERE status = ? ORDER BY created_at, id",
                (CaseStatus.ABS_CHECKING,),
            ).fetchall()
            case_ids = [str(row["id"]) for row in rows]
            if not case_ids:
                return 0
            placeholders = ",".join("?" for _ in case_ids)
            connection.execute(
                f"""
                UPDATE cases
                SET status = ?, abs_status = ?, updated_at = ?
                WHERE id IN ({placeholders})
                """,
                (
                    CaseStatus.READY_FOR_ABS,
                    AbsStatus.TECHNICAL_ERROR,
                    now,
                    *case_ids,
                ),
            )
            for case_id in case_ids:
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('case', ?, 'abs_interrupted_requeued',
                              'system', ?, ?)
                    """,
                    (
                        case_id,
                        json.dumps(
                            {"next_status": CaseStatus.READY_FOR_ABS},
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
        self._clear_abs_session()
        return len(case_ids)

    def abs_login_required(self) -> bool:
        if self.abs_session_active():
            return False
        return bool(
            self.db.fetch_one(
                """
                SELECT id
                FROM cases
                WHERE status = ?
                  AND abs_status IN (?, ?, ?)
                LIMIT 1
                """,
                (
                    CaseStatus.READY_FOR_ABS,
                    AbsStatus.AUTH_ERROR,
                    AbsStatus.UNAVAILABLE,
                    AbsStatus.TECHNICAL_ERROR,
                ),
            )
        )

    def abs_is_fake(self) -> bool:
        return bool(getattr(self.abs, "is_fake", False))

    def abs_tls_verification_disabled(self) -> bool:
        return bool(getattr(self.abs, "tls_verification_disabled", False))

    def abs_session_supported(self) -> bool:
        return bool(getattr(self.abs, "supports_session", False))

    def reconcile_recipient_display_names(self) -> int:
        changed = 0
        for case in self.db.fetch_all(
            """
            SELECT id, recipient_full_name, recipient_display_name
            FROM cases
            WHERE recipient_full_name IS NOT NULL
              AND recipient_full_name != ''
            """
        ):
            expected = self.names.recipient_display(case["recipient_full_name"])
            if expected and expected != (case.get("recipient_display_name") or ""):
                self.db.execute(
                    """
                    UPDATE cases
                    SET recipient_display_name = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (expected, utc_now(), case["id"]),
                )
                changed += 1
        return changed

    def lookup_registry(self, inn: str) -> RegistryLookupResult:
        return self.registry.lookup_by_inn(inn)

    def get_period_threshold(self) -> date:
        row = self.db.fetch_one(
            "SELECT value FROM settings WHERE key = ?",
            (self.PERIOD_THRESHOLD_SETTING,),
        )
        if row:
            try:
                return date.fromisoformat(str(row["value"]))
            except ValueError:
                pass
        return self.settings.period_threshold

    def get_inbox_dir(self) -> Path:
        row = self.db.fetch_one(
            "SELECT value FROM settings WHERE key = ?",
            (self.INBOX_DIR_SETTING,),
        )
        value = str(row["value"]).strip() if row else ""
        return Path(value).resolve() if value else self.settings.inbox_dir.resolve()

    def get_registry_priority(self) -> str:
        row = self.db.fetch_one(
            "SELECT value FROM settings WHERE key = ?",
            (self.REGISTRY_PRIORITY_SETTING,),
        )
        value = str(row["value"]).strip() if row else ""
        if value in (
            CompositeRegistryClient.PRIMARY_OSOO,
            CompositeRegistryClient.PRIMARY_REESTR_KG,
        ):
            return value
        return CompositeRegistryClient.PRIMARY_OSOO

    def get_processing_workers(self) -> int:
        row = self.db.fetch_one(
            "SELECT value FROM settings WHERE key = ?",
            (self.PROCESSING_WORKERS_SETTING,),
        )
        try:
            value = int(row["value"]) if row else self.settings.processing_workers
        except (TypeError, ValueError):
            value = self.settings.processing_workers
        return max(1, min(2, value))

    def get_scanner_settings(self) -> dict[str, int | str]:
        rows = self.db.fetch_all(
            "SELECT key, value FROM settings WHERE key IN (?, ?)",
            (self.SCANNER_DPI_SETTING, self.SCANNER_COLOR_MODE_SETTING),
        )
        values = {str(row["key"]): str(row["value"]) for row in rows}
        try:
            dpi = int(values.get(self.SCANNER_DPI_SETTING, DEFAULT_SCANNER_DPI))
        except (TypeError, ValueError):
            dpi = DEFAULT_SCANNER_DPI
        if dpi not in SCANNER_DPI_CHOICES:
            dpi = DEFAULT_SCANNER_DPI
        color_mode = values.get(
            self.SCANNER_COLOR_MODE_SETTING,
            DEFAULT_SCANNER_COLOR_MODE,
        ).strip().casefold()
        if color_mode not in SCANNER_COLOR_INTENTS:
            color_mode = DEFAULT_SCANNER_COLOR_MODE
        return {"dpi": dpi, "color_mode": color_mode}

    def update_operational_settings(
        self,
        *,
        period_threshold: str,
        inbox_dir: str,
        registry_priority: str = CompositeRegistryClient.PRIMARY_OSOO,
        processing_workers: int = 2,
        scanner_dpi: int = DEFAULT_SCANNER_DPI,
        scanner_color_mode: str = DEFAULT_SCANNER_COLOR_MODE,
        actor: str = "Сотрудник",
    ) -> None:
        try:
            threshold = date.fromisoformat(period_threshold)
        except ValueError as exc:
            raise WorkflowValidationError("Укажите корректный порог ОДБ") from exc
        folder = Path(inbox_dir.strip()).resolve()
        try:
            folder.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise WorkflowValidationError(
                "Не удалось открыть или создать папку писем"
            ) from exc
        if registry_priority not in (
            CompositeRegistryClient.PRIMARY_OSOO,
            CompositeRegistryClient.PRIMARY_REESTR_KG,
        ):
            registry_priority = CompositeRegistryClient.PRIMARY_OSOO
        workers = max(1, min(2, int(processing_workers)))
        try:
            scanner_dpi = int(scanner_dpi)
        except (TypeError, ValueError) as exc:
            raise WorkflowValidationError(
                "Укажите корректное разрешение сканирования"
            ) from exc
        if scanner_dpi not in SCANNER_DPI_CHOICES:
            raise WorkflowValidationError(
                "Разрешение сканирования должно быть 150, 200 или 300 DPI"
            )
        scanner_color_mode = str(scanner_color_mode).strip().casefold()
        if scanner_color_mode not in SCANNER_COLOR_INTENTS:
            raise WorkflowValidationError("Выберите режим изображения сканера")
        now = utc_now()
        self.db.executemany(
            """
            INSERT INTO settings(key, value, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(key) DO UPDATE
            SET value = excluded.value, updated_at = excluded.updated_at
            """,
            [
                (self.PERIOD_THRESHOLD_SETTING, threshold.isoformat(), now),
                (self.INBOX_DIR_SETTING, str(folder), now),
                (self.REGISTRY_PRIORITY_SETTING, registry_priority, now),
                (self.PROCESSING_WORKERS_SETTING, str(workers), now),
                (self.SCANNER_DPI_SETTING, str(scanner_dpi), now),
                (self.SCANNER_COLOR_MODE_SETTING, scanner_color_mode, now),
            ],
        )
        self.db.audit(
            "settings",
            "processing",
            "processing_settings_updated",
            {
                "period_threshold": threshold.isoformat(),
                "inbox_dir": str(folder),
                "registry_priority": registry_priority,
                "processing_workers": workers,
                "scanner_dpi": scanner_dpi,
                "scanner_color_mode": scanner_color_mode,
            },
            actor=actor,
        )

    def get_ui_preferences(self) -> dict[str, bool]:
        keys = tuple(self.UI_SETTING_DEFAULTS)
        placeholders = ", ".join("?" for _ in keys)
        rows = self.db.fetch_all(
            f"SELECT key, value FROM settings WHERE key IN ({placeholders})",
            keys,
        )
        result = dict(self.UI_SETTING_DEFAULTS)
        for row in rows:
            result[row["key"]] = str(row["value"]).casefold() in {
                "1", "true", "yes", "on"
            }
        return result

    def update_ui_preferences(
        self,
        *,
        allow_multiple_taxpayers: bool,
        show_recipient_salutation: bool = True,
        require_review_checkbox: bool,
        actor: str = "Сотрудник",
    ) -> dict[str, bool]:
        values = {
            "allow_multiple_taxpayers": allow_multiple_taxpayers,
            "show_recipient_salutation": show_recipient_salutation,
            "require_review_checkbox": require_review_checkbox,
        }
        now = utc_now()
        self.db.executemany(
            """
            INSERT INTO settings(key, value, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(key) DO UPDATE
            SET value = excluded.value, updated_at = excluded.updated_at
            """,
            [(key, "1" if value else "0", now) for key, value in values.items()],
        )
        self.db.audit(
            "settings",
            "review_interface",
            "review_interface_settings_updated",
            values,
            actor=actor,
        )
        return values

    def recipient_suggestions(
        self,
        query: str,
        limit: int = 6,
    ) -> list[dict[str, str]]:
        normalized_query = " ".join(query.casefold().split())
        if len(normalized_query) < 2:
            return []
        rows = self.db.fetch_all(
            """
            SELECT recipient_full_name, recipient_display_name,
                   recipient_position, district_place, updated_at
            FROM cases
            WHERE fields_confirmed = 1
              AND recipient_full_name IS NOT NULL
              AND recipient_full_name != ''
            ORDER BY updated_at DESC
            """
        )
        ranked: list[tuple[float, dict[str, str]]] = []
        seen: set[tuple[str, str]] = set()
        for row in rows:
            full_name = " ".join((row.get("recipient_full_name") or "").split())
            district = clean_location(row.get("district_place") or "")
            key = (full_name.casefold(), district.casefold())
            if not full_name or key in seen:
                continue
            searchable = f"{full_name} {district}".casefold()
            if normalized_query in searchable:
                score = 2.0 + len(normalized_query) / max(len(searchable), 1)
            else:
                score = SequenceMatcher(None, normalized_query, searchable).ratio()
                if score < 0.46:
                    continue
            seen.add(key)
            ranked.append(
                (
                    score,
                    {
                        "full_name": full_name,
                        "display_name": row.get("recipient_display_name") or "",
                        "position": row.get("recipient_position") or "",
                        "district_place": district,
                    },
                )
            )
        ranked.sort(key=lambda item: item[0], reverse=True)
        return [item for _score, item in ranked[:limit]]

    def search_registry(
        self,
        query: str,
        search_mode: str,
    ) -> RegistryLookupResult:
        if search_mode == "name":
            return self.registry.search_by_name(query)
        return self.registry.lookup_by_inn(query)

    def initialize_employee_profiles(self) -> None:
        configured = self.settings.employee_name.strip()
        if not configured:
            return
        name = self._normalize_employee_name(configured)
        name_key = name.casefold()
        now = utc_now()
        self.db.execute(
            """
            INSERT INTO employee_profiles(
                name, name_key, created_at, updated_at
            ) VALUES (?, ?, ?, ?)
            ON CONFLICT(name_key) DO NOTHING
            """,
            (name, name_key, now, now),
        )
        if not self._active_employee_from_database():
            self._store_active_employee(name_key)
            self._apply_employee_to_open_cases(name)

    @staticmethod
    def _normalize_office_text(value: str) -> str:
        normalized = value.casefold().replace("ё", "е")
        normalized = normalized.replace("ысык", "иссык")
        # Город Джалал-Абад переименован в Манас. Заменяем только
        # самостоятельное название города, но не «Джалал-Абадская область».
        normalized = re.sub(
            r"\b(?:джалал|жалал)[-\s]+абад\b",
            "манас",
            normalized,
        )
        normalized = re.sub(r"\b(?:городу|города|город)\b", "г", normalized)
        normalized = re.sub(r"[^\w\s-]", " ", normalized)
        return " ".join(normalized.split())

    def replace_gns_offices(
        self,
        records: list[dict[str, Any]],
        actor: str = "Сотрудник",
    ) -> int:
        prepared: list[tuple[str, str, str, str, str, str, str]] = []
        legacy_locations: dict[str, str] = {}
        seen: set[str] = set()
        now = utc_now()
        for record in records:
            office_name = " ".join(
                str(record.get("office_name") or "").split()
            )
            district_place = clean_location(
                str(record.get("district_place") or "")
            )
            postal_address = " ".join(
                str(record.get("postal_address") or "").split()
            )
            raw_aliases = record.get("aliases") or []
            if isinstance(raw_aliases, str):
                raw_aliases = [raw_aliases]
            aliases = [
                " ".join(str(alias).split())
                for alias in raw_aliases
                if " ".join(str(alias).split())
            ]
            if not office_name or not district_place:
                raise WorkflowValidationError(
                    "У каждого налогового органа нужны название и район/место"
                )
            office_key = self._normalize_office_text(
                f"{office_name} {district_place}"
            )
            if office_key in seen:
                raise WorkflowValidationError(
                    "В справочнике повторяется один налоговый орган"
                )
            seen.add(office_key)
            prepared.append(
                (
                    office_key,
                    office_name,
                    district_place,
                    postal_address,
                    json.dumps(aliases, ensure_ascii=False),
                    now,
                    now,
                )
            )
            for legacy_value in [office_name, *aliases]:
                legacy = clean_location(
                    legacy_value.removeprefix("УГНС ")
                    if legacy_value.startswith("УГНС ")
                    else legacy_value
                )
                if legacy and legacy != district_place:
                    legacy_locations[legacy.casefold()] = district_place
        with self.db.connect() as connection:
            connection.execute("DELETE FROM gns_offices")
            connection.executemany(
                """
                INSERT INTO gns_offices(
                    office_key, office_name, district_place,
                    postal_address, aliases_json, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                prepared,
            )
            case_rows = connection.execute(
                """
                SELECT id, district_place FROM cases
                WHERE district_place IS NOT NULL AND district_place != ''
                """
            ).fetchall()
            for case_row in case_rows:
                original = clean_location(case_row["district_place"])
                replacement = legacy_locations.get(original.casefold())
                if not replacement:
                    continue
                connection.execute(
                    "UPDATE cases SET district_place = ?, updated_at = ? "
                    "WHERE id = ?",
                    (replacement, now, case_row["id"]),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        "case",
                        case_row["id"],
                        "gns_office_location_expanded",
                        actor,
                        json.dumps(
                            {"before": original, "after": replacement},
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
        self._invalidate_gns_office_cache()
        self.db.audit(
            "settings",
            "gns_offices",
            "gns_offices_replaced",
            {"office_count": len(prepared)},
            actor=actor,
        )
        return len(prepared)

    @staticmethod
    def read_gns_offices_csv(path: Path) -> list[dict[str, Any]]:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            rows = list(csv.DictReader(source))
        records: list[dict[str, Any]] = []
        for row in rows:
            office_name = " ".join(
                str(row.get("Наименование УГНС") or "").split()
            )
            postal_address = " ".join(
                str(row.get("Юридический адрес") or "").split()
            )
            if not office_name:
                continue
            district_place = " ".join(
                str(row.get("Район и место") or "").split()
            )
            aliases = [
                " ".join(alias.split())
                for alias in str(row.get("Алиасы") or "").split(";")
                if " ".join(alias.split())
            ]
            if not district_place:
                district_place = (
                    office_name.removeprefix("УГНС ")
                    if office_name.startswith("УГНС ")
                    else office_name
                )
            records.append(
                {
                    "office_name": office_name,
                    "district_place": district_place,
                    "postal_address": postal_address,
                    "aliases": aliases,
                }
            )
        return records

    def initialize_gns_offices(self) -> int:
        if not self.DEFAULT_GNS_OFFICES_PATH.is_file():
            return 0
        source_hash = "district-v5:" + hashlib.sha256(
            self.DEFAULT_GNS_OFFICES_PATH.read_bytes()
        ).hexdigest()
        row = self.db.fetch_one("SELECT COUNT(*) AS count FROM gns_offices")
        signature = self.db.fetch_one(
            "SELECT value FROM settings WHERE key = ?",
            (self.GNS_OFFICES_SOURCE_SETTING,),
        )
        if (
            row
            and int(row["count"])
            and signature
            and signature["value"] == source_hash
        ):
            return 0
        preserved = self.db.fetch_all(
            "SELECT * FROM gns_offices WHERE user_modified = 1"
        )
        count = self.replace_gns_offices(
            self.read_gns_offices_csv(self.DEFAULT_GNS_OFFICES_PATH),
            actor="Система",
        )
        if preserved:
            with self.db.connect() as connection:
                connection.executemany(
                    """
                    INSERT INTO gns_offices(
                        office_key, office_name, district_place,
                        postal_address, email_address, backup_emails_json,
                        aliases_json, active, user_modified, is_custom,
                        created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?)
                    ON CONFLICT(office_key) DO UPDATE SET
                        office_name = excluded.office_name,
                        district_place = excluded.district_place,
                        postal_address = excluded.postal_address,
                        email_address = excluded.email_address,
                        backup_emails_json = excluded.backup_emails_json,
                        aliases_json = excluded.aliases_json,
                        active = excluded.active,
                        user_modified = 1,
                        is_custom = excluded.is_custom,
                        updated_at = excluded.updated_at
                    """,
                    [
                        (
                            office["office_key"],
                            office["office_name"],
                            office["district_place"],
                            office.get("postal_address"),
                            office.get("email_address"),
                            office.get("backup_emails_json") or "[]",
                            office.get("aliases_json") or "[]",
                            int(office.get("active", 1)),
                            int(office.get("is_custom", 0)),
                            office["created_at"],
                            office["updated_at"],
                        )
                        for office in preserved
                    ],
                )
            self._invalidate_gns_office_cache()
        self.db.execute(
            """
            INSERT INTO settings(key, value, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(key) DO UPDATE
            SET value = excluded.value, updated_at = excluded.updated_at
            """,
            (self.GNS_OFFICES_SOURCE_SETTING, source_hash, utc_now()),
        )
        return count

    def initialize_gns_office_emails(self) -> dict[str, Any]:
        path = self.DEFAULT_GNS_EMAILS_PATH
        if not path.is_file():
            return {"updated": 0, "unmatched": []}
        source_hash = "office-alias-v3:" + hashlib.sha256(
            path.read_bytes()
        ).hexdigest()
        signature = self.db.fetch_one(
            "SELECT value FROM settings WHERE key = ?",
            (self.GNS_EMAILS_SOURCE_SETTING,),
        )
        if signature and signature["value"] == source_hash:
            return {"updated": 0, "unmatched": []}
        offices = self.list_gns_offices()
        by_name: dict[str, list[dict[str, Any]]] = {}
        for office in offices:
            if office.get("user_modified"):
                continue
            keys = [office.get("office_name") or "", *(office.get("aliases") or [])]
            for key in keys:
                normalized_key = self._normalize_office_text(key)
                candidates = by_name.setdefault(normalized_key, [])
                if not any(candidate["id"] == office["id"] for candidate in candidates):
                    candidates.append(office)
        updated = 0
        unmatched: list[str] = []
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            for row in csv.DictReader(source):
                official_name = " ".join((row.get("Наименование УГНС") or "").split())
                email = (row.get("Email") or "").strip().casefold()
                if not official_name or not re.fullmatch(
                    r"[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}", email
                ):
                    unmatched.append(official_name or "Пустая строка")
                    continue
                candidates = by_name.get(self._normalize_office_text(official_name), [])
                if len(candidates) != 1:
                    unmatched.append(official_name)
                    continue
                self.db.execute(
                    "UPDATE gns_offices SET email_address = ?, updated_at = ? "
                    "WHERE id = ?",
                    (email, utc_now(), candidates[0]["id"]),
                )
                updated += 1
        self.db.execute(
            """
            INSERT INTO settings(key, value, updated_at) VALUES (?, ?, ?)
            ON CONFLICT(key) DO UPDATE
            SET value = excluded.value, updated_at = excluded.updated_at
            """,
            (self.GNS_EMAILS_SOURCE_SETTING, source_hash, utc_now()),
        )
        self.db.audit(
            "settings",
            "gns_office_emails",
            "gns_office_emails_updated",
            {
                "updated": updated,
                "unmatched": unmatched,
                "source_url": self.GNS_EMAILS_SOURCE_URL,
            },
            actor="Система",
        )
        self._invalidate_gns_office_cache()
        return {"updated": updated, "unmatched": unmatched}

    def _invalidate_gns_office_cache(self) -> None:
        self._gns_offices_cache = None
        self._gns_office_terms_cache = None
        self._gns_office_exact_cache = None

    def list_gns_offices(self) -> list[dict[str, Any]]:
        if self._gns_offices_cache is not None:
            return self._gns_offices_cache
        rows = self.db.fetch_all(
            """
            SELECT * FROM gns_offices
            WHERE active = 1
            ORDER BY district_place, office_name
            """
        )
        for row in rows:
            try:
                row["aliases"] = json.loads(
                    row.get("aliases_json") or "[]"
                )
            except json.JSONDecodeError:
                row["aliases"] = []
            try:
                backup_emails = json.loads(
                    row.get("backup_emails_json") or "[]"
                )
                row["backup_emails"] = [
                    str(email).strip().casefold()
                    for email in backup_emails
                    if self.EMAIL_RE.fullmatch(
                        str(email).strip().casefold()
                    )
                ]
            except (json.JSONDecodeError, TypeError):
                row["backup_emails"] = []
            row["delivery_email"] = (
                str(row.get("email_address") or "").strip().casefold()
                or next(iter(row["backup_emails"]), "")
            )
        self._gns_offices_cache = rows
        return self._gns_offices_cache

    @classmethod
    def _normalize_office_emails(
        cls,
        value: str | list[str] | tuple[str, ...],
    ) -> list[str]:
        raw_values = (
            re.split(r"[,;\n]+", value)
            if isinstance(value, str)
            else list(value)
        )
        result: list[str] = []
        for raw in raw_values:
            email = str(raw or "").strip().casefold()
            if not email:
                continue
            if not cls.EMAIL_RE.fullmatch(email):
                raise WorkflowValidationError(f"Некорректная почта: {email}")
            if email not in result:
                result.append(email)
        return result

    @staticmethod
    def office_delivery_email(office: dict[str, Any] | None) -> str:
        if not office:
            return ""
        primary = str(office.get("email_address") or "").strip().casefold()
        if primary:
            return primary
        backups = office.get("backup_emails") or []
        return str(backups[0]).strip().casefold() if backups else ""

    def save_gns_office(
        self,
        district_place: str,
        email_address: str = "",
        backup_emails: str | list[str] = "",
        aliases: str | list[str] = "",
        *,
        office_id: int | None = None,
        actor: str = "Сотрудник",
    ) -> int:
        district = clean_location(district_place)
        if not district:
            raise WorkflowValidationError("Укажите район или город")
        if len(district) > 240:
            raise WorkflowValidationError("Название района слишком длинное")
        primary_values = self._normalize_office_emails(email_address)
        if len(primary_values) > 1:
            raise WorkflowValidationError("Укажите одну основную почту")
        primary = primary_values[0] if primary_values else ""
        reserves = [
            email
            for email in self._normalize_office_emails(backup_emails)
            if email != primary
        ]
        if len(reserves) > 5:
            raise WorkflowValidationError("Можно указать до пяти резервных почт")
        raw_aliases = (
            re.split(r"[,;\n]+", aliases)
            if isinstance(aliases, str)
            else list(aliases)
        )
        saved_aliases: list[str] = []
        for raw_alias in raw_aliases:
            alias = " ".join(str(raw_alias or "").split())
            if not alias:
                continue
            if len(alias) > 240:
                raise WorkflowValidationError("Алиас слишком длинный")
            if all(
                self._normalize_office_text(existing_alias)
                != self._normalize_office_text(alias)
                for existing_alias in saved_aliases
            ):
                saved_aliases.append(alias)
        if len(saved_aliases) > 20:
            raise WorkflowValidationError("Можно указать до двадцати алиасов")

        current = None
        if office_id is not None:
            current = self.db.fetch_one(
                "SELECT * FROM gns_offices WHERE id = ? AND active = 1",
                (office_id,),
            )
            if not current:
                raise WorkflowValidationError("Район не найден")

        normalized_district = self._normalize_office_text(district)
        district_is_unchanged = bool(
            current
            and self._normalize_office_text(current["district_place"])
            == normalized_district
        )
        supplied_emails = {primary, *reserves} - {""}
        for office in self.list_gns_offices():
            if current and office["id"] == current["id"]:
                continue
            if (
                not district_is_unchanged
                and self._normalize_office_text(office["district_place"])
                == normalized_district
            ):
                raise WorkflowValidationError("Такой район уже есть")
            office_emails = {
                str(office.get("email_address") or "").strip().casefold(),
                *(office.get("backup_emails") or []),
            } - {""}
            if supplied_emails & office_emails:
                raise WorkflowValidationError(
                    "Эта почта уже указана у другого района"
                )
            other_terms = [
                office.get("office_name") or "",
                office.get("district_place") or "",
                *(office.get("aliases") or []),
            ]
            normalized_other_terms = {
                self._normalize_office_text(term)
                for term in other_terms
                if self._normalize_office_text(term)
            }
            if (
                not district_is_unchanged
                and normalized_district in normalized_other_terms
            ):
                raise WorkflowValidationError(
                    "Такое название уже используется другим районом"
                )
            if any(
                self._normalize_office_text(alias) in normalized_other_terms
                for alias in saved_aliases
            ):
                raise WorkflowValidationError(
                    "Этот алиас уже используется другим районом"
                )

        now = utc_now()
        if current:
            old_district = clean_location(current["district_place"])
            if (
                self._normalize_office_text(old_district) != normalized_district
                and all(
                    self._normalize_office_text(alias)
                    != self._normalize_office_text(old_district)
                    for alias in saved_aliases
                )
            ):
                saved_aliases.append(old_district)
            self.db.execute(
                """
                UPDATE gns_offices
                SET district_place = ?, email_address = ?,
                    backup_emails_json = ?, aliases_json = ?, user_modified = 1,
                    updated_at = ?
                WHERE id = ? AND active = 1
                """,
                (
                    district,
                    primary or None,
                    json.dumps(reserves, ensure_ascii=False),
                    json.dumps(saved_aliases, ensure_ascii=False),
                    now,
                    current["id"],
                ),
            )
            saved_id = int(current["id"])
            event_type = "gns_office_updated"
        else:
            office_key = "manual:" + hashlib.sha256(
                normalized_district.encode("utf-8")
            ).hexdigest()[:24]
            office_name = (
                f"УГНС {district}"
                if district.casefold().startswith("по ")
                else f"УГНС по {district}"
            )
            with self.db.connect() as connection:
                cursor = connection.execute(
                    """
                    INSERT INTO gns_offices(
                        office_key, office_name, district_place,
                        postal_address, email_address, backup_emails_json,
                        aliases_json, active, user_modified, is_custom,
                        created_at, updated_at
                    ) VALUES (?, ?, ?, '', ?, ?, ?, 1, 1, 1, ?, ?)
                    """,
                    (
                        office_key,
                        office_name,
                        district,
                        primary or None,
                        json.dumps(reserves, ensure_ascii=False),
                        json.dumps(saved_aliases, ensure_ascii=False),
                        now,
                        now,
                    ),
                )
                saved_id = int(cursor.lastrowid)
            event_type = "gns_office_added"

        self._invalidate_gns_office_cache()
        self.db.audit(
            "gns_office",
            str(saved_id),
            event_type,
            {
                "primary_email_configured": bool(primary),
                "backup_email_count": len(reserves),
                "alias_count": len(saved_aliases),
            },
            actor=actor,
        )
        return saved_id

    def delete_gns_office(
        self,
        office_id: int,
        *,
        actor: str = "Сотрудник",
    ) -> None:
        office = self.db.fetch_one(
            "SELECT id FROM gns_offices WHERE id = ? AND active = 1",
            (office_id,),
        )
        if not office:
            raise WorkflowValidationError("Район не найден")
        self.db.execute(
            """
            UPDATE gns_offices
            SET active = 0, user_modified = 1, updated_at = ?
            WHERE id = ?
            """,
            (utc_now(), office_id),
        )
        self._invalidate_gns_office_cache()
        self.db.audit(
            "gns_office",
            str(office_id),
            "gns_office_deleted",
            actor=actor,
        )

    def _gns_office_match_index(
        self,
    ) -> tuple[
        list[tuple[dict[str, Any], frozenset[str]]],
        dict[str, tuple[dict[str, Any], ...]],
    ]:
        if (
            self._gns_office_terms_cache is not None
            and self._gns_office_exact_cache is not None
        ):
            return self._gns_office_terms_cache, self._gns_office_exact_cache
        terms_by_office: list[
            tuple[dict[str, Any], frozenset[str]]
        ] = []
        exact: dict[str, list[dict[str, Any]]] = {}
        for office in self.list_gns_offices():
            raw_terms = [
                office.get("office_name") or "",
                office.get("district_place") or "",
                office.get("postal_address") or "",
                *(office.get("aliases") or []),
            ]
            all_normalized_terms = frozenset(
                normalized
                for term in raw_terms
                if (normalized := self._normalize_office_text(term))
            )
            # Короткая форма «по г. Ош» допустима только как точное и
            # уникальное значение. В поиске подстроки она слишком общая и
            # могла бы захватить более длинный, другой адресат.
            normalized_terms = frozenset(
                term for term in all_normalized_terms if len(term) >= 8
            )
            terms_by_office.append((office, normalized_terms))
            for term in all_normalized_terms:
                candidates = exact.setdefault(term, [])
                if not any(
                    candidate["id"] == office["id"]
                    for candidate in candidates
                ):
                    candidates.append(office)
        self._gns_office_terms_cache = terms_by_office
        self._gns_office_exact_cache = {
            term: tuple(candidates) for term, candidates in exact.items()
        }
        return self._gns_office_terms_cache, self._gns_office_exact_cache

    def match_gns_office(self, text: str) -> dict[str, Any] | None:
        normalized_text = self._normalize_office_text(text)
        if not normalized_text:
            return None
        terms_by_office, exact = self._gns_office_match_index()
        direct = exact.get(normalized_text, ())
        if len(direct) == 1:
            return direct[0]
        if len(direct) > 1:
            return None
        matches: list[dict[str, Any]] = []
        for office, normalized_terms in terms_by_office:
            # A character substring can turn an official QR phrase into a
            # different office.  Accept only complete normalized phrases from
            # the directory or its approved aliases.
            padded_text = f" {normalized_text} "
            if any(f" {term} " in padded_text for term in normalized_terms):
                matches.append(office)
        return matches[0] if len(matches) == 1 else None

    def suggest_gns_offices(
        self,
        text: str,
        limit: int = 1,
    ) -> list[dict[str, Any]]:
        best = self.best_gns_office(text)
        return [best] if best else []

    def best_gns_office(self, text: str) -> dict[str, Any] | None:
        exact = self.match_gns_office(text)
        if exact:
            return exact
        normalized_text = self._normalize_office_text(text)
        if not normalized_text:
            return None
        generic_tokens = {
            "угнс",
            "цоп",
            "уккн",
            "по",
            "г",
            "району",
            "района",
            "области",
            "город",
            "города",
            "бишкек",
            "управление",
            "государственной",
            "налоговой",
            "службы",
        }
        query_tokens = [
            token
            for token in normalized_text.replace("-", " ").split()
            if len(token) >= 3 and token not in generic_tokens
        ]
        ranked: list[tuple[float, dict[str, Any]]] = []
        for office in self.list_gns_offices():
            normalized_name = self._normalize_office_text(
                office.get("office_name") or ""
            )
            distinctive = {
                token
                for token in normalized_name.replace("-", " ").split()
                if len(token) >= 3 and token not in generic_tokens
            }
            if not distinctive or not query_tokens:
                continue
            score = sum(
                max(
                    SequenceMatcher(None, wanted, seen).ratio()
                    for seen in query_tokens
                )
                for wanted in distinctive
            ) / len(distinctive)
            ranked.append((score, office))
        ranked.sort(
            key=lambda item: (
                -item[0],
                item[1]["office_name"].casefold(),
            )
        )
        if not ranked or ranked[0][0] < 0.84:
            return None
        if len(ranked) > 1 and ranked[0][0] - ranked[1][0] < 0.08:
            return None
        return ranked[0][1]

    def canonical_gns_office(
        self,
        text: str,
    ) -> dict[str, Any] | None:
        """Return one directory office for an exact or safe fuzzy match.

        The local directory is the only source of the canonical district
        value. A fuzzy result is accepted only when it clears both the
        confidence threshold and the gap from the runner-up.
        """
        if not self.list_gns_offices():
            self.initialize_gns_offices()
        return self.best_gns_office(clean_location(text))

    def reconcile_gns_office_districts(self) -> dict[str, int]:
        """Canonicalize saved cases and unsigned generated responses."""
        now = utc_now()
        case_count = 0
        group_count = 0
        for case in self.db.fetch_all(
            """
            SELECT id, district_place
            FROM cases
            WHERE district_place IS NOT NULL AND district_place != ''
            ORDER BY created_at, id
            """
        ):
            before = clean_location(case["district_place"])
            office = self.canonical_gns_office(before)
            if not office:
                continue
            after = office["district_place"]
            if before == after:
                continue
            self.db.execute(
                "UPDATE cases SET district_place = ?, updated_at = ? WHERE id = ?",
                (after, now, case["id"]),
            )
            self.db.audit(
                "case",
                case["id"],
                "gns_office_location_canonicalized",
                {
                    "before": before,
                    "after": after,
                    "office_id": office["id"],
                    "source": "unique_directory_match",
                },
            )
            case_count += 1

        for group in self.db.fetch_all(
            """
            SELECT *
            FROM response_groups
            WHERE status = 'created'
              AND district_place IS NOT NULL
              AND district_place != ''
            ORDER BY created_at, id
            """
        ):
            before = clean_location(group["district_place"])
            office = self.canonical_gns_office(before)
            if not office:
                continue
            after = office["district_place"]
            if before == after:
                continue
            active_scan = self.db.fetch_one(
                """
                SELECT signed_response_scans.id
                FROM signed_response_scans
                JOIN response_letters
                  ON response_letters.id =
                     signed_response_scans.response_letter_id
                WHERE response_letters.response_group_id = ?
                  AND signed_response_scans.status != 'superseded'
                LIMIT 1
                """,
                (group["id"],),
            )
            if active_scan:
                self.db.audit(
                    "response_group",
                    group["id"],
                    "gns_office_canonicalization_needs_review",
                    {
                        "office_id": office["id"],
                        "reason": "signed_response_exists",
                    },
                )
                continue

            taxpayers = self.db.fetch_all(
                """
                SELECT name, inn
                FROM response_group_taxpayers
                WHERE response_group_id = ?
                ORDER BY display_order
                """,
                (group["id"],),
            )
            letters = self.db.fetch_all(
                """
                SELECT outgoing_number
                FROM response_letters
                WHERE response_group_id = ?
                ORDER BY letter_order
                """,
                (group["id"],),
            )
            output = ensure_within(
                Path(group["response_path"]),
                self.settings.responses_dir,
            )
            if not taxpayers or not letters or not output.exists():
                continue
            temporary = output.with_name(
                f".{output.stem}-district-{uuid4().hex}.docx"
            )
            try:
                _, likely_overflow = self.word.render_pages(
                    temporary,
                    {
                        "district_place": after,
                        "recipient_position": group["recipient_position"],
                        "recipient_display_name": group["recipient_display_name"],
                        "employee_name": group["employee_name"],
                    },
                    taxpayers,
                    int(group["taxpayers_per_letter"] or len(taxpayers)),
                    outgoing_numbers=[
                        str(letter["outgoing_number"] or "")
                        for letter in letters
                    ],
                )
                temporary.replace(output)
            except (OSError, WordTemplateError, ValueError):
                temporary.unlink(missing_ok=True)
                continue
            self.db.execute(
                """
                UPDATE response_groups
                SET district_place = ?, response_page_overflow = ?,
                    opened_for_print_at = NULL, word_reopen_required = 1,
                    updated_at = ?
                WHERE id = ?
                """,
                (after, int(likely_overflow), now, group["id"]),
            )
            self.db.audit(
                "response_group",
                group["id"],
                "gns_office_location_canonicalized",
                {
                    "before": before,
                    "after": after,
                    "office_id": office["id"],
                    "filename": output.name,
                },
            )
            group_count += 1
        return {"cases": case_count, "response_groups": group_count}

    def reconcile_gns_office_hints(self) -> int:
        rows = self.db.fetch_all(
            """
            SELECT cases.id AS case_id, cases.district_place,
                   pages.extracted_text
            FROM cases
            JOIN pages ON pages.case_id = cases.id
            WHERE cases.fields_confirmed = 0
              AND pages.page_type = 'letter'
              AND pages.manual_confirmed = 0
              AND pages.extracted_text IS NOT NULL
              AND pages.extracted_text != ''
            ORDER BY cases.id, pages.page_number
            """
        )
        matches: dict[str, dict[str, dict[str, Any]]] = {}
        originals: dict[str, str] = {}
        for row in rows:
            office = self.best_gns_office(row["extracted_text"])
            if not office:
                continue
            matches.setdefault(row["case_id"], {})[office["office_key"]] = office
            originals[row["case_id"]] = row.get("district_place") or ""

        updated = 0
        for case_id, candidates in matches.items():
            if len(candidates) != 1:
                continue
            office = next(iter(candidates.values()))
            replacement = office["district_place"]
            if clean_location(originals.get(case_id)) == replacement:
                continue
            self.db.execute(
                "UPDATE cases SET district_place = ?, updated_at = ? WHERE id = ?",
                (replacement, utc_now(), case_id),
            )
            self.db.audit(
                "case",
                case_id,
                "gns_office_suggested",
                {
                    "office_id": office["id"],
                    "district_place": replacement,
                    "source": "unique_fuzzy_ocr_directory_match",
                },
            )
            updated += 1
        return updated

    def list_employee_profiles(self) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            """
            SELECT id, name, name_key, created_at, updated_at
            FROM employee_profiles
            ORDER BY name COLLATE NOCASE
            """
        )

    def get_active_employee(self) -> str | None:
        active = self._active_employee_from_database()
        if active:
            return active
        configured = self.settings.employee_name.strip()
        return configured or None

    def add_employee(self, employee_name: str) -> str:
        name = self._normalize_employee_name(employee_name)
        name_key = name.casefold()
        existing = self.db.fetch_one(
            "SELECT name FROM employee_profiles WHERE name_key = ?",
            (name_key,),
        )
        added = existing is None
        if added:
            now = utc_now()
            self.db.execute(
                """
                INSERT INTO employee_profiles(
                    name, name_key, created_at, updated_at
                ) VALUES (?, ?, ?, ?)
                """,
                (name, name_key, now, now),
            )
        else:
            name = existing["name"]

        self._store_active_employee(name_key)
        self._apply_employee_to_open_cases(name)
        if added:
            self.db.audit(
                "settings",
                "employee_profiles",
                "employee_profile_added",
                {"employee_name": name},
            )
        self.db.audit(
            "settings",
            "active_employee",
            "active_employee_selected",
            {"employee_name": name},
        )
        return name

    def select_employee(self, employee_key: str) -> str:
        row = self.db.fetch_one(
            """
            SELECT name, name_key
            FROM employee_profiles
            WHERE name_key = ?
            """,
            (employee_key.strip().casefold(),),
        )
        if not row:
            raise WorkflowValidationError(
                "Выбранный исполнитель не найден в локальном справочнике"
            )
        self._store_active_employee(row["name_key"])
        self._apply_employee_to_open_cases(row["name"])
        self.db.audit(
            "settings",
            "active_employee",
            "active_employee_selected",
            {"employee_name": row["name"]},
        )
        return row["name"]

    def _active_employee_from_database(self) -> str | None:
        row = self.db.fetch_one(
            """
            SELECT employee_profiles.name
            FROM settings
            JOIN employee_profiles
              ON employee_profiles.name_key = settings.value
            WHERE settings.key = ?
            """,
            (self.ACTIVE_EMPLOYEE_SETTING,),
        )
        return row["name"] if row else None

    def _store_active_employee(self, name_key: str) -> None:
        self.db.execute(
            """
            INSERT INTO settings(key, value, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(key) DO UPDATE
            SET value = excluded.value, updated_at = excluded.updated_at
            """,
            (self.ACTIVE_EMPLOYEE_SETTING, name_key, utc_now()),
        )

    def _apply_employee_to_open_cases(self, employee_name: str) -> None:
        self.db.execute(
            """
            UPDATE cases
            SET employee_name = ?, updated_at = ?
            WHERE response_path IS NULL
              AND status NOT IN (?, ?)
            """,
            (
                employee_name,
                utc_now(),
                CaseStatus.RESPONSE_CREATED,
                CaseStatus.COMPLETED,
            ),
        )
        self.reconcile_official_qr_pages()

    @staticmethod
    def _normalize_employee_name(employee_name: str) -> str:
        name = " ".join(employee_name.split())
        if len(name) < 2:
            raise WorkflowValidationError(
                "Введите имя или ФИО исполнителя"
            )
        if len(name) > 120:
            raise WorkflowValidationError(
                "Имя исполнителя не должно быть длиннее 120 символов"
            )
        if not any(character.isalpha() for character in name):
            raise WorkflowValidationError(
                "Имя исполнителя должно содержать буквы"
            )
        if any(ord(character) < 32 for character in name):
            raise WorkflowValidationError(
                "Имя исполнителя содержит недопустимые символы"
            )
        return name

    def create_upload(
        self,
        original_filename: str,
        stream: BinaryIO,
        *,
        intake_source: str = "manual_upload",
    ) -> str:
        upload_id = uuid4().hex
        safe_name = sanitize_filename(original_filename)
        suffix = Path(safe_name).suffix.casefold()
        if suffix not in SUPPORTED_INPUT_EXTENSIONS:
            raise WorkflowValidationError(
                "Разрешены PDF, BMP, PNG, JPG и JPEG"
            )
        if intake_source not in self.INTAKE_SOURCES:
            raise WorkflowValidationError("Неизвестный источник документа")

        upload_dir = self.settings.uploads_dir / upload_id
        stored_path = upload_dir / "source.pdf"
        original_path = stored_path
        if suffix == ".pdf":
            digest, _ = save_pdf_stream(
                stream,
                stored_path,
                self.settings.max_upload_bytes,
            )
        else:
            original_path = upload_dir / f"source{suffix}"
            digest, _ = save_input_stream(
                stream,
                original_path,
                self.settings.max_upload_bytes,
                suffix,
            )
        try:
            if suffix in IMAGE_INPUT_EXTENSIONS:
                self.pdf.image_to_pdf(original_path, stored_path)
            page_count = self.pdf.page_count(stored_path)
        except Exception:
            stored_path.unlink(missing_ok=True)
            original_path.unlink(missing_ok=True)
            if upload_dir.exists():
                upload_dir.rmdir()
            raise

        created_at = utc_now()
        self.db.execute(
            """
            INSERT INTO uploads(
                id, original_filename, stored_path, sha256,
                intake_source, page_count, status, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                upload_id,
                safe_name,
                str(stored_path),
                digest,
                intake_source,
                page_count,
                UploadStatus.REGISTERED,
                created_at,
            ),
        )
        page_rows = [
            (
                uuid4().hex,
                upload_id,
                page_number,
                PageType.UNKNOWN,
                PageStatus.REGISTERED,
                created_at,
                created_at,
            )
            for page_number in range(1, page_count + 1)
        ]
        self.db.executemany(
            """
            INSERT INTO pages(
                id, upload_id, page_number, page_type,
                status, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            page_rows,
        )
        self.db.audit(
            "upload",
            upload_id,
            "upload_registered",
            {
                "filename": safe_name,
                "page_count": page_count,
                "sha256": digest,
                "intake_source": intake_source,
            },
        )
        return upload_id

    def _original_input_image(self, upload: dict[str, Any]) -> Path | None:
        suffix = Path(
            str(upload.get("original_filename") or "")
        ).suffix.casefold()
        if suffix not in IMAGE_INPUT_EXTENSIONS:
            return None
        candidate = Path(upload["stored_path"]).parent / f"source{suffix}"
        try:
            candidate = ensure_within(candidate, self.settings.uploads_dir)
        except ValueError:
            return None
        return candidate if candidate.is_file() else None

    def import_inbox(self, max_files: int = 200) -> dict[str, Any]:
        if not self.get_active_employee():
            raise WorkflowValidationError(
                "Сначала выберите или добавьте исполнителя"
            )
        inbox = self.get_inbox_dir()
        inbox.mkdir(parents=True, exist_ok=True)
        candidates = sorted(
            (
                path
                for path in inbox.rglob("*")
                if path.is_file()
                and path.suffix.casefold() in SUPPORTED_INPUT_EXTENSIONS
            ),
            key=lambda path: str(path).casefold(),
        )[:max_files]
        imported: list[str] = []
        skipped: list[str] = []
        errors: list[dict[str, str]] = []

        for source in candidates:
            try:
                safe_source = ensure_within(source, inbox)
                digest = hashlib.sha256()
                total = 0
                with safe_source.open("rb") as stream:
                    while chunk := stream.read(1024 * 1024):
                        total += len(chunk)
                        if total > self.settings.max_upload_bytes:
                            raise WorkflowValidationError(
                                "Документ превышает допустимый размер"
                            )
                        digest.update(chunk)
                existing = self.db.fetch_one(
                    "SELECT id FROM uploads WHERE sha256 = ?",
                    (digest.hexdigest(),),
                )
                if existing:
                    skipped.append(safe_source.name)
                    continue
                with safe_source.open("rb") as stream:
                    imported.append(
                        self.create_upload(
                            safe_source.name,
                            stream,
                            intake_source="inbox_folder",
                        )
                    )
            except Exception as exc:
                errors.append(
                    {
                        "filename": source.name,
                        "message": str(exc)[:300],
                    }
                )

        self.db.audit(
            "settings",
            "inbox",
            "inbox_scanned",
            {
                "folder": str(inbox),
                "found": len(candidates),
                "imported": len(imported),
                "skipped_duplicates": len(skipped),
                "errors": errors,
            },
        )
        return {
            "folder": str(inbox),
            "found": len(candidates),
            "imported": imported,
            "skipped": skipped,
            "errors": errors,
        }

    def reset_processing_data(self) -> dict[str, Any]:
        processing = self.db.fetch_one(
            "SELECT COUNT(*) AS count FROM uploads WHERE status = ?",
            (UploadStatus.PROCESSING,),
        ) or {"count": 0}
        if processing["count"]:
            raise WorkflowValidationError(
                "Сейчас обрабатываются PDF. Дождитесь завершения и повторите сброс."
            )

        runtime = self.settings.runtime_dir.resolve()
        data_directories = (
            self.settings.uploads_dir,
            self.settings.previews_dir,
            self.settings.official_dir,
            self.settings.responses_dir,
        )
        safe_directories: list[Path] = []
        for directory in data_directories:
            resolved = directory.resolve()
            if resolved == runtime or runtime not in resolved.parents:
                raise WorkflowValidationError(
                    "Рабочая папка сброса находится вне каталога приложения"
                )
            safe_directories.append(resolved)

        counts = {
            "uploads": int(
                (self.db.fetch_one("SELECT COUNT(*) AS n FROM uploads") or {})
                .get("n", 0)
            ),
            "cases": int(
                (self.db.fetch_one("SELECT COUNT(*) AS n FROM cases") or {})
                .get("n", 0)
            ),
            "pages": int(
                (self.db.fetch_one("SELECT COUNT(*) AS n FROM pages") or {})
                .get("n", 0)
            ),
        }
        with self.db.connect() as connection:
            connection.execute("DELETE FROM response_groups")
            connection.execute("DELETE FROM uploads")
            connection.execute("DELETE FROM audit_events")

        cleanup_errors: list[str] = []
        for directory in safe_directories:
            directory.mkdir(parents=True, exist_ok=True)
            for child in directory.iterdir():
                try:
                    if child.is_symlink() or child.is_file():
                        child.unlink()
                    else:
                        shutil.rmtree(child)
                except OSError:
                    cleanup_errors.append(str(child))

        self._clear_abs_session()
        self.db.audit(
            "settings",
            "processing_data",
            "processing_data_reset",
            {
                **counts,
                "inbox_preserved": str(self.get_inbox_dir()),
                "cleanup_errors": len(cleanup_errors),
            },
            actor=self.get_active_employee() or "Сотрудник",
        )
        return {**counts, "cleanup_errors": cleanup_errors}

    def process_upload(
        self, upload_id: str, only_registered_pages: bool = False
    ) -> None:
        upload = self.get_upload(upload_id)
        if not upload:
            return
        self.db.execute(
            "UPDATE uploads SET status = ? WHERE id = ?",
            (UploadStatus.PROCESSING, upload_id),
        )
        if only_registered_pages:
            pages = self.db.fetch_all(
                """
                SELECT * FROM pages
                WHERE upload_id = ?
                  AND status IN (?, ?, ?)
                  AND manual_confirmed = 0
                ORDER BY page_number
                """,
                (
                    upload_id,
                    PageStatus.REGISTERED,
                    PageStatus.PREVIEW_READY,
                    PageStatus.PROCESSING,
                ),
            )
        else:
            pages = self.db.fetch_all(
                "SELECT * FROM pages WHERE upload_id = ? ORDER BY page_number",
                (upload_id,),
            )
        def process_page(page: dict[str, Any]) -> None:
            try:
                self._process_page(upload, page)
            except Exception as exc:
                self._record_page_processing_error(page["id"], exc)

        worker_count = min(self.get_processing_workers(), len(pages))
        if worker_count > 1:
            with ThreadPoolExecutor(
                max_workers=worker_count,
                thread_name_prefix="gns-page",
            ) as executor:
                futures = [executor.submit(process_page, page) for page in pages]
                for future in as_completed(futures):
                    future.result()
        else:
            for page in pages:
                process_page(page)
        self.reconcile_official_qr_pages(upload_id)
        self._complete_qr_companions(upload_id)
        self._require_letter_for_scan_only_decisions(upload_id)
        self._remove_orphan_cases(upload_id)
        self._refresh_upload_status(upload_id)

    def _record_page_processing_error(
        self, page_id: str, exc: Exception
    ) -> None:
        message = str(exc)[:800]
        self.db.execute(
            """
            UPDATE pages
            SET status = ?, issue_code = ?, issue_message = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                PageStatus.TECHNICAL_ERROR,
                "page_processing_error",
                message,
                utc_now(),
                page_id,
            ),
        )
        self.db.audit(
            "page",
            page_id,
            "page_processing_error",
            {"message": message},
        )
        record_exception(
            "processing",
            "page",
            exc,
        )

    def interrupted_upload_ids(self) -> list[str]:
        return [
            row["id"]
            for row in self.db.fetch_all(
                """
                SELECT id FROM uploads
                WHERE status IN (?, ?)
                ORDER BY created_at, id
                """,
                (UploadStatus.REGISTERED, UploadStatus.PROCESSING),
            )
        ]

    def _require_letter_for_scan_only_decisions(self, upload_id: str) -> int:
        if self._upload_has_confirmed_letter(upload_id):
            return 0
        decisions = self.db.fetch_all(
            """
            SELECT id, page_number FROM pages
            WHERE upload_id = ?
              AND page_type = ?
              AND status = ?
              AND qr_status != ?
              AND manual_confirmed = 0
            ORDER BY page_number
            """,
            (
                upload_id,
                PageType.DECISION,
                PageStatus.COMPLETED,
                QrStatus.FOUND,
            ),
        )
        if not decisions:
            return 0
        now = utc_now()
        for decision in decisions:
            self.db.execute(
                """
                UPDATE pages
                SET status = ?, issue_code = ?, issue_message = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    PageStatus.NEEDS_REVIEW,
                    "confirmed_letter_missing",
                    (
                        "Страница уверенно похожа на решение, но в PDF не "
                        "найдено ни одного письма. Проверьте тип, чтобы письмо "
                        "не было пропущено."
                    ),
                    now,
                    decision["id"],
                ),
            )
            self.db.audit(
                "page",
                decision["id"],
                "decision_without_letter_review_required",
                {"page_number": decision["page_number"]},
            )
        return len(decisions)

    def _upload_has_confirmed_letter(self, upload_id: str) -> bool:
        return bool(
            self.db.fetch_one(
                """
                SELECT pages.id
                FROM pages
                LEFT JOIN cases ON cases.id = pages.case_id
                WHERE pages.upload_id = ?
                  AND (
                    pages.page_type = ?
                    OR (
                      pages.qr_status = ?
                      AND cases.source_kind = ?
                      AND cases.fields_confirmed = 1
                      AND cases.official_document_path IS NOT NULL
                    )
                  )
                LIMIT 1
                """,
                (
                    upload_id,
                    PageType.LETTER,
                    QrStatus.FOUND,
                    ValueSource.QR_OFFICIAL,
                ),
            )
        )

    def reconcile_confident_scan_decisions(
        self,
        upload_id: str | None = None,
    ) -> int:
        """Safely re-evaluate legacy decision pages after rule updates."""
        params: list[Any] = [
            PageType.DECISION,
            PageStatus.NEEDS_REVIEW,
            PageStatus.COMPLETED,
            PageType.UNKNOWN,
            PageStatus.NEEDS_REVIEW,
            QrStatus.FOUND,
        ]
        upload_filter = ""
        if upload_id:
            upload_filter = " AND upload_id = ?"
            params.append(upload_id)
        pages = self.db.fetch_all(
            f"""
            SELECT * FROM pages
            WHERE (
                    (page_type = ? AND status IN (?, ?))
                    OR (page_type = ? AND status = ?)
                  )
              AND qr_status != ?
              AND manual_confirmed = 0
              {upload_filter}
            ORDER BY upload_id, page_number
            """,
            tuple(params),
        )
        changed = 0
        affected_uploads: set[str] = set()
        letter_cache: dict[str, bool] = {}
        for page in pages:
            visual_evidence = self._visual_evidence_for_page(page)
            type_evidence_text = page.get("type_evidence_text") or ""
            classification = self.classifier.classify(
                "\n".join(
                    filter(
                        None,
                        (
                            page.get("extracted_text") or "",
                            type_evidence_text,
                        ),
                    )
                ),
                float(page.get("quality_score") or 0.0),
                visual_evidence,
            )
            new_type_evidence = False
            if (
                visual_evidence
                and visual_evidence.decision_layout
                and not classification.automatic_terminal
                and not type_evidence_text
            ):
                preview_path = self._preview_path_for_page(page)
                if preview_path:
                    type_evidence_text = self.ocr.recognize_type_markers(
                        preview_path
                    )
                if type_evidence_text:
                    new_type_evidence = True
                    classification = self.classifier.classify(
                        "\n".join(
                            filter(
                                None,
                                (
                                    page.get("extracted_text") or "",
                                    type_evidence_text,
                                ),
                            )
                        ),
                        float(page.get("quality_score") or 0.0),
                        visual_evidence,
                    )
            if (
                page.get("page_type") == PageType.UNKNOWN
                and not (
                    classification.page_type == PageType.DECISION
                    and classification.automatic_terminal
                )
            ):
                if new_type_evidence:
                    self.db.execute(
                        """
                        UPDATE pages
                        SET type_evidence_text = ?,
                            type_evidence_method = ?, updated_at = ?
                        WHERE id = ?
                        """,
                        (
                            type_evidence_text,
                            "local_header_ocr",
                            utc_now(),
                            page["id"],
                        ),
                    )
                continue
            current_upload_id = page["upload_id"]
            if current_upload_id not in letter_cache:
                letter_cache[current_upload_id] = (
                    self._upload_has_confirmed_letter(current_upload_id)
                )

            if classification.automatic_terminal and letter_cache[
                current_upload_id
            ]:
                status = PageStatus.COMPLETED
                issue_code = None
                issue_message = None
                event_type = "confident_decision_reconciled"
            elif not letter_cache[current_upload_id]:
                status = PageStatus.NEEDS_REVIEW
                issue_code = "confirmed_letter_missing"
                issue_message = (
                    "Это решение распознано уверенно, но во всём PDF не "
                    "найдено письмо. Проверьте, чтобы основной документ не "
                    "был пропущен."
                )
                event_type = "decision_without_letter_review_required"
            else:
                status = PageStatus.NEEDS_REVIEW
                issue_code = "decision_type_not_confident"
                issue_message = (
                    "Страница похожа на решение, но для автоматического "
                    "завершения не хватило независимых признаков формы. "
                    "Проверьте тип страницы."
                )
                event_type = "decision_confidence_review_required"

            if (
                page.get("status") == status
                and page.get("issue_code") == issue_code
                and page.get("issue_message") == issue_message
                and not new_type_evidence
            ):
                continue
            self.db.execute(
                """
                UPDATE pages
                SET page_type = ?, status = ?, issue_code = ?, issue_message = ?,
                    type_confidence = ?, type_evidence_text = ?,
                    type_evidence_method = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    classification.page_type,
                    status,
                    issue_code,
                    issue_message,
                    classification.confidence,
                    type_evidence_text or None,
                    "local_header_ocr" if type_evidence_text else None,
                    utc_now(),
                    page["id"],
                ),
            )
            self.db.audit(
                "page",
                page["id"],
                event_type,
                {
                    "page_number": page["page_number"],
                    "automatic_terminal": (
                        classification.automatic_terminal
                    ),
                    "reasons": classification.reasons,
                    "visual_evidence": (
                        {
                            "decision_layout": visual_evidence.decision_layout,
                            "confidence": visual_evidence.confidence,
                            "horizontal_line_groups": (
                                visual_evidence.horizontal_line_groups
                            ),
                            "vertical_line_groups": (
                                visual_evidence.vertical_line_groups
                            ),
                        }
                        if visual_evidence
                        else None
                    ),
                },
            )
            changed += 1
            affected_uploads.add(current_upload_id)

        for affected_upload_id in affected_uploads:
            self._refresh_upload_status(affected_upload_id)
        return changed

    def _visual_evidence_for_page(self, page: dict[str, Any]):
        preview_path = self._preview_path_for_page(page)
        if preview_path:
            return self.visual.analyze(preview_path)
        return None

    def _preview_path_for_page(
        self, page: dict[str, Any]
    ) -> Path | None:
        candidates = [
            page.get("enhanced_preview_path"),
            page.get("preview_path"),
            self.settings.previews_dir
            / page["upload_id"]
            / f"page-{int(page['page_number']):04d}-enhanced.jpg",
            self.settings.previews_dir
            / page["upload_id"]
            / f"page-{int(page['page_number']):04d}.jpg",
        ]
        for candidate in candidates:
            if candidate and Path(candidate).is_file():
                return Path(candidate)
        return None

    def reconcile_official_qr_pages(
        self, upload_id: str | None = None
    ) -> int:
        parameters: list[Any] = [
            PageStatus.NEEDS_REVIEW,
            QrStatus.FOUND,
            ValueSource.QR_OFFICIAL,
        ]
        upload_filter = ""
        if upload_id:
            upload_filter = "AND pages.upload_id = ?"
            parameters.append(upload_id)
        pages = self.db.fetch_all(
            f"""
            SELECT pages.id, pages.upload_id, pages.page_number
            FROM pages
            JOIN cases ON cases.id = pages.case_id
            WHERE pages.status = ?
              AND pages.manual_confirmed = 0
              AND pages.qr_status = ?
              AND cases.source_kind = ?
              AND cases.fields_confirmed = 1
              AND cases.official_document_path IS NOT NULL
              {upload_filter}
            ORDER BY pages.upload_id, pages.page_number
            """,
            parameters,
        )
        if not pages:
            return 0

        now = utc_now()
        for page in pages:
            self.db.execute(
                """
                UPDATE pages
                SET status = ?, issue_code = NULL, issue_message = NULL,
                    updated_at = ?
                WHERE id = ?
                """,
                (PageStatus.COMPLETED, now, page["id"]),
            )
            self.db.audit(
                "page",
                page["id"],
                "official_qr_auto_completed",
                {
                    "page_number": page["page_number"],
                    "reason": "official_fields_already_confirmed",
                },
            )

        for affected_upload_id in {
            page["upload_id"] for page in pages
        }:
            self._refresh_upload_status(affected_upload_id)
        return len(pages)

    def reprocess_incomplete_official_documents(self) -> int:
        cases = self.db.fetch_all(
            """
            SELECT id, official_document_path
            FROM cases
            WHERE source_kind = ?
              AND official_document_path IS NOT NULL
              AND fields_confirmed = 0
              AND official_parse_version < ?
            ORDER BY created_at
            """,
            (ValueSource.QR_OFFICIAL, self.OFFICIAL_PARSER_VERSION),
        )
        completed = 0
        for case in cases:
            try:
                path = ensure_within(
                    Path(case["official_document_path"]),
                    self.settings.official_dir,
                )
                if path.is_file() and self._apply_official_document(
                    case["id"], path
                ):
                    completed += 1
            except (OSError, ValueError):
                self.db.execute(
                    """
                    UPDATE cases
                    SET official_parse_version = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        self.OFFICIAL_PARSER_VERSION,
                        utc_now(),
                        case["id"],
                    ),
                )
        if completed:
            self.reconcile_official_qr_pages()
        return completed

    def _complete_qr_companions(self, upload_id: str) -> None:
        self.db.execute(
            """
            UPDATE pages
            SET status = ?, issue_code = ?, issue_message = ?, updated_at = ?
            WHERE upload_id = ? AND page_type = ? AND qr_status = ?
              AND qr_payload_hash IS NOT NULL
              AND EXISTS (
                  SELECT 1
                  FROM pages AS letter
                  WHERE letter.upload_id = pages.upload_id
                    AND letter.qr_payload_hash = pages.qr_payload_hash
                    AND letter.page_type = ?
              )
            """,
            (
                PageStatus.COMPLETED,
                "qr_group_accounted",
                (
                    "Страница учтена по тому же проверенному QR, "
                    "что и распознанное письмо."
                ),
                utc_now(),
                upload_id,
                PageType.UNKNOWN,
                QrStatus.FOUND,
                PageType.LETTER,
            ),
        )

    def _remove_orphan_cases(self, upload_id: str) -> None:
        row = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM cases
            WHERE upload_id = ?
              AND source_kind = ?
              AND official_document_path IS NULL
              AND fields_confirmed = 0
              AND status IN (?, ?, ?)
              AND NOT EXISTS (
                  SELECT 1 FROM pages WHERE pages.case_id = cases.id
              )
              AND NOT EXISTS (
                  SELECT 1 FROM response_group_cases
                  WHERE response_group_cases.case_id = cases.id
              )
            """,
            (
                upload_id,
                ValueSource.OCR_SCAN,
                CaseStatus.COLLECTING,
                CaseStatus.NEEDS_REVIEW,
                CaseStatus.TECHNICAL_ERROR,
            ),
        )
        count = int(row["count"] if row else 0)
        if not count:
            return
        self.db.execute(
            """
            DELETE FROM cases
            WHERE upload_id = ?
              AND source_kind = ?
              AND official_document_path IS NULL
              AND fields_confirmed = 0
              AND status IN (?, ?, ?)
              AND NOT EXISTS (
                  SELECT 1 FROM pages WHERE pages.case_id = cases.id
              )
              AND NOT EXISTS (
                  SELECT 1 FROM response_group_cases
                  WHERE response_group_cases.case_id = cases.id
              )
            """,
            (
                upload_id,
                ValueSource.OCR_SCAN,
                CaseStatus.COLLECTING,
                CaseStatus.NEEDS_REVIEW,
                CaseStatus.TECHNICAL_ERROR,
            ),
        )
        self.db.audit(
            "upload",
            upload_id,
            "orphan_cases_removed",
            {"case_count": count},
        )

    def request_problem_reprocess(self, upload_id: str) -> int:
        if not self.get_upload(upload_id):
            raise WorkflowValidationError("PDF не найден")
        row = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM pages
            WHERE upload_id = ? AND manual_confirmed = 0
              AND status IN (?, ?)
            """,
            (
                upload_id,
                PageStatus.NEEDS_REVIEW,
                PageStatus.TECHNICAL_ERROR,
            ),
        )
        count = int(row["count"] if row else 0)
        if count == 0:
            raise WorkflowValidationError(
                "Нет неподтверждённых проблемных страниц для повторной обработки"
            )
        self.db.execute(
            """
            UPDATE pages
            SET status = ?, issue_code = NULL, issue_message = NULL,
                updated_at = ?
            WHERE upload_id = ? AND manual_confirmed = 0
              AND status IN (?, ?)
            """,
            (
                PageStatus.REGISTERED,
                utc_now(),
                upload_id,
                PageStatus.NEEDS_REVIEW,
                PageStatus.TECHNICAL_ERROR,
            ),
        )
        self.db.execute(
            "UPDATE uploads SET status = ?, completed_at = NULL WHERE id = ?",
            (UploadStatus.PROCESSING, upload_id),
        )
        self.db.audit(
            "upload",
            upload_id,
            "upload_reprocess_requested",
            {"page_count": count},
        )
        return count

    def request_page_reprocess(self, page_id: str) -> str:
        page = self.get_page(page_id)
        if not page:
            raise WorkflowValidationError("Страница не найдена")
        if page["manual_confirmed"]:
            raise WorkflowValidationError(
                "Подтверждённую сотрудником страницу нельзя перезапускать"
            )
        if page["status"] not in {
            PageStatus.NEEDS_REVIEW,
            PageStatus.TECHNICAL_ERROR,
        }:
            raise WorkflowValidationError(
                "Эта страница не требует повторной обработки"
            )
        self.db.execute(
            """
            UPDATE pages
            SET status = ?, issue_code = NULL, issue_message = NULL,
                updated_at = ?
            WHERE id = ?
            """,
            (PageStatus.REGISTERED, utc_now(), page_id),
        )
        self.db.execute(
            """
            UPDATE uploads
            SET status = ?, completed_at = NULL
            WHERE id = ?
            """,
            (UploadStatus.PROCESSING, page["upload_id"]),
        )
        self.db.audit(
            "page",
            page_id,
            "page_reprocess_requested",
            {"page_number": page["page_number"]},
        )
        return page["upload_id"]

    def request_precise_ocr(self, page_id: str) -> str:
        page = self.get_page(page_id)
        if not page:
            raise WorkflowValidationError("Страница не найдена")
        if page["manual_confirmed"]:
            raise WorkflowValidationError(
                "Подтверждённую сотрудником страницу нельзя распознавать повторно"
            )
        if page["status"] != PageStatus.NEEDS_REVIEW:
            raise WorkflowValidationError(
                "Точный OCR доступен только для страницы ручной проверки"
            )
        if not self.ocr._model_available(self.settings.ocr_best_data_dir):
            raise WorkflowValidationError(
                "Точная локальная OCR-модель best не установлена"
            )
        self.db.execute(
            "UPDATE pages SET status = ?, updated_at = ? WHERE id = ?",
            (PageStatus.PROCESSING, utc_now(), page_id),
        )
        self.db.audit(
            "page",
            page_id,
            "precise_ocr_requested",
            {"page_number": page["page_number"], "model": "best"},
        )
        return page["upload_id"]

    def process_precise_ocr(self, page_id: str) -> None:
        page = self.get_page(page_id)
        if not page:
            return
        upload = self.get_upload(page["upload_id"])
        if not upload:
            return
        output_path = (
            self.settings.previews_dir
            / upload["id"]
            / f".page-{page['page_number']:04d}-{page_id}-best.png"
        )
        remove_output = True
        try:
            original_image = self._original_input_image(upload)
            if original_image is None:
                self.pdf.render_page_high_resolution(
                    Path(upload["stored_path"]),
                    page["page_number"],
                    output_path,
                )
            else:
                output_path = original_image
                remove_output = False
            result = self.ocr.recognize(
                Path(upload["stored_path"]),
                page["page_number"],
                output_path,
                model_name="best",
            )
            if page.get("case_id") and result.text.strip():
                fields = self.extractor.extract_scan_letter(result.text)
                with self._case_lock:
                    self._prefill_scan_case(
                        page["case_id"],
                        fields,
                        office_suggestion=self.best_gns_office(result.text),
                        allow_ocr_suggestions=True,
                        perform_registry_check=False,
                    )
            self.db.execute(
                """
                UPDATE pages
                SET ocr_status = ?, ocr_confidence = ?, extracted_text = ?,
                    status = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    result.status,
                    result.confidence,
                    result.text,
                    PageStatus.NEEDS_REVIEW,
                    utc_now(),
                    page_id,
                ),
            )
            self.db.audit(
                "page",
                page_id,
                "precise_ocr_completed",
                {
                    "model": "best",
                    "status": result.status,
                    "confidence": result.confidence,
                },
            )
        except Exception as exc:
            self.db.execute(
                "UPDATE pages SET status = ?, updated_at = ? WHERE id = ?",
                (PageStatus.NEEDS_REVIEW, utc_now(), page_id),
            )
            self.db.audit(
                "page",
                page_id,
                "precise_ocr_failed",
                {"model": "best", "error": str(exc)[:300]},
            )
        finally:
            if remove_output:
                output_path.unlink(missing_ok=True)

    def _process_page(
        self, upload: dict[str, Any], page: dict[str, Any]
    ) -> None:
        processing_started = monotonic_time.monotonic()
        page_id = page["id"]
        page_number = page["page_number"]
        record_event(
            "processing",
            "page",
            "started",
            details={"page_number": page_number},
        )
        self.db.execute(
            "UPDATE pages SET status = ?, updated_at = ? WHERE id = ?",
            (PageStatus.PROCESSING, utc_now(), page_id),
        )

        preview_dir = self.settings.previews_dir / upload["id"]
        preview_path = preview_dir / f"page-{page_number:04d}.jpg"
        enhanced_path = preview_dir / f"page-{page_number:04d}-enhanced.jpg"
        rendered = self.pdf.render_page(
            Path(upload["stored_path"]),
            page_number,
            preview_path,
            enhanced_path,
        )
        # Превью является самостоятельным результатом этапа. Сохраняем пути
        # до QR и сетевой загрузки, чтобы таймаут официального сервера не
        # оставлял сотрудника без изображения проблемной страницы.
        self.db.execute(
            """
            UPDATE pages
            SET preview_path = ?, enhanced_preview_path = ?,
                quality_score = ?, status = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                str(rendered.preview_path),
                str(rendered.enhanced_preview_path),
                rendered.quality_score,
                PageStatus.PREVIEW_READY,
                utc_now(),
                page_id,
            ),
        )

        source_pdf = Path(upload["stored_path"])
        original_image = self._original_input_image(upload)
        qr_result = self.qr.decode(preview_path)
        if qr_result.status == QrStatus.NOT_FOUND:
            enhanced_qr = self.qr.decode(enhanced_path)
            if enhanced_qr.status != QrStatus.NOT_FOUND:
                enhanced_qr.method = (
                    f"enhanced-{enhanced_qr.method}"
                    if enhanced_qr.method
                    else "enhanced"
                )
                qr_result = enhanced_qr
        if qr_result.status != QrStatus.FOUND and original_image is not None:
            original_qr = self.qr.decode(original_image)
            if original_qr.status != QrStatus.NOT_FOUND:
                original_qr.method = (
                    f"original-image-{original_qr.method}"
                    if original_qr.method
                    else "original-image"
                )
                qr_result = original_qr
        high_resolution_path = preview_dir / (
            f".page-{page_number:04d}-{page_id}-highres.png"
        )
        high_resolution_ready = False
        high_resolution_attempted = False

        def ensure_high_resolution(audit_event: str) -> Path | None:
            nonlocal high_resolution_ready, high_resolution_attempted
            if high_resolution_ready:
                return high_resolution_path
            if high_resolution_attempted:
                return None
            high_resolution_attempted = True
            try:
                self.pdf.render_page_high_resolution(
                    source_pdf,
                    page_number,
                    high_resolution_path,
                )
                high_resolution_ready = True
                return high_resolution_path
            except (PdfProcessingError, OSError) as exc:
                self.db.audit(
                    "page",
                    page_id,
                    audit_event,
                    {"error": str(exc)[:300]},
                )
                return None

        if qr_result.status == QrStatus.NOT_FOUND:
            qr_image = ensure_high_resolution(
                "qr_high_resolution_retry_failed"
            )
            if qr_image:
                high_resolution_qr = self.qr.decode_high_resolution(
                    qr_image
                )
                if high_resolution_qr.status != QrStatus.NOT_FOUND:
                    qr_result = high_resolution_qr

        case_id: str | None = None
        official_complete = False
        official_issue: str | None = None
        if qr_result.status == QrStatus.FOUND:
            # Карточки и официальные документы связываются под блокировкой:
            # две страницы одного пакета могут иметь одинаковый QR.
            with self._case_lock:
                case_id, needs_official = self._ensure_qr_case(
                    upload["id"], qr_result.payload_hash or ""
                )
                if (
                    self.settings.auto_download_official
                    and needs_official
                    and qr_result.payload
                ):
                    try:
                        official_path = self.official.download(
                            qr_result.payload,
                            self.settings.official_dir / case_id,
                        )
                        official_complete = self._apply_official_document(
                            case_id, official_path
                        )
                        if not official_complete:
                            official_issue = (
                                "Официальная версия получена по QR. "
                                "Не удалось надёжно извлечь все "
                                "обязательные поля."
                            )
                    except (
                        OfficialDocumentError,
                        OSError,
                        ValueError,
                    ) as exc:
                        official_issue = str(exc)
                elif self.settings.auto_download_official and not needs_official:
                    existing_case = self.get_case(case_id)
                    if existing_case and existing_case.get(
                        "official_document_path"
                    ):
                        official_complete = bool(
                            existing_case.get("fields_confirmed")
                        )
                        if not official_complete:
                            official_complete = self._apply_official_document(
                                case_id,
                                Path(existing_case["official_document_path"]),
                            )
                        if not official_complete:
                            official_issue = (
                                "Официальная версия уже получена по QR. "
                                "Не удалось надёжно извлечь все "
                                "обязательные поля."
                            )
                refreshed_case = self.get_case(case_id)
                if (
                    refreshed_case
                    and refreshed_case.get("source_kind")
                    == ValueSource.QR_OFFICIAL
                    and refreshed_case.get("fields_confirmed")
                    and refreshed_case.get("official_document_path")
                ):
                    official_complete = True
                    official_issue = None

        if official_complete:
            ocr_result = OcrResult(
                status=OcrStatus.SKIPPED_OFFICIAL,
                text="",
                confidence=0.0,
                language="не требуется: официальный документ по QR",
                issue=None,
            )
            high_resolution_path.unlink(missing_ok=True)
        else:
            ocr_image_path = original_image or ensure_high_resolution(
                "ocr_high_resolution_render_failed"
            ) or rendered.preview_path
            try:
                ocr_result = self.ocr.recognize(
                    source_pdf,
                    page_number,
                    ocr_image_path,
                )
            finally:
                high_resolution_path.unlink(missing_ok=True)

        visual_evidence = self.visual.analyze(
            rendered.enhanced_preview_path
        )
        type_evidence_text = ""
        classification = self.classifier.classify(
            ocr_result.text,
            rendered.quality_score,
            visual_evidence,
        )
        if (
            qr_result.status != QrStatus.FOUND
            and visual_evidence.decision_layout
            and not classification.automatic_terminal
        ):
            type_evidence_text = self.ocr.recognize_type_markers(
                rendered.enhanced_preview_path
            )
            if type_evidence_text:
                classification = self.classifier.classify(
                    f"{ocr_result.text}\n{type_evidence_text}",
                    rendered.quality_score,
                    visual_evidence,
                )

        extracted = None
        ocr_suggestions_allowed = bool(
            ocr_result.text.strip()
            and ocr_result.status
            not in {
                OcrStatus.NOT_STARTED,
                OcrStatus.REQUIRES_ENGINE,
                OcrStatus.ERROR,
            }
        )
        if (
            classification.page_type in {PageType.LETTER, PageType.UNKNOWN}
            and ocr_suggestions_allowed
        ):
            extracted = self.extractor.extract_scan_letter(ocr_result.text)
            office_suggestion = self.best_gns_office(ocr_result.text)
            has_structured_hints = bool(
                extracted.taxpayers
                or extracted.period_start
                or extracted.period_end
                or extracted.recipient_position
                or extracted.recipient_full_name
                or office_suggestion
            )
            if classification.page_type == PageType.LETTER or has_structured_hints:
                with self._case_lock:
                    if case_id is None:
                        case_id = self._create_scan_case(upload["id"], page_id)
                    self._prefill_scan_case(
                        case_id,
                        extracted,
                        ocr_result.critical_fields_agree,
                        office_suggestion,
                        recipient_fields_agree=ocr_result.recipient_fields_agree,
                        taxpayer_fields_agree=ocr_result.taxpayer_fields_agree,
                        period_fields_agree=ocr_result.period_fields_agree,
                        allow_ocr_suggestions=ocr_suggestions_allowed,
                    )

        page_status, issue_code, issue_message = self._page_outcome(
            classification.page_type,
            classification.confidence,
            qr_result.status,
            official_complete,
            official_issue,
            classification.automatic_terminal,
        )
        if (
            classification.page_type == PageType.LETTER
            and not official_complete
            and qr_result.status != QrStatus.FOUND
            and not ocr_result.critical_fields_agree
        ):
            issue_code = "ocr_critical_fields_disagree"
            if ocr_suggestions_allowed and extracted and (
                extracted.taxpayers
                or (extracted.period_start and extracted.period_end)
            ) or (
                extracted
                and (
                    ocr_result.taxpayer_fields_agree
                    and extracted.taxpayers
                    or ocr_result.period_fields_agree
                    and extracted.period_start
                    and extracted.period_end
                )
            ):
                issue_message = (
                    "Найденные OCR значения уже подставлены как подсказки. "
                    "Сверьте их с изображением и исправьте только "
                    "расхождения."
                )
            else:
                issue_message = (
                    "Локальный OCR не подтверждает наименование, ИНН и "
                    "период автоматически. Критические поля нужно "
                    "посимвольно сверить с изображением."
                )
        self.db.execute(
            """
            UPDATE pages
            SET case_id = ?, preview_path = ?, enhanced_preview_path = ?,
                page_type = ?, type_confidence = ?, quality_score = ?,
                qr_status = ?, qr_payload_hash = ?, qr_safe_url = ?,
                qr_method = ?, ocr_status = ?, ocr_confidence = ?,
                extracted_text = ?, type_evidence_text = ?,
                type_evidence_method = ?, status = ?, issue_code = ?,
                issue_message = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                case_id,
                str(rendered.preview_path),
                str(rendered.enhanced_preview_path),
                classification.page_type,
                classification.confidence,
                rendered.quality_score,
                qr_result.status,
                qr_result.payload_hash,
                qr_result.safe_url,
                qr_result.method,
                ocr_result.status,
                ocr_result.confidence,
                ocr_result.text,
                type_evidence_text or None,
                "local_header_ocr" if type_evidence_text else None,
                page_status,
                issue_code,
                issue_message,
                utc_now(),
                page_id,
            ),
        )
        self.db.audit(
            "page",
            page_id,
            "page_processed",
            {
                "page_number": page_number,
                "page_type": classification.page_type,
                "type_confidence": classification.confidence,
                "quality_score": rendered.quality_score,
                "qr_status": qr_result.status,
                "qr_method": qr_result.method,
                "ocr_status": ocr_result.status,
                "ocr_confidence": ocr_result.confidence,
                "automatic_terminal": classification.automatic_terminal,
                "visual_decision_layout": visual_evidence.decision_layout,
                "visual_confidence": visual_evidence.confidence,
                "visual_horizontal_line_groups": (
                    visual_evidence.horizontal_line_groups
                ),
                "visual_vertical_line_groups": (
                    visual_evidence.vertical_line_groups
                ),
                "status": page_status,
            },
        )
        record_event(
            "processing",
            "page",
            "completed",
            details={
                "duration_ms": int(
                    (monotonic_time.monotonic() - processing_started) * 1000
                ),
                "page_number": page_number,
                "page_type": str(classification.page_type),
                "page_status": str(page_status),
                "ocr_status": str(ocr_result.status),
                "ocr_text_length": len(ocr_result.text),
                "ocr_confidence": ocr_result.confidence,
                "structured_party_count": (
                    len(extracted.taxpayers) if extracted else 0
                ),
                "period_hint_present": bool(
                    extracted
                    and extracted.period_start
                    and extracted.period_end
                ),
                "case_created": bool(case_id),
            },
        )

    def _page_outcome(
        self,
        page_type: PageType,
        type_confidence: float,
        qr_status: QrStatus,
        official_complete: bool,
        official_issue: str | None,
        automatic_terminal: bool = False,
    ) -> tuple[PageStatus, str | None, str | None]:
        if official_complete:
            return PageStatus.COMPLETED, None, None
        if official_issue:
            return (
                PageStatus.NEEDS_REVIEW,
                "official_confirmation_required",
                official_issue,
            )
        if page_type == PageType.UNKNOWN:
            return (
                PageStatus.NEEDS_REVIEW,
                "page_type_unknown",
                "Не удалось уверенно определить тип страницы.",
            )
        if qr_status == QrStatus.FOUND:
            return (
                PageStatus.NEEDS_REVIEW,
                "official_pending",
                "QR прочитан, но обязательные данные ещё не подтверждены.",
            )
        if page_type == PageType.LETTER:
            return (
                PageStatus.NEEDS_REVIEW,
                "scan_confirmation_required",
                "Письмо распознано по скану и требует подтверждения полей.",
            )
        if page_type == PageType.DECISION and automatic_terminal:
            return PageStatus.COMPLETED, None, None
        if page_type == PageType.DECISION:
            return (
                PageStatus.NEEDS_REVIEW,
                "decision_type_not_confident",
                (
                    "Страница похожа на решение, но для автоматического "
                    "завершения не хватило независимых признаков формы. "
                    "Проверьте тип страницы."
                ),
            )
        return (
            PageStatus.NEEDS_REVIEW,
            "manual_review_required",
            "Страница требует ручной проверки.",
        )

    def _ensure_qr_case(
        self, upload_id: str, payload_hash: str
    ) -> tuple[str, bool]:
        existing = self.db.fetch_one(
            """
            SELECT * FROM cases
            WHERE upload_id = ? AND qr_payload_hash = ?
            """,
            (upload_id, payload_hash),
        )
        if existing:
            return existing["id"], not bool(
                existing.get("official_document_path")
            )

        case_id = uuid4().hex
        now = utc_now()
        self.db.execute(
            """
            INSERT INTO cases(
                id, upload_id, status, source_kind, qr_payload_hash,
                employee_name, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                case_id,
                upload_id,
                CaseStatus.COLLECTING,
                ValueSource.QR_LINK,
                payload_hash,
                self.get_active_employee(),
                now,
                now,
            ),
        )
        return case_id, True

    def _create_scan_case(self, upload_id: str, page_id: str) -> str:
        case_id = uuid4().hex
        now = utc_now()
        self.db.execute(
            """
            INSERT INTO cases(
                id, upload_id, status, source_kind,
                employee_name, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                case_id,
                upload_id,
                CaseStatus.NEEDS_REVIEW,
                ValueSource.OCR_SCAN,
                self.get_active_employee(),
                now,
                now,
            ),
        )
        self.db.audit(
            "case",
            case_id,
            "scan_case_created",
            {"source_page_id": page_id},
        )
        return case_id

    def _prefill_scan_case(
        self,
        case_id: str,
        fields,
        critical_fields_agree: bool = False,
        office_suggestion: dict[str, Any] | None = None,
        recipient_fields_agree: bool = False,
        taxpayer_fields_agree: bool = False,
        period_fields_agree: bool = False,
        allow_ocr_suggestions: bool = False,
        perform_registry_check: bool = True,
    ) -> None:
        current = self.get_case(case_id)
        if not current:
            return
        if (
            current.get("source_kind") == ValueSource.QR_OFFICIAL
            and current.get("fields_confirmed")
        ):
            return
        if office_suggestion and not current.get("district_place"):
            self.db.execute(
                """
                UPDATE cases
                SET district_place = COALESCE(district_place, ?),
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    office_suggestion["district_place"],
                    utc_now(),
                    case_id,
                ),
            )
            self.db.audit(
                "case",
                case_id,
                "gns_office_suggested",
                {
                    "office_id": office_suggestion["id"],
                    "district_place": office_suggestion["district_place"],
                    "source": "ocr_exact_directory_match",
                },
            )
        if recipient_fields_agree or allow_ocr_suggestions:
            position = (
                fields.recipient_position
                if fields.recipient_position
                and not current.get("recipient_position")
                else None
            )
            full_name = (
                fields.recipient_full_name
                if fields.recipient_full_name
                and not current.get("recipient_full_name")
                else None
            )
            recipient_display = (
                self.names.recipient_display(full_name) if full_name else None
            )
        else:
            position = None
            full_name = None
            recipient_display = None
        if position or full_name:
            self.db.execute(
                """
                UPDATE cases SET
                    recipient_position = COALESCE(recipient_position, ?),
                    recipient_full_name = COALESCE(recipient_full_name, ?),
                    recipient_display_name = COALESCE(recipient_display_name, ?),
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    position,
                    full_name,
                    recipient_display,
                    utc_now(),
                    case_id,
                ),
            )
            self.db.audit(
                "case",
                case_id,
                "ocr_recipient_suggested",
                {
                    "position": position,
                    "full_name": full_name,
                    "source": (
                        "two_local_ocr_passes"
                        if recipient_fields_agree
                        else "structured_ocr_hint"
                    ),
                },
            )
        suggest_period = bool(
            critical_fields_agree
            or period_fields_agree
            or allow_ocr_suggestions
        )
        suggest_taxpayers = bool(
            critical_fields_agree
            or taxpayer_fields_agree
            or allow_ocr_suggestions
        )
        existing_taxpayers_changed = (
            self._collapse_duplicate_unconfirmed_ocr_taxpayers(case_id)
        )
        self.db.execute(
            """
            UPDATE cases
            SET period_start = COALESCE(period_start, ?),
                period_end = COALESCE(period_end, ?),
                status = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                fields.period_start if suggest_period else None,
                fields.period_end if suggest_period else None,
                CaseStatus.NEEDS_REVIEW,
                utc_now(),
                case_id,
            ),
        )
        if (
            suggest_period
            and fields.period_start
            and fields.period_end
            and not current.get("period_start")
            and not current.get("period_end")
        ):
            self.db.audit(
                "case",
                case_id,
                "ocr_period_suggested",
                {
                    "period_start": fields.period_start,
                    "period_end": fields.period_end,
                    "confirmed": False,
                },
            )
        if (
            suggest_taxpayers
            and fields.taxpayers
            and not self.get_taxpayers(case_id)
        ):
            rows = []
            now = utc_now()
            suggestions: dict[str, str] = {}
            conflicting_inns: set[str] = set()
            for taxpayer in fields.taxpayers:
                if not taxpayer.name or not taxpayer.inn:
                    continue
                inn = re.sub(r"\D", "", taxpayer.inn)
                name = clean_taxpayer_name(taxpayer.name)
                if len(inn) != 14 or not name:
                    continue
                previous = suggestions.get(inn)
                if previous and previous.casefold() != name.casefold():
                    conflicting_inns.add(inn)
                else:
                    suggestions.setdefault(inn, name)
            for index, (inn, name) in enumerate(suggestions.items(), 1):
                rows.append(
                    (
                        uuid4().hex,
                        case_id,
                        index,
                        "" if inn in conflicting_inns else name,
                        inn,
                        ValueSource.OCR_SCAN,
                        ValueSource.OCR_SCAN,
                        0,
                        now,
                        now,
                    )
                )
            if rows:
                self.db.executemany(
                    """
                    INSERT INTO taxpayers(
                        id, case_id, display_order, name, inn,
                        name_source, inn_source, manually_confirmed,
                        created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    rows,
                )
                self.db.audit(
                    "case",
                    case_id,
                    "ocr_taxpayers_suggested",
                    {"count": len(rows), "confirmed": False},
                )
                if conflicting_inns:
                    self.db.audit(
                        "case",
                        case_id,
                        "ocr_taxpayer_name_conflict",
                        {
                            "inns": sorted(conflicting_inns),
                            "confirmed": False,
                        },
                    )
                if perform_registry_check:
                    self.check_registry_case(case_id, automatic=True)
        elif existing_taxpayers_changed:
            if perform_registry_check:
                self.check_registry_case(case_id, automatic=True)

    def reconcile_structured_ocr_hints(self) -> int:
        rows = self.db.fetch_all(
            """
            SELECT pages.case_id, pages.extracted_text
            FROM pages
            JOIN cases ON cases.id = pages.case_id
            WHERE pages.status IN ('needs_review', 'technical_error')
              AND pages.page_type = 'letter'
              AND pages.manual_confirmed = 0
              AND cases.fields_confirmed = 0
              AND pages.extracted_text IS NOT NULL
              AND pages.extracted_text != ''
            ORDER BY pages.created_at, pages.page_number
            """
        )
        changed = 0
        for row in rows:
            before = self.get_case(row["case_id"])
            before_taxpayers = self.get_taxpayers(row["case_id"])
            fields = self.extractor.extract_scan_letter(row["extracted_text"])
            needs_prefill = bool(
                before
                and (
                    fields.recipient_position
                    and not before.get("recipient_position")
                    or fields.recipient_full_name
                    and not before.get("recipient_full_name")
                    or fields.period_start
                    and not before.get("period_start")
                    or fields.period_end
                    and not before.get("period_end")
                    or fields.taxpayers
                    and not before_taxpayers
                )
            )
            if not needs_prefill:
                continue
            self._prefill_scan_case(
                row["case_id"],
                fields,
                office_suggestion=self.best_gns_office(row["extracted_text"]),
                allow_ocr_suggestions=True,
                perform_registry_check=False,
            )
            after = self.get_case(row["case_id"])
            after_taxpayers = self.get_taxpayers(row["case_id"])
            if before != after or before_taxpayers != after_taxpayers:
                changed += 1
        return changed

    def _collapse_duplicate_unconfirmed_ocr_taxpayers(
        self, case_id: str
    ) -> bool:
        taxpayers = self.get_taxpayers(case_id)
        groups: dict[str, list[dict[str, Any]]] = {}
        for taxpayer in taxpayers:
            inn = re.sub(r"\D", "", taxpayer.get("inn") or "")
            groups.setdefault(inn, []).append(taxpayer)

        changed = False
        for inn, duplicates in groups.items():
            if len(duplicates) < 2 or len(inn) != 14:
                continue
            if any(
                taxpayer.get("manually_confirmed")
                or taxpayer.get("name_source") != ValueSource.OCR_SCAN
                or taxpayer.get("inn_source") != ValueSource.OCR_SCAN
                for taxpayer in duplicates
            ):
                continue
            keep = duplicates[0]
            names = {
                " ".join((taxpayer.get("name") or "").split()).casefold()
                for taxpayer in duplicates
                if (taxpayer.get("name") or "").strip()
            }
            if len(names) > 1:
                self.db.execute(
                    """
                    UPDATE taxpayers
                    SET name = '', registry_status = NULL,
                        registry_name = NULL, registry_director = NULL,
                        registry_checked_at = NULL, registry_provider = NULL,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (utc_now(), keep["id"]),
                )
                self.db.audit(
                    "case",
                    case_id,
                    "ocr_taxpayer_name_conflict",
                    {"inns": [inn], "confirmed": False},
                )
            for duplicate in duplicates[1:]:
                self.db.execute(
                    "DELETE FROM taxpayers WHERE id = ?",
                    (duplicate["id"],),
                )
            changed = True
        return changed

    def _apply_official_document(
        self, case_id: str, official_path: Path
    ) -> bool:
        text = self._official_text(official_path)
        fields = self.extractor.extract_official_letter(text)
        raw_district = clean_location(fields.district_place or "")
        office = self.canonical_gns_office(raw_district) if raw_district else None
        district_place = office["district_place"] if office else raw_district
        if raw_district and not office:
            fields.issues.append(
                "Район и место не удалось однозначно выбрать из справочника ГНС."
            )
        recipient_display = (
            self.names.recipient_display(fields.recipient_full_name or "")
            if fields.recipient_full_name
            else None
        )
        complete = bool(
            fields.confidence >= 0.95
            and district_place
            and office
            and fields.recipient_position
            and fields.recipient_full_name
            and recipient_display
            and fields.period_start
            and fields.period_end
            and fields.taxpayers
        )
        target_status = (
            CaseStatus.READY_FOR_ABS if complete else CaseStatus.NEEDS_REVIEW
        )
        employee = self.get_active_employee()
        self.db.execute(
            """
            UPDATE cases
            SET official_document_path = ?, source_kind = ?,
                district_place = ?,
                recipient_position = ?, recipient_full_name = ?,
                recipient_display_name = ?, period_start = ?,
                period_end = ?, employee_name = ?,
                fields_confirmed = ?, status = ?,
                official_parse_version = ?,
                match_review_previous_status = CASE
                    WHEN match_review_previous_status IS NOT NULL
                    THEN ? ELSE NULL END,
                updated_at = ?
            WHERE id = ?
            """,
            (
                str(official_path),
                ValueSource.QR_OFFICIAL,
                district_place,
                fields.recipient_position,
                fields.recipient_full_name,
                recipient_display,
                fields.period_start,
                fields.period_end,
                employee,
                1 if complete else 0,
                target_status,
                self.OFFICIAL_PARSER_VERSION,
                target_status,
                utc_now(),
                case_id,
            ),
        )
        self.db.execute("DELETE FROM taxpayers WHERE case_id = ?", (case_id,))
        rows = []
        now = utc_now()
        for index, taxpayer in enumerate(fields.taxpayers, 1):
            if taxpayer.name and taxpayer.inn:
                rows.append(
                    (
                        uuid4().hex,
                        case_id,
                        index,
                        clean_taxpayer_name(taxpayer.name),
                        taxpayer.inn,
                        ValueSource.QR_OFFICIAL,
                        ValueSource.QR_OFFICIAL,
                        1,
                        now,
                        now,
                    )
                )
        if rows:
            self.db.executemany(
                """
                INSERT INTO taxpayers(
                    id, case_id, display_order, name, inn,
                    name_source, inn_source, manually_confirmed,
                    created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                rows,
            )
        self.db.audit(
            "case",
            case_id,
            "official_document_processed",
            {
                "extension": official_path.suffix.lower(),
                "fields_complete": complete,
                "employee_profile_present": bool(employee),
                "issues": fields.issues,
                "gns_office_id": office["id"] if office else None,
                "district_canonicalized": bool(
                    office and raw_district != district_place
                ),
            },
        )
        if office and raw_district != district_place:
            self.db.audit(
                "case",
                case_id,
                "gns_office_location_canonicalized",
                {
                    "before": raw_district,
                    "after": district_place,
                    "office_id": office["id"],
                    "source": "official_document_unique_directory_match",
                },
            )
        self.reconcile_case_match_reviews()
        if complete and not self.has_pending_case_match_review(case_id):
            self.auto_check_abs(case_id)
        return complete

    @staticmethod
    def _official_text(path: Path) -> str:
        if path.suffix.casefold() == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if path.suffix.casefold() == ".docx":
            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
        raise WorkflowValidationError(
            "Неподдерживаемый формат официального документа"
        )

    def confirm_page(
        self,
        page_id: str,
        *,
        page_type: str,
        district_place: str = "",
        recipient_position: str = "",
        recipient_full_name: str = "",
        recipient_display_name: str = "",
        period_start: str = "",
        period_end: str = "",
        period_route: str = "",
        employee_name: str = "",
        taxpayers: list[dict[str, str]] | None = None,
        critical_fields_verified: bool = False,
        actor: str = "Сотрудник",
    ) -> str | None:
        page = self.get_page(page_id)
        if not page:
            raise WorkflowValidationError("Страница не найдена")
        if page["status"] not in {
            PageStatus.NEEDS_REVIEW,
            PageStatus.TECHNICAL_ERROR,
        }:
            raise WorkflowValidationError(
                "Страница уже обработана и не требует подтверждения"
            )
        try:
            selected_type = PageType(page_type)
        except ValueError as exc:
            raise WorkflowValidationError("Неизвестный тип страницы") from exc
        if selected_type == PageType.UNKNOWN:
            raise WorkflowValidationError("Нужно обозначить тип страницы")

        case_id = page.get("case_id")
        if selected_type == PageType.LETTER:
            if not critical_fields_verified:
                raise WorkflowValidationError(
                    "Подтвердите, что ИНН, наименования и период сверены "
                    "с изображением страницы"
                )
            employee_name = (
                employee_name.strip() or self.get_active_employee() or ""
            )
            district_place = clean_location(district_place)
            office = self.canonical_gns_office(district_place)
            if not office:
                raise WorkflowValidationError(
                    "Район и место не найдены однозначно в справочнике ГНС"
                )
            original_district_place = district_place
            district_place = office["district_place"]
            if period_start and period_end:
                period_route = ""
            clean_taxpayers = self._validate_manual_fields(
                district_place,
                recipient_position,
                recipient_full_name,
                period_start,
                period_end,
                period_route,
                employee_name,
                taxpayers or [],
            )
            if not case_id:
                case_id = self._create_scan_case(
                    page["upload_id"], page_id
                )
            display_name = (
                recipient_display_name.strip()
                or self.names.recipient_display(recipient_full_name)
            )
            if not display_name:
                raise WorkflowValidationError(
                    "Не удалось сформировать обращение к адресату"
                )

            self.db.execute(
                """
                UPDATE cases
                SET district_place = ?, recipient_position = ?,
                    recipient_full_name = ?, recipient_display_name = ?,
                    period_start = ?, period_end = ?, period_route = ?,
                    employee_name = ?,
                    source_kind = ?, fields_confirmed = 1,
                    status = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    district_place,
                    recipient_position.strip(),
                    recipient_full_name.strip(),
                    display_name,
                    period_start,
                    period_end,
                    period_route or None,
                    employee_name,
                    ValueSource.MANUAL,
                    CaseStatus.READY_FOR_ABS,
                    utc_now(),
                    case_id,
                ),
            )
            self._replace_taxpayers(case_id, clean_taxpayers)
            if original_district_place != district_place:
                self.db.audit(
                    "case",
                    case_id,
                    "gns_office_location_canonicalized",
                    {
                        "before": original_district_place,
                        "after": district_place,
                        "office_id": office["id"],
                        "source": "manual_unique_directory_match",
                    },
                    actor=actor,
                )
            if page.get("qr_status") != QrStatus.FOUND:
                self.check_registry_case(case_id, automatic=True)
        elif case_id:
            linked_case = self.get_case(case_id)
            if linked_case and not linked_case.get("fields_confirmed") and not (
                linked_case.get("official_document_path")
            ):
                # Черновик, созданный из-за ошибочной классификации письмом,
                # не должен удерживать PDF в состоянии ручной проверки.
                case_id = None

        self.db.execute(
            """
            UPDATE pages
            SET case_id = ?, page_type = ?, type_confidence = 1,
                manual_confirmed = 1, status = ?,
                issue_code = NULL, issue_message = NULL, updated_at = ?
            WHERE id = ?
            """,
            (
                case_id,
                selected_type,
                PageStatus.MANUALLY_CONFIRMED,
                utc_now(),
                page_id,
            ),
        )
        self.db.audit(
            "page",
            page_id,
            "page_manually_confirmed",
            {
                "page_type": selected_type,
                "case_id": case_id,
                "taxpayer_count": len(taxpayers or []),
            },
            actor=actor,
        )
        self._remove_orphan_cases(page["upload_id"])
        self._refresh_upload_status(page["upload_id"])
        if selected_type == PageType.LETTER and case_id:
            self.reconcile_case_match_reviews()
        if (
            selected_type == PageType.LETTER
            and case_id
            and not self.has_pending_case_match_review(case_id)
        ):
            self.auto_check_abs(case_id)
        return case_id

    def mark_page_type_from_queue(
        self,
        page_id: str,
        page_type: str,
        actor: str = "Сотрудник",
    ) -> bool:
        try:
            selected = PageType(page_type)
        except ValueError as exc:
            raise WorkflowValidationError("Неизвестный тип страницы") from exc
        if selected == PageType.UNKNOWN:
            raise WorkflowValidationError("Выберите тип страницы")
        if selected != PageType.LETTER:
            self.confirm_page(page_id, page_type=str(selected), actor=actor)
            return True
        page = self.get_page(page_id)
        if not page or page.get("status") not in {
            PageStatus.NEEDS_REVIEW,
            PageStatus.TECHNICAL_ERROR,
        }:
            raise WorkflowValidationError("Страница уже обработана")
        self.db.execute(
            """
            UPDATE pages
            SET page_type = ?, type_confidence = 1, updated_at = ?
            WHERE id = ?
            """,
            (PageType.LETTER, utc_now(), page_id),
        )
        self.db.audit(
            "page",
            page_id,
            "page_type_selected_from_queue",
            {"page_type": str(PageType.LETTER)},
            actor=actor,
        )
        return False

    def reopen_page_type_review(
        self,
        page_id: str,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        """Return a manually classified non-letter page to the review queue."""
        page = self.get_page(page_id)
        if not page:
            raise WorkflowValidationError("Страница не найдена")
        if (
            page.get("status") != PageStatus.MANUALLY_CONFIRMED
            or not page.get("manual_confirmed")
            or page.get("page_type") == PageType.LETTER
        ):
            raise WorkflowValidationError(
                "Изменить тип можно только у вручную подтверждённого "
                "решения, приложения или прочей страницы"
            )
        previous_type = str(page.get("page_type") or PageType.UNKNOWN)
        self.db.execute(
            """
            UPDATE pages
            SET page_type = ?, type_confidence = 0,
                manual_confirmed = 0, status = ?,
                issue_code = ?, issue_message = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                PageType.UNKNOWN,
                PageStatus.NEEDS_REVIEW,
                "page_type_reopened",
                "Выберите правильный тип страницы.",
                utc_now(),
                page_id,
            ),
        )
        self.db.audit(
            "page",
            page_id,
            "page_type_reopened",
            {"previous_page_type": previous_type},
            actor=actor,
        )
        self._refresh_upload_status(page["upload_id"])
        return self.get_page(page_id) or page

    @staticmethod
    def _registry_name_key(value: str) -> str:
        normalized = value.casefold().replace("ё", "е")
        normalized = re.sub(
            r"\bобщество\s+с\s+ограниченной\s+ответственностью\b",
            "осоо",
            normalized,
        )
        normalized = re.sub(
            r"\bиндивидуальный\s+предприниматель\b",
            "ип",
            normalized,
        )
        return "".join(character for character in normalized if character.isalnum())

    def check_registry_case(
        self,
        case_id: str,
        *,
        automatic: bool = False,
    ) -> dict[str, int]:
        if automatic and not self.settings.auto_registry_check:
            return {"skipped": 1}
        taxpayers = self.get_taxpayers(case_id)
        summary: dict[str, int] = {}
        for taxpayer in taxpayers:
            status = "error"
            official_name = None
            director = None
            provider = None
            kind = classify_taxpayer(taxpayer["name"], taxpayer["inn"])
            if kind == TaxpayerKind.INDIVIDUAL:
                status = "not_applicable"
            elif kind == TaxpayerKind.UNKNOWN:
                status = "classification_uncertain"
            else:
                try:
                    result = self.registry.lookup_by_inn(taxpayer["inn"])
                    official_name = clean_taxpayer_name(result.official_name)
                    director = result.director
                    provider = result.provider
                    if result.status == "found" and official_name:
                        status = (
                            "match"
                            if self._registry_name_key(taxpayer["name"])
                            == self._registry_name_key(official_name)
                            else "mismatch"
                        )
                    else:
                        status = result.status
                except RegistryLookupError:
                    status = "error"
                    provider = None

            checked_at = utc_now()
            self.db.execute(
                """
                UPDATE taxpayers
                SET registry_status = ?, registry_name = ?,
                    registry_director = ?, registry_checked_at = ?,
                    registry_provider = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    status,
                    official_name,
                    director,
                    checked_at,
                    provider,
                    checked_at,
                    taxpayer["id"],
                ),
            )
            summary[status] = summary.get(status, 0) + 1
            self.db.audit(
                "taxpayer",
                taxpayer["id"],
                "registry_checked",
                {
                    "automatic": automatic,
                    "status": status,
                    "provider": provider or "",
                },
            )
        case = self.get_case(case_id)
        needs_confirmation = bool(
            automatic
            and case
            and case.get("source_kind") != ValueSource.QR_OFFICIAL
            and any(
                status not in {"match", "not_applicable"}
                for status in summary
            )
        )
        if needs_confirmation:
            self.db.execute(
                "UPDATE cases SET status = ?, updated_at = ? WHERE id = ?",
                (CaseStatus.NEEDS_REVIEW, utc_now(), case_id),
            )
        return summary

    def accept_registry_variance(
        self,
        case_id: str,
        actor: str = "Сотрудник",
    ) -> None:
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if case.get("source_kind") == ValueSource.QR_OFFICIAL:
            raise WorkflowValidationError(
                "Официальные данные QR не требуют подтверждения по реестру"
            )
        taxpayers = self.get_taxpayers(case_id)
        if not case.get("fields_confirmed") or not taxpayers:
            raise WorkflowValidationError(
                "Сначала подтвердите данные письма и налогоплательщиков"
            )
        if not any(
            taxpayer.get("registry_status")
            and taxpayer.get("registry_status")
            not in {"match", "not_applicable"}
            for taxpayer in taxpayers
        ):
            raise WorkflowValidationError(
                "В карточке нет расхождений ОсОО.KG для подтверждения"
            )
        self.db.execute(
            "UPDATE cases SET status = ?, updated_at = ? WHERE id = ?",
            (CaseStatus.READY_FOR_ABS, utc_now(), case_id),
        )
        self.db.audit(
            "case",
            case_id,
            "registry_variance_accepted",
            {
                "results": [
                    {
                        "inn": taxpayer["inn"],
                        "status": taxpayer.get("registry_status"),
                    }
                    for taxpayer in taxpayers
                ]
            },
            actor=actor,
        )
        self.auto_check_abs(case_id)

    @staticmethod
    def _validate_manual_fields(
        district_place: str,
        recipient_position: str,
        recipient_full_name: str,
        period_start: str,
        period_end: str,
        period_route: str,
        employee_name: str,
        taxpayers: list[dict[str, str]],
    ) -> list[dict[str, str]]:
        fields = {
            "Район и место": district_place,
            "Должность адресата": recipient_position,
            "ФИО адресата": recipient_full_name,
            "Исполнитель банка": employee_name,
        }
        missing = [label for label, value in fields.items() if not value.strip()]
        if missing:
            raise WorkflowValidationError(
                "Не заполнены поля: " + ", ".join(missing)
            )
        if period_start or period_end:
            if not period_start or not period_end:
                raise WorkflowValidationError(
                    "Укажите обе даты периода или выберите маршрут ОДБ"
                )
            try:
                start = date.fromisoformat(period_start)
                end = date.fromisoformat(period_end)
            except ValueError as exc:
                raise WorkflowValidationError("Период содержит неверную дату") from exc
            if start > end:
                raise WorkflowValidationError(
                    "Начало периода не может быть позже окончания"
                )
        elif period_route not in {"odb", "no_odb"}:
            raise WorkflowValidationError(
                "Если даты не читаются, выберите: проверять в ОДБ или нет"
            )
        if not taxpayers:
            raise WorkflowValidationError(
                "Нужно указать хотя бы одного налогоплательщика"
            )

        clean: list[dict[str, str]] = []
        for index, taxpayer in enumerate(taxpayers, 1):
            name = clean_taxpayer_name(taxpayer.get("name", ""))
            inn = re.sub(r"\D", "", taxpayer.get("inn", ""))
            if not name or not inn:
                raise WorkflowValidationError(
                    f"Налогоплательщик {index}: нужны наименование и ИНН"
                )
            clean.append({"name": name, "inn": inn})
        return clean

    def _replace_taxpayers(
        self, case_id: str, taxpayers: list[dict[str, str]]
    ) -> None:
        self.db.execute("DELETE FROM taxpayers WHERE case_id = ?", (case_id,))
        now = utc_now()
        self.db.executemany(
            """
            INSERT INTO taxpayers(
                id, case_id, display_order, name, inn,
                name_source, inn_source, manually_confirmed,
                created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    uuid4().hex,
                    case_id,
                    index,
                    taxpayer["name"],
                    taxpayer["inn"],
                    ValueSource.MANUAL,
                    ValueSource.MANUAL,
                    1,
                    now,
                    now,
                )
                for index, taxpayer in enumerate(taxpayers, 1)
            ],
        )

    def check_abs(
        self,
        case_id: str,
        username: str = "",
        password: str = "",
    ):
        with self._abs_check_lock:
            return self._check_abs_locked(case_id, username, password)

    def _check_abs_locked(
        self,
        case_id: str,
        username: str = "",
        password: str = "",
    ):
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if case["status"] != CaseStatus.READY_FOR_ABS:
            raise WorkflowValidationError(
                "Обращение уже не ожидает проверку АБС"
            )
        if not case["fields_confirmed"]:
            raise WorkflowValidationError(
                "Перед АБС нужно подтвердить все поля письма"
            )
        if self.has_pending_case_match_review(case_id):
            raise WorkflowValidationError(
                "Сначала подтвердите совпадение обращений в ручной проверке"
            )
        taxpayers = self.get_taxpayers(case_id)
        if not taxpayers:
            raise WorkflowValidationError("Нет налогоплательщиков")

        valid_taxpayers = [
            taxpayer
            for taxpayer in taxpayers
            if len(re.sub(r"\D", "", taxpayer.get("inn") or "")) == 14
        ]
        invalid_taxpayers = [
            taxpayer
            for taxpayer in taxpayers
            if len(re.sub(r"\D", "", taxpayer.get("inn") or "")) != 14
        ]
        if not valid_taxpayers:
            now = utc_now()
            threshold = self.get_period_threshold()
            start = (
                date.fromisoformat(case["period_start"])
                if case.get("period_start")
                else None
            )
            next_status = (
                CaseStatus.MANUAL_PERIOD_RULE
                if case.get("period_route") == "odb"
                or (start is not None and start < threshold)
                else CaseStatus.READY_FOR_RESPONSE
            )
            result = AbsCheckResult(
                status=AbsStatus.INVALID_INN,
                taxpayers=[
                    {
                        "inn": taxpayer["inn"],
                        "name": taxpayer["name"],
                        "result": AbsStatus.INVALID_INN,
                    }
                    for taxpayer in invalid_taxpayers
                ],
                message=(
                    "Проверка АБС пропущена: ИНН не содержит 14 цифр. "
                    "Сотрудник должен сверить его с письмом."
                ),
                is_fake=self.abs_is_fake(),
            )
            with self.db.connect() as connection:
                claimed = connection.execute(
                    """
                    UPDATE cases
                    SET status = ?, abs_status = ?, updated_at = ?
                    WHERE id = ? AND status = ?
                    """,
                    (
                        next_status,
                        AbsStatus.INVALID_INN,
                        now,
                        case_id,
                        CaseStatus.READY_FOR_ABS,
                    ),
                )
                if claimed.rowcount != 1:
                    raise WorkflowValidationError(
                        "Обращение уже обрабатывается в АБС"
                    )
                connection.execute(
                    """
                    UPDATE taxpayers
                    SET abs_result = ?, abs_account_result = NULL,
                        abs_active_account_count = NULL,
                        abs_closed_account_count = NULL, odb_result = NULL,
                        updated_at = ?
                    WHERE case_id = ?
                    """,
                    (AbsStatus.INVALID_INN, now, case_id),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('case', ?, 'abs_skipped_invalid_inn',
                              'system', ?, ?)
                    """,
                    (
                        case_id,
                        json.dumps(
                            {
                                "taxpayer_count": len(invalid_taxpayers),
                                "next_status": next_status,
                            },
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
            return result

        username = username.strip()
        session_reused = False
        if not username or not password:
            if self.abs_session_supported() and self._abs_session_credentials:
                username, password = self._abs_session_credentials
                session_reused = True
            else:
                raise WorkflowValidationError(
                    "Введите логин и пароль для проверки АБС"
                )

        check_started_at = utc_now()
        with self.db.connect() as connection:
            claimed = connection.execute(
                """
                UPDATE cases
                SET status = ?, abs_status = ?, updated_at = ?
                WHERE id = ? AND status = ?
                """,
                (
                    CaseStatus.ABS_CHECKING,
                    AbsStatus.CHECKING,
                    check_started_at,
                    case_id,
                    CaseStatus.READY_FOR_ABS,
                ),
            )
            if claimed.rowcount != 1:
                raise WorkflowValidationError(
                    "Обращение уже обрабатывается в АБС"
                )
        abs_started = monotonic_time.monotonic()
        record_event(
            "abs",
            "check",
            "started",
            details={
                "mode": "fake" if self.abs_is_fake() else "tolubay",
                "record_count": len(taxpayers),
                "session_reused": session_reused,
            },
        )
        failed_statuses = {
            AbsStatus.AUTH_ERROR,
            AbsStatus.UNAVAILABLE,
            AbsStatus.TECHNICAL_ERROR,
        }

        def requeue_failed_check(status: AbsStatus) -> None:
            failed_at = utc_now()
            with self.db.connect() as connection:
                current = connection.execute(
                    "SELECT status FROM cases WHERE id = ?",
                    (case_id,),
                ).fetchone()
                if not current or current["status"] != CaseStatus.ABS_CHECKING:
                    return
                connection.execute(
                    """
                    UPDATE cases
                    SET status = ?, abs_status = ?, updated_at = ?
                    WHERE id = ? AND status = ?
                    """,
                    (
                        CaseStatus.READY_FOR_ABS,
                        status,
                        failed_at,
                        case_id,
                        CaseStatus.ABS_CHECKING,
                    ),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('case', ?, ?, 'system', ?, ?)
                    """,
                    (
                        case_id,
                        "fake_abs_checked" if self.abs_is_fake() else "abs_checked",
                        json.dumps(
                            {
                                "status": status,
                                "completed_taxpayers_preserved": True,
                                "next_status": CaseStatus.READY_FOR_ABS,
                            },
                            ensure_ascii=False,
                        ),
                        failed_at,
                    ),
                )

        checkpoint_statuses = {
            AbsStatus.FOUND,
            AbsStatus.NOT_FOUND,
            AbsStatus.MULTIPLE,
        }
        checkpoint_results: dict[str, dict[str, Any]] = {}
        for taxpayer in valid_taxpayers:
            if taxpayer.get("abs_result") in checkpoint_statuses:
                checkpoint_results[str(taxpayer["id"])] = {
                    "inn": taxpayer["inn"],
                    "name": taxpayer["name"],
                    "result": taxpayer["abs_result"],
                    "active_account_count": taxpayer.get(
                        "abs_active_account_count"
                    ),
                    "closed_account_count": taxpayer.get(
                        "abs_closed_account_count"
                    ),
                }

        def persist_checkpoint(
            taxpayer: dict[str, Any], result_item: dict[str, Any]
        ) -> None:
            saved_at = utc_now()
            with self.db.connect() as connection:
                current = connection.execute(
                    "SELECT status FROM cases WHERE id = ?", (case_id,)
                ).fetchone()
                if not current or current["status"] != CaseStatus.ABS_CHECKING:
                    raise WorkflowValidationError(
                        "Состояние обращения изменилось во время проверки АБС"
                    )
                updated = connection.execute(
                    """
                    UPDATE taxpayers
                    SET abs_result = ?, abs_account_result = NULL,
                        abs_active_account_count = ?,
                        abs_closed_account_count = ?, odb_result = NULL,
                        updated_at = ?
                    WHERE id = ? AND case_id = ? AND abs_result IS NULL
                    """,
                    (
                        result_item["result"],
                        result_item.get("active_account_count"),
                        result_item.get("closed_account_count"),
                        saved_at,
                        taxpayer["id"],
                        case_id,
                    ),
                )
                if updated.rowcount != 1:
                    raise WorkflowValidationError(
                        "Результат этого лица уже изменён; обновите страницу."
                    )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('taxpayer', ?, ?, 'system', ?, ?)
                    """,
                    (
                        taxpayer["id"],
                        "fake_abs_taxpayer_checked"
                        if self.abs_is_fake()
                        else "abs_taxpayer_checked",
                        json.dumps(
                            {"status": result_item["result"]},
                            ensure_ascii=False,
                        ),
                        saved_at,
                    ),
                )

        try:
            failed_result: AbsCheckResult | None = None
            for taxpayer in valid_taxpayers:
                taxpayer_id = str(taxpayer["id"])
                if taxpayer_id in checkpoint_results:
                    continue
                one_result = self.abs.check(username, password, [taxpayer])
                if one_result.status in failed_statuses:
                    failed_result = one_result
                    break
                if len(one_result.taxpayers) != 1:
                    failed_result = AbsCheckResult(
                        status=AbsStatus.TECHNICAL_ERROR,
                        taxpayers=[],
                        message=(
                            "АБС вернула неполный или противоречивый ответ. "
                            "Выполните проверку незавершённых лиц заново."
                        ),
                        is_fake=self.abs_is_fake(),
                    )
                    break
                item = one_result.taxpayers[0]
                item_status = item.get("result")
                if (
                    str(item.get("inn") or "") != str(taxpayer["inn"])
                    or item_status not in checkpoint_statuses
                    or (
                        item_status == AbsStatus.FOUND
                        and not one_result.is_fake
                        and any(
                            not isinstance(item.get(key), int)
                            or isinstance(item.get(key), bool)
                            or item.get(key) < 0
                            for key in (
                                "active_account_count",
                                "closed_account_count",
                            )
                        )
                    )
                ):
                    failed_result = AbsCheckResult(
                        status=AbsStatus.TECHNICAL_ERROR,
                        taxpayers=[],
                        message=(
                            "АБС вернула неполный или противоречивый ответ. "
                            "Выполните проверку незавершённых лиц заново."
                        ),
                        is_fake=one_result.is_fake,
                    )
                    break
                persist_checkpoint(taxpayer, item)
                checkpoint_results[taxpayer_id] = item

            if failed_result is not None:
                result = failed_result
            else:
                ordered_results = [
                    checkpoint_results[str(taxpayer["id"])]
                    for taxpayer in valid_taxpayers
                ]
                ordered_results.extend(
                    {
                        "inn": taxpayer["inn"],
                        "name": taxpayer["name"],
                        "result": AbsStatus.INVALID_INN,
                    }
                    for taxpayer in invalid_taxpayers
                )
                result_values = {item["result"] for item in ordered_results}
                result = AbsCheckResult(
                    status=(
                        AbsStatus.MULTIPLE
                        if AbsStatus.MULTIPLE in result_values
                        else AbsStatus.FOUND
                        if AbsStatus.FOUND in result_values
                        else AbsStatus.NOT_FOUND
                    ),
                    taxpayers=ordered_results,
                    message="Проверка АБС завершена.",
                    is_fake=self.abs_is_fake(),
                )
        except Exception as exc:
            self._clear_abs_session()
            requeue_failed_check(AbsStatus.TECHNICAL_ERROR)
            username = ""
            password = ""
            record_exception(
                "abs",
                "check",
                exc,
                details={
                    "duration_ms": int(
                        (monotonic_time.monotonic() - abs_started) * 1000
                    ),
                    "mode": "fake" if self.abs_is_fake() else "tolubay",
                    "record_count": len(taxpayers),
                },
            )
            raise

        validated_results: list[tuple[str, dict[str, Any]]] = []
        if result.status not in failed_statuses:
            expected_taxpayers = valid_taxpayers + invalid_taxpayers
            anomaly = len(result.taxpayers) != len(expected_taxpayers)
            valid_item_statuses = {
                AbsStatus.FOUND,
                AbsStatus.NOT_FOUND,
                AbsStatus.MULTIPLE,
            }
            if not anomaly:
                for expected, actual in zip(
                    expected_taxpayers,
                    result.taxpayers,
                    strict=True,
                ):
                    actual_status = actual.get("result")
                    expected_invalid = expected in invalid_taxpayers
                    if (
                        str(actual.get("inn") or "") != str(expected["inn"])
                        or (
                            expected_invalid
                            and actual_status != AbsStatus.INVALID_INN
                        )
                        or (
                            not expected_invalid
                            and actual_status not in valid_item_statuses
                        )
                    ):
                        anomaly = True
                        break
                    if actual_status == AbsStatus.FOUND and not result.is_fake:
                        for key in (
                            "active_account_count",
                            "closed_account_count",
                        ):
                            count = actual.get(key)
                            if (
                                not isinstance(count, int)
                                or isinstance(count, bool)
                                or count < 0
                            ):
                                anomaly = True
                                break
                    if anomaly:
                        break
                    validated_results.append((str(expected["id"]), actual))

            valid_status_values = [
                item.get("result")
                for item in result.taxpayers[: len(valid_taxpayers)]
            ]
            expected_case_status = (
                AbsStatus.MULTIPLE
                if AbsStatus.MULTIPLE in valid_status_values
                else AbsStatus.FOUND
                if AbsStatus.FOUND in valid_status_values
                else AbsStatus.NOT_FOUND
            )
            if result.status != expected_case_status:
                anomaly = True
            if anomaly:
                self._clear_abs_session()
                result = AbsCheckResult(
                    status=AbsStatus.TECHNICAL_ERROR,
                    taxpayers=[],
                    message=(
                        "АБС вернула неполный или противоречивый ответ. "
                        "Результаты не сохранены; выполните проверку заново."
                    ),
                    is_fake=self.abs_is_fake(),
                )
                validated_results = []

        if result.status in failed_statuses or not self.abs_session_supported():
            self._clear_abs_session()
        else:
            self._abs_session_credentials = (username, password)
        # В БД и журнал учётные данные не передаются.
        username = ""
        password = ""
        record_event(
            "abs",
            "check",
            str(result.status),
            details={
                "duration_ms": int(
                    (monotonic_time.monotonic() - abs_started) * 1000
                ),
                "mode": "fake" if result.is_fake else "tolubay",
                "record_count": len(taxpayers),
                "result_count": len(result.taxpayers),
            },
        )

        threshold = self.get_period_threshold()
        start = (
            date.fromisoformat(case["period_start"])
            if case.get("period_start")
            else None
        )
        if result.status in failed_statuses:
            next_status = CaseStatus.READY_FOR_ABS
        elif any(
            item.get("result") == AbsStatus.FOUND
            for item in result.taxpayers
        ):
            next_status = CaseStatus.NEEDS_REVIEW
        elif case.get("period_route") == "odb" or (
            start is not None and start < threshold
        ):
            next_status = CaseStatus.MANUAL_PERIOD_RULE
        elif result.status == AbsStatus.NOT_FOUND:
            next_status = CaseStatus.READY_FOR_RESPONSE
        else:
            next_status = CaseStatus.NEEDS_REVIEW

        if result.status in failed_statuses:
            requeue_failed_check(result.status)
        else:
            completed_at = utc_now()
            audit_payload = {
                "status": result.status,
                "is_fake": result.is_fake,
                "taxpayers": result.taxpayers,
                "next_status": next_status,
            }
            try:
                with self.db.connect() as connection:
                    current = connection.execute(
                        "SELECT status FROM cases WHERE id = ?",
                        (case_id,),
                    ).fetchone()
                    if (
                        not current
                        or current["status"] != CaseStatus.ABS_CHECKING
                    ):
                        raise WorkflowValidationError(
                            "Состояние обращения изменилось во время проверки АБС"
                        )
                    for taxpayer_id, taxpayer_result in validated_results:
                        updated = connection.execute(
                            """
                            UPDATE taxpayers
                            SET abs_result = ?, abs_account_result = NULL,
                                abs_active_account_count = ?,
                                abs_closed_account_count = ?,
                                odb_result = NULL, updated_at = ?
                            WHERE id = ? AND case_id = ?
                            """,
                            (
                                taxpayer_result["result"],
                                taxpayer_result.get("active_account_count"),
                                taxpayer_result.get("closed_account_count"),
                                completed_at,
                                taxpayer_id,
                                case_id,
                            ),
                        )
                        if updated.rowcount != 1:
                            raise WorkflowValidationError(
                                "Состав лиц изменился во время проверки АБС"
                            )
                    updated_case = connection.execute(
                        """
                        UPDATE cases
                        SET status = ?, abs_status = ?, updated_at = ?
                        WHERE id = ? AND status = ?
                        """,
                        (
                            next_status,
                            result.status,
                            completed_at,
                            case_id,
                            CaseStatus.ABS_CHECKING,
                        ),
                    )
                    if updated_case.rowcount != 1:
                        raise WorkflowValidationError(
                            "Состояние обращения изменилось во время проверки АБС"
                        )
                    connection.execute(
                        """
                        INSERT INTO audit_events(
                            entity_type, entity_id, event_type,
                            actor, payload_json, created_at
                        ) VALUES ('case', ?, ?, 'system', ?, ?)
                        """,
                        (
                            case_id,
                            "fake_abs_checked" if result.is_fake else "abs_checked",
                            json.dumps(audit_payload, ensure_ascii=False),
                            completed_at,
                        ),
                    )
            except Exception:
                self._clear_abs_session()
                requeue_failed_check(AbsStatus.TECHNICAL_ERROR)
                raise
        return result

    def auto_check_abs(self, case_id: str) -> bool:
        case = self.get_case(case_id)
        if not case or case.get("status") != CaseStatus.READY_FOR_ABS:
            return False
        taxpayers = self.get_taxpayers(case_id)
        credentials: tuple[str, str] | None = None
        clear_after = False
        if taxpayers and all(
            len(re.sub(r"\D", "", taxpayer.get("inn") or "")) != 14
            for taxpayer in taxpayers
        ):
            credentials = ("", "")
        elif self.abs_is_fake():
            credentials = ("local-auto", "local-auto")
            clear_after = True
        elif self.abs_session_active():
            # Реальная сеть выполняется только фоновым координатором веб-
            # приложения, чтобы пользовательский POST не ожидал Tolubay.
            return False
        else:
            return False
        try:
            self.check_abs(case_id, *credentials)
        except WorkflowValidationError:
            refreshed = self.get_case(case_id)
            if refreshed and refreshed.get("status") != CaseStatus.READY_FOR_ABS:
                return False
            raise
        finally:
            if clear_after:
                self._clear_abs_session()
        return True

    def confirm_odb(
        self,
        case_id: str,
        results: list[dict[str, str]],
        actor: str = "Сотрудник",
    ) -> str:
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if case["status"] != CaseStatus.MANUAL_PERIOD_RULE:
            raise WorkflowValidationError(
                "Проверка ОДБ для этого обращения не требуется"
            )
        threshold = self.get_period_threshold()
        if case.get("period_route") != "odb" and (
            not case.get("period_start")
            or date.fromisoformat(case["period_start"]) >= threshold
        ):
            raise WorkflowValidationError(
                "Период обращения не относится к старой АБС"
            )

        taxpayers = self.get_taxpayers(case_id)
        expected = {taxpayer["inn"] for taxpayer in taxpayers}
        supplied = {item.get("inn", "") for item in results}
        if expected != supplied:
            raise WorkflowValidationError(
                "Нужно отметить результат ОДБ для каждого налогоплательщика"
            )
        allowed = {OdbStatus.FOUND, OdbStatus.NOT_FOUND}
        normalized: dict[str, OdbStatus] = {}
        for item in results:
            try:
                normalized[item["inn"]] = OdbStatus(item["result"])
            except (KeyError, ValueError) as exc:
                raise WorkflowValidationError(
                    "Неизвестный результат проверки ОДБ"
                ) from exc
        if set(normalized.values()) - allowed:
            raise WorkflowValidationError(
                "Неизвестный результат проверки ОДБ"
            )

        now = utc_now()
        for inn, odb_result in normalized.items():
            self.db.execute(
                """
                UPDATE taxpayers
                SET odb_result = ?, updated_at = ?
                WHERE case_id = ? AND inn = ?
                """,
                (odb_result, now, case_id, inn),
            )

        refreshed = self.get_taxpayers(case_id)
        all_absent = all(
            (
                taxpayer.get("abs_result")
                in {AbsStatus.NOT_FOUND, AbsStatus.INVALID_INN}
                or taxpayer.get("abs_account_result") == "not_found"
            )
            and taxpayer.get("odb_result") == OdbStatus.NOT_FOUND
            for taxpayer in refreshed
        )
        next_status = (
            CaseStatus.READY_FOR_RESPONSE
            if all_absent
            else CaseStatus.MANUAL_PERIOD_RULE
        )
        self.db.execute(
            "UPDATE cases SET status = ?, updated_at = ? WHERE id = ?",
            (next_status, now, case_id),
        )
        self.db.audit(
            "case",
            case_id,
            "odb_checked",
            {
                "results": [
                    {"inn": inn, "result": result}
                    for inn, result in normalized.items()
                ],
                "next_status": next_status,
            },
            actor=actor,
        )
        return next_status

    def confirm_odb_taxpayer(
        self,
        case_id: str,
        inn: str,
        result: str,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if case["status"] != CaseStatus.MANUAL_PERIOD_RULE:
            raise WorkflowValidationError(
                "Проверка ОДБ для этого обращения не требуется"
            )
        threshold = self.get_period_threshold()
        if case.get("period_route") != "odb" and (
            not case.get("period_start")
            or date.fromisoformat(case["period_start"]) >= threshold
        ):
            raise WorkflowValidationError(
                "Период обращения не относится к старой АБС"
            )

        taxpayers = self.get_taxpayers(case_id)
        taxpayer = next(
            (item for item in taxpayers if item["inn"] == inn),
            None,
        )
        if not taxpayer:
            raise WorkflowValidationError(
                "Налогоплательщик не относится к этому обращению"
            )
        try:
            odb_result = OdbStatus(result)
        except ValueError as exc:
            raise WorkflowValidationError(
                "Неизвестный результат проверки ОДБ"
            ) from exc
        if odb_result not in {OdbStatus.FOUND, OdbStatus.NOT_FOUND}:
            raise WorkflowValidationError(
                "Неизвестный результат проверки ОДБ"
            )

        now = utc_now()
        self.db.execute(
            """
            UPDATE taxpayers
            SET odb_result = ?, updated_at = ?
            WHERE case_id = ? AND inn = ?
            """,
            (odb_result, now, case_id, inn),
        )
        refreshed = self.get_taxpayers(case_id)
        pending_count = sum(
            1 for item in refreshed if not item.get("odb_result")
        )
        all_absent = all(
            (
                item.get("abs_result")
                in {AbsStatus.NOT_FOUND, AbsStatus.INVALID_INN}
                or item.get("abs_account_result") == "not_found"
            )
            and item.get("odb_result") == OdbStatus.NOT_FOUND
            for item in refreshed
        )
        next_status = (
            CaseStatus.READY_FOR_RESPONSE
            if not pending_count and all_absent
            else CaseStatus.MANUAL_PERIOD_RULE
        )
        self.db.execute(
            "UPDATE cases SET status = ?, updated_at = ? WHERE id = ?",
            (next_status, now, case_id),
        )
        has_found = any(
            item.get("odb_result") == OdbStatus.FOUND for item in refreshed
        )
        self.db.audit(
            "case",
            case_id,
            "odb_checked",
            {
                "results": [{"inn": inn, "result": odb_result}],
                "remaining": pending_count,
                "next_status": next_status,
            },
            actor=actor,
        )
        return {
            "case_id": case_id,
            "inn": inn,
            "result": odb_result,
            "remaining": pending_count,
            "next_status": next_status,
            "manual_response": has_found,
        }

    def confirm_abs_account_taxpayer(
        self,
        case_id: str,
        inn: str,
        result: str,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        case = self.get_case(case_id)
        taxpayer = next(
            (item for item in self.get_taxpayers(case_id) if item["inn"] == inn),
            None,
        )
        if not case or not taxpayer:
            raise WorkflowValidationError("Налогоплательщик не найден")
        if (
            case.get("status") != CaseStatus.NEEDS_REVIEW
            or case.get("abs_status")
            not in {AbsStatus.FOUND, AbsStatus.MULTIPLE}
        ):
            raise WorkflowValidationError(
                "Сначала завершите проверку обращения в АБС"
            )
        if taxpayer.get("abs_result") != AbsStatus.FOUND:
            raise WorkflowValidationError(
                "Отметка счёта нужна только для найденной анкеты АБС"
            )
        if result not in {"found", "not_found"}:
            raise WorkflowValidationError("Неизвестный результат проверки счёта")
        now = utc_now()
        self.db.execute(
            "UPDATE taxpayers SET abs_account_result = ?, updated_at = ? "
            "WHERE case_id = ? AND inn = ?",
            (result, now, case_id, inn),
        )
        refreshed = self.get_taxpayers(case_id)
        pending = any(
            item.get("abs_result") == AbsStatus.FOUND
            and not item.get("abs_account_result")
            for item in refreshed
        )
        has_account = any(
            item.get("abs_account_result") == "found" for item in refreshed
        )
        all_absent = all(
            item.get("abs_result")
            in {AbsStatus.NOT_FOUND, AbsStatus.INVALID_INN}
            or item.get("abs_account_result") == "not_found"
            for item in refreshed
        )
        threshold = self.get_period_threshold()
        old_period = case.get("period_route") == "odb" or bool(
            case.get("period_start")
            and date.fromisoformat(case["period_start"]) < threshold
        )
        if pending or has_account or not all_absent:
            next_status = CaseStatus.NEEDS_REVIEW
        elif old_period:
            next_status = CaseStatus.MANUAL_PERIOD_RULE
        else:
            next_status = CaseStatus.READY_FOR_RESPONSE
        self.db.execute(
            "UPDATE cases SET status = ?, updated_at = ? WHERE id = ?",
            (next_status, now, case_id),
        )
        self.db.audit(
            "taxpayer",
            taxpayer["id"],
            "abs_account_manually_checked",
            {"result": result, "next_status": next_status},
            actor=actor,
        )
        return {"result": result, "next_status": next_status}

    def check_abs_today(
        self,
        username: str = "",
        password: str = "",
        business_date: date | None = None,
        progress_callback: Callable[[dict[str, Any]], None] | None = None,
    ) -> dict[str, Any]:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        initial = self.db.fetch_one(
            "SELECT COUNT(*) AS count FROM cases WHERE status = ?",
            (CaseStatus.READY_FOR_ABS,),
        ) or {"count": 0}
        total_count = int(initial["count"] or 0)
        if not total_count:
            raise WorkflowValidationError(
                "Нет подтверждённых обращений, ожидающих АБС"
            )

        counts: dict[str, int] = {}
        processed_count = 0
        processed_taxpayer_count = 0
        requires_login = False
        error_message = ""
        while True:
            case = self.db.fetch_one(
                """
                SELECT id FROM cases
                WHERE status = ?
                ORDER BY created_at, id
                LIMIT 1
                """,
                (CaseStatus.READY_FOR_ABS,),
            )
            if not case:
                break
            try:
                result = self.check_abs(case["id"], username, password)
            except WorkflowValidationError:
                refreshed = self.get_case(case["id"])
                if refreshed and refreshed.get("status") != CaseStatus.READY_FOR_ABS:
                    continue
                raise
            processed_count += 1
            processed_taxpayer_count += len(result.taxpayers)
            key = str(result.status)
            counts[key] = counts.get(key, 0) + 1
            if progress_callback:
                progress_callback(
                    {
                        "processed_cases": processed_count,
                        "processed_taxpayers": processed_taxpayer_count,
                        "initial_cases": total_count,
                        "last_status": key,
                    }
                )
            if result.status in {
                AbsStatus.AUTH_ERROR,
                AbsStatus.UNAVAILABLE,
                AbsStatus.TECHNICAL_ERROR,
            }:
                requires_login = True
                error_message = result.message
                break
        # Локальные параметры больше не нужны. Успешный сеанс продолжает
        # работать через защищённое состояние объектов только в памяти.
        username = ""
        password = ""
        self.db.audit(
            "settings",
            f"abs-batch-{day.isoformat()}",
            (
                "fake_abs_batch_checked"
                if self.abs_is_fake()
                else "abs_batch_checked"
            ),
            {
                "business_date": day.isoformat(),
                "case_count": processed_count,
                "taxpayer_count": processed_taxpayer_count,
                "results": counts,
                "is_fake": self.abs_is_fake(),
            },
        )
        remaining = self.db.fetch_one(
            "SELECT COUNT(*) AS count FROM cases WHERE status = ?",
            (CaseStatus.READY_FOR_ABS,),
        ) or {"count": 0}
        return {
            "business_date": day.isoformat(),
            "case_count": processed_count,
            "taxpayer_count": processed_taxpayer_count,
            "initial_case_count": total_count,
            "remaining_case_count": int(remaining["count"] or 0),
            "results": counts,
            "requires_login": requires_login,
            "error_message": error_message,
        }

    def today_overview(
        self,
        business_date: date | None = None,
        *,
        view: str = "all",
    ) -> dict[str, Any]:
        if view not in {"all", "prepare", "created", "manual"}:
            raise WorkflowValidationError("Неизвестный раздел ответов")
        self.reconcile_case_match_reviews()
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        counts = self.db.fetch_one(
            """
            SELECT
                SUM(CASE WHEN cases.status = 'ready_for_abs'
                    THEN 1 ELSE 0 END) AS ready_abs,
                SUM(CASE WHEN cases.status = 'ready_for_abs'
                          AND cases.source_kind = 'qr_official'
                    THEN 1 ELSE 0 END) AS ready_qr,
                SUM(CASE WHEN cases.status = 'ready_for_response'
                    THEN 1 ELSE 0 END) AS ready_response,
                SUM(CASE
                    WHEN cases.status = 'manual_period_rule'
                     AND EXISTS (
                        SELECT 1
                        FROM taxpayers AS odb_taxpayers
                        WHERE odb_taxpayers.case_id = cases.id
                          AND odb_taxpayers.odb_result IS NULL
                     )
                    THEN 1 ELSE 0
                END) AS odb_pending
            FROM cases
            """
        ) or {}
        unresolved = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM pages
            WHERE pages.status IN ('needs_review', 'technical_error')
            """
        ) or {"count": 0}
        unresolved_matches = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM case_match_reviews
            WHERE status = 'pending'
            """
        ) or {"count": 0}
        generated: list[dict[str, Any]] = []
        if view in {"all", "created"}:
            # Созданный Word ещё не означает, что письмо отправлено. Пока
            # Outlook надёжно фиксирует только открытие черновика, старые
            # незавершённые ответы нельзя скрывать фильтром текущего дня.
            generated = self._generated_response_groups(None)
        pending_generated = [
            group for group in generated if not group["is_fully_sent"]
        ]
        sent_generated = [
            group for group in generated if group["is_fully_sent"]
        ]

        odb_cases = (
            self._list_odb_cases() if view == "all" else []
        )

        return {
            "business_date": day.isoformat(),
            "ready_abs": int(counts.get("ready_abs") or 0),
            "ready_qr": int(counts.get("ready_qr") or 0),
            "ready_response": int(counts.get("ready_response") or 0),
            "odb_pending": int(counts.get("odb_pending") or 0),
            "unresolved_pages": int(unresolved["count"]),
            "unresolved_matches": int(unresolved_matches["count"]),
            "odb_cases": odb_cases,
            "not_found_groups": (
                self._daily_response_groups(day, AbsStatus.NOT_FOUND)
                if view in {"all", "prepare"}
                else []
            ),
            "found_groups": (
                self._daily_response_groups(day, AbsStatus.FOUND)
                if view in {"all", "manual"}
                else []
            ),
            "generated_groups": generated,
            "pending_generated_groups": pending_generated,
            "sent_generated_groups": sent_generated,
            "unnumbered_letters": (
                self.list_outgoing_letters(None, only_unnumbered=True)
                if view in {"all", "created"}
                else []
            ),
        }

    def _generated_response_groups(
        self,
        business_date: date | None,
    ) -> list[dict[str, Any]]:
        parameters: tuple[Any, ...] = ()
        where = "WHERE response_groups.status = 'created'"
        if business_date is not None:
            where += " AND response_groups.business_date = ?"
            parameters = (business_date.isoformat(),)
        generated = self.db.fetch_all(
            f"""
            SELECT response_groups.*,
                   (
                       SELECT MIN(CAST(response_letters.outgoing_number AS INTEGER))
                       FROM response_letters
                       WHERE response_letters.response_group_id = response_groups.id
                         AND response_letters.outgoing_number IS NOT NULL
                         AND response_letters.outgoing_number != ''
                   ) AS first_outgoing_number
            FROM response_groups
            {where}
            ORDER BY CASE WHEN first_outgoing_number IS NULL THEN 1 ELSE 0 END,
                     first_outgoing_number, created_at, rowid
            """,
            parameters,
        )
        if not generated:
            return []

        generated_by_id = {group["id"]: group for group in generated}
        group_ids = tuple(generated_by_id)
        placeholders = ",".join("?" for _ in group_ids)
        generated_letters = self.db.fetch_all(
            f"""
            SELECT response_letters.*
            FROM response_letters
            JOIN response_groups
              ON response_groups.id = response_letters.response_group_id
            WHERE response_groups.id IN ({placeholders})
            ORDER BY response_groups.created_at,
                     response_groups.rowid,
                     CASE WHEN response_letters.outgoing_number IS NULL
                               OR response_letters.outgoing_number = ''
                          THEN 1 ELSE 0 END,
                     CAST(response_letters.outgoing_number AS INTEGER),
                     response_letters.letter_order
            """,
            group_ids,
        )
        for group in generated:
            group["letters"] = []
            group["taxpayers"] = []

        generated_taxpayers = self.db.fetch_all(
            f"""
            SELECT response_group_id, display_order, source_case_id, name, inn
            FROM response_group_taxpayers
            WHERE response_group_id IN ({placeholders})
            ORDER BY response_group_id, display_order
            """,
            group_ids,
        )
        for taxpayer in generated_taxpayers:
            group = generated_by_id.get(taxpayer["response_group_id"])
            if group is not None:
                group["taxpayers"].append(taxpayer)

        letter_ids = tuple(letter["id"] for letter in generated_letters)
        scans_by_letter: dict[str, dict[str, Any]] = {}
        sessions_by_letter: dict[str, dict[str, Any]] = {}
        outgoing_by_scan: dict[str, dict[str, Any]] = {}
        if letter_ids:
            letter_placeholders = ",".join("?" for _ in letter_ids)
            active_scans = self.db.fetch_all(
                f"""
                SELECT * FROM signed_response_scans
                WHERE response_letter_id IN ({letter_placeholders})
                  AND status IN ('needs_confirmation', 'confirmed')
                ORDER BY created_at DESC
                """,
                letter_ids,
            )
            scans_by_letter = {
                scan["response_letter_id"]: scan for scan in active_scans
            }
            active_scan_sessions = self.db.fetch_all(
                f"""
                SELECT * FROM signed_scan_sessions
                WHERE response_letter_id IN ({letter_placeholders})
                  AND status IN ('collecting', 'technical_error')
                ORDER BY created_at DESC
                """,
                letter_ids,
            )
            session_ids = tuple(
                session["id"] for session in active_scan_sessions
            )
            pages_by_session: dict[str, list[dict[str, Any]]] = {}
            if session_ids:
                session_placeholders = ",".join("?" for _ in session_ids)
                session_pages = self.db.fetch_all(
                    f"""
                    SELECT * FROM signed_scan_session_pages
                    WHERE session_id IN ({session_placeholders})
                      AND status = 'active'
                    ORDER BY session_id, page_order, created_at, id
                    """,
                    session_ids,
                )
                for page in session_pages:
                    pages_by_session.setdefault(
                        page["session_id"], []
                    ).append(page)
            for session in active_scan_sessions:
                session["pages"] = pages_by_session.get(session["id"], [])
                sessions_by_letter.setdefault(
                    session["response_letter_id"], session
                )
            scan_ids = tuple(scan["id"] for scan in active_scans)
            if scan_ids:
                scan_placeholders = ",".join("?" for _ in scan_ids)
                outgoing_messages = self.db.fetch_all(
                    f"""
                    SELECT * FROM outlook_outgoing_messages
                    WHERE signed_scan_id IN ({scan_placeholders})
                    ORDER BY created_at DESC
                    """,
                    scan_ids,
                )
                outgoing_by_scan = {
                    message["signed_scan_id"]: message
                    for message in outgoing_messages
                }

        for letter in generated_letters:
            letter["signed_scan"] = scans_by_letter.get(letter["id"])
            letter["scan_session"] = sessions_by_letter.get(letter["id"])
            if letter["signed_scan"]:
                letter["outlook_message"] = outgoing_by_scan.get(
                    letter["signed_scan"]["id"]
                )
            else:
                letter["outlook_message"] = None
            group = generated_by_id.get(letter["response_group_id"])
            if group is not None:
                office = self.match_gns_office(group["district_place"])
                letter["gns_email"] = self.office_delivery_email(office)
                group["letters"].append(letter)
        for group in generated:
            group["letter_count"] = len(group["letters"])
            group["numbered_letter_count"] = sum(
                1 for letter in group["letters"] if letter["outgoing_number"]
            )
            group["sent_letter_count"] = sum(
                1
                for letter in group["letters"]
                if letter.get("outlook_message")
                and letter["outlook_message"].get("status") == "sent"
            )
            group["is_fully_sent"] = bool(group["letters"]) and (
                group["sent_letter_count"] == group["letter_count"]
            )
        return generated

    def _list_odb_cases(self) -> list[dict[str, Any]]:
        odb_cases = self.db.fetch_all(
            """
            SELECT cases.id, cases.recipient_display_name,
                   cases.district_place, cases.period_start,
                   COUNT(taxpayers.id) AS taxpayer_count,
                   SUM(CASE WHEN taxpayers.odb_result IS NULL
                       THEN 1 ELSE 0 END) AS pending_count
            FROM cases
            JOIN taxpayers ON taxpayers.case_id = cases.id
            WHERE cases.status = 'manual_period_rule'
            GROUP BY cases.id
            HAVING SUM(CASE WHEN taxpayers.odb_result IS NULL
                THEN 1 ELSE 0 END) > 0
            ORDER BY cases.created_at
            """
        )
        if not odb_cases:
            return []
        taxpayers_by_case: dict[str, list[dict[str, Any]]] = {}
        odb_taxpayers = self.db.fetch_all(
            """
            SELECT taxpayers.case_id, taxpayers.name, taxpayers.inn,
                   taxpayers.odb_result
            FROM taxpayers
            JOIN cases ON cases.id = taxpayers.case_id
            WHERE cases.status = 'manual_period_rule'
              AND taxpayers.odb_result IS NULL
            ORDER BY taxpayers.case_id, taxpayers.display_order
            """
        )
        for taxpayer in odb_taxpayers:
            taxpayers_by_case.setdefault(
                taxpayer["case_id"], []
            ).append(taxpayer)
        for odb_case in odb_cases:
            odb_case["taxpayers"] = taxpayers_by_case.get(
                odb_case["id"], []
            )
        return odb_cases

    @staticmethod
    def _filter_group_taxpayers(
        groups: list[dict[str, Any]],
        predicate: Callable[[dict[str, Any]], bool],
    ) -> list[dict[str, Any]]:
        filtered_groups: list[dict[str, Any]] = []
        for source_group in groups:
            taxpayers = [
                dict(taxpayer)
                for taxpayer in source_group.get("taxpayers", [])
                if predicate(taxpayer)
            ]
            if not taxpayers:
                continue
            group = dict(source_group)
            group["taxpayers"] = taxpayers
            group["taxpayer_count"] = len(taxpayers)
            case_ids = {
                str(taxpayer.get("source_case_id") or "")
                for taxpayer in taxpayers
                if taxpayer.get("source_case_id")
            }
            group["case_ids"] = sorted(case_ids)
            group["case_count"] = len(case_ids)
            group["can_generate"] = False
            filtered_groups.append(group)
        return filtered_groups

    def work_counts(
        self,
        business_date: date | None = None,
    ) -> dict[str, int]:
        # Сначала убираем сравнения, которые уже перестали быть актуальными.
        # Иначе навигация считает старую pending-запись, а открытая очередь
        # сразу после этого очищает её и выглядит пустой при счётчике 1.
        self.reconcile_case_match_reviews()
        row = self.db.fetch_one(
            """
            SELECT
                (SELECT COUNT(*) FROM uploads) AS incoming,
                (SELECT COUNT(*) FROM uploads
                 WHERE status = 'processing') AS processing,
                (SELECT COUNT(*) FROM pages
                 WHERE status IN ('needs_review', 'technical_error'))
                    AS review_pages,
                (SELECT COUNT(*) FROM case_match_reviews
                 WHERE status = 'pending') AS review_matches,
                (SELECT COUNT(*)
                 FROM taxpayers
                 JOIN cases ON cases.id = taxpayers.case_id
                 WHERE cases.status = 'needs_review'
                   AND taxpayers.abs_result = 'found'
                   AND taxpayers.abs_account_result IS NULL)
                    AS review_accounts,
                (SELECT COUNT(*)
                 FROM taxpayers
                 JOIN cases ON cases.id = taxpayers.case_id
                 WHERE cases.status = 'manual_period_rule'
                   AND taxpayers.odb_result IS NULL)
                    AS review_odb,
                (SELECT COUNT(*)
                 FROM cases
                 WHERE cases.status IN ('needs_review', 'technical_error')
                   AND NOT EXISTS (
                       SELECT 1 FROM pages
                       WHERE pages.case_id = cases.id
                         AND pages.status IN (
                             'needs_review', 'technical_error'
                         )
                   )
                   AND NOT EXISTS (
                       SELECT 1 FROM case_match_reviews
                       WHERE status = 'pending'
                         AND (left_case_id = cases.id
                              OR right_case_id = cases.id)
                   )
                   AND NOT EXISTS (
                       SELECT 1 FROM taxpayers
                       WHERE taxpayers.case_id = cases.id
                         AND taxpayers.abs_result = 'found'
                         AND taxpayers.abs_account_result IS NULL
                   )
                   AND NOT EXISTS (
                       SELECT 1 FROM taxpayers
                       WHERE taxpayers.case_id = cases.id
                         AND (taxpayers.abs_account_result = 'found'
                              OR taxpayers.odb_result = 'found')
                   )) AS review_cases,
                (SELECT COUNT(*) FROM cases
                 WHERE status = 'ready_for_abs') AS ready_abs,
                (SELECT COUNT(*) FROM cases
                 WHERE status = 'ready_for_response') AS ready_responses,
                (SELECT COUNT(*) FROM response_groups
                 WHERE status = 'created')
                    AS created_responses,
                (SELECT COUNT(DISTINCT cases.id)
                 FROM cases
                 JOIN taxpayers ON taxpayers.case_id = cases.id
                 WHERE taxpayers.abs_account_result = 'found'
                    OR taxpayers.odb_result = 'found') AS manual_responses
            """
        ) or {}
        review_count = sum(
            int(row.get(key) or 0)
            for key in (
                "review_pages",
                "review_matches",
                "review_accounts",
                "review_odb",
                "review_cases",
            )
        )
        return {
            "incoming": int(row.get("incoming") or 0),
            "processing": int(row.get("processing") or 0),
            "review": review_count,
            "ready_abs": int(row.get("ready_abs") or 0),
            "ready_responses": int(row.get("ready_responses") or 0),
            "created_responses": int(row.get("created_responses") or 0),
            "manual_responses": int(row.get("manual_responses") or 0),
        }

    def manual_review_overview(
        self,
        business_date: date | None = None,
    ) -> dict[str, Any]:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        found_groups = self._daily_response_groups(day, AbsStatus.FOUND)
        account_groups = self._filter_group_taxpayers(
            found_groups,
            lambda taxpayer: (
                taxpayer.get("abs_result") == "found"
                and not taxpayer.get("abs_account_result")
            ),
        )
        generic_cases = self.db.fetch_all(
            """
            SELECT cases.*, uploads.original_filename,
                   (SELECT COUNT(*) FROM taxpayers
                    WHERE taxpayers.case_id = cases.id) AS taxpayer_count
            FROM cases
            JOIN uploads ON uploads.id = cases.upload_id
            WHERE cases.status IN ('needs_review', 'technical_error')
              AND NOT EXISTS (
                  SELECT 1 FROM pages
                  WHERE pages.case_id = cases.id
                    AND pages.status IN ('needs_review', 'technical_error')
              )
              AND NOT EXISTS (
                  SELECT 1 FROM case_match_reviews
                  WHERE status = 'pending'
                    AND (left_case_id = cases.id OR right_case_id = cases.id)
              )
              AND NOT EXISTS (
                  SELECT 1 FROM taxpayers
                  WHERE taxpayers.case_id = cases.id
                    AND taxpayers.abs_result = 'found'
                    AND taxpayers.abs_account_result IS NULL
              )
              AND NOT EXISTS (
                  SELECT 1 FROM taxpayers
                  WHERE taxpayers.case_id = cases.id
                    AND (taxpayers.abs_account_result = 'found'
                         OR taxpayers.odb_result = 'found')
              )
            ORDER BY cases.created_at, cases.id
            """
        )
        return {
            "pages": self.list_review_pages(),
            "match_reviews": self.list_case_match_reviews(),
            "account_groups": account_groups,
            "odb_cases": self._list_odb_cases(),
            "case_tasks": generic_cases,
            "review_history": self._manual_review_history(),
        }

    def _manual_review_history(self, limit: int = 100) -> list[dict[str, Any]]:
        """Return ABS accounts found and all completed period-based ODB checks."""
        safe_limit = max(1, min(int(limit), 200))
        events = self.db.fetch_all(
            """
            SELECT id, entity_type, entity_id, event_type, actor,
                   payload_json, created_at
            FROM audit_events
            WHERE event_type IN (
                'abs_account_manually_checked',
                'odb_checked'
            )
            ORDER BY id DESC
            LIMIT ?
            """,
            (safe_limit * 2,),
        )
        if not events:
            return []

        taxpayer_ids = tuple(
            event["entity_id"]
            for event in events
            if event["event_type"] == "abs_account_manually_checked"
        )
        odb_case_ids = tuple(
            event["entity_id"]
            for event in events
            if event["event_type"] == "odb_checked"
        )

        abs_context: dict[str, dict[str, Any]] = {}
        if taxpayer_ids:
            placeholders = ",".join("?" for _ in taxpayer_ids)
            abs_context = {
                row["taxpayer_id"]: row
                for row in self.db.fetch_all(
                    f"""
                    SELECT taxpayers.id AS taxpayer_id, taxpayers.name,
                           taxpayers.inn, cases.id AS case_id,
                           cases.district_place, uploads.original_filename
                    FROM taxpayers
                    JOIN cases ON cases.id = taxpayers.case_id
                    JOIN uploads ON uploads.id = cases.upload_id
                    WHERE taxpayers.id IN ({placeholders})
                    """,
                    taxpayer_ids,
                )
            }

        odb_context: dict[tuple[str, str], dict[str, Any]] = {}
        if odb_case_ids:
            placeholders = ",".join("?" for _ in odb_case_ids)
            for row in self.db.fetch_all(
                f"""
                SELECT cases.id AS case_id, cases.district_place,
                       uploads.original_filename, taxpayers.name,
                       taxpayers.inn
                FROM cases
                JOIN uploads ON uploads.id = cases.upload_id
                JOIN taxpayers ON taxpayers.case_id = cases.id
                WHERE cases.id IN ({placeholders})
                """,
                odb_case_ids,
            ):
                odb_context[(row["case_id"], row["inn"])] = row

        def payload(event: dict[str, Any]) -> dict[str, Any]:
            try:
                value = json.loads(event["payload_json"])
            except (TypeError, json.JSONDecodeError):
                return {}
            return value if isinstance(value, dict) else {}

        def history_item(
            event: dict[str, Any],
            *,
            kind: str,
            label: str,
            name: str,
            inn: str = "",
            district_place: str = "",
            result: str,
            result_tone: str,
            suffix: str = "",
        ) -> dict[str, Any]:
            return {
                "id": f"{event['id']}{suffix}",
                "kind": kind,
                "kind_label": label,
                "name": name,
                "inn": inn,
                "district_place": district_place,
                "result": result,
                "result_tone": result_tone,
                "actor": event.get("actor") or "",
                "created_at": event["created_at"],
            }

        history: list[dict[str, Any]] = []
        for event in events:
            data = payload(event)
            if event["event_type"] == "abs_account_manually_checked":
                context = abs_context.get(event["entity_id"])
                if not context or data.get("result") != "found":
                    continue
                history.append(
                    history_item(
                        event,
                        kind="abs",
                        label="АБС",
                        name=context["name"],
                        inn=context["inn"],
                        district_place=context.get("district_place") or "",
                        result="Счёт есть",
                        result_tone="found",
                    )
                )
            elif event["event_type"] == "odb_checked":
                results = data.get("results")
                if not isinstance(results, list):
                    continue
                for index, decision in enumerate(results):
                    if not isinstance(decision, dict):
                        continue
                    inn = str(decision.get("inn") or "")
                    context = odb_context.get((event["entity_id"], inn))
                    if not context:
                        continue
                    found = decision.get("result") == "found"
                    history.append(
                        history_item(
                            event,
                            kind="odb",
                            label="ОДБ",
                            name=context["name"],
                            inn=inn,
                            district_place=context.get("district_place") or "",
                            result="Найден" if found else "Не найден",
                            result_tone="found" if found else "not-found",
                            suffix=f"-{index}",
                        )
                    )
            if len(history) >= safe_limit:
                break
        return history[:safe_limit]

    def manual_response_groups(
        self,
        business_date: date | None = None,
    ) -> list[dict[str, Any]]:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        return self._filter_group_taxpayers(
            self._daily_response_groups(day, AbsStatus.FOUND),
            lambda taxpayer: (
                taxpayer.get("abs_account_result") == "found"
                or taxpayer.get("odb_result") == "found"
            ),
        )

    def _daily_response_groups(
        self,
        business_date: date,
        abs_bucket: AbsStatus,
    ) -> list[dict[str, Any]]:
        if abs_bucket == AbsStatus.NOT_FOUND:
            extra_where = (
                "cases.status = 'ready_for_response' "
                "AND (taxpayers.abs_result = 'not_found' "
                "OR taxpayers.abs_result = 'invalid_inn' "
                "OR taxpayers.abs_account_result = 'not_found')"
            )
        else:
            extra_where = (
                "cases.status IN ('needs_review', 'manual_period_rule') "
                "AND (taxpayers.abs_account_result = 'found' "
                "OR (taxpayers.abs_result = 'found' "
                "AND taxpayers.abs_account_result IS NULL) "
                "OR taxpayers.odb_result = 'found') "
                "AND ("
                "cases.status != 'manual_period_rule' "
                "OR NOT EXISTS ("
                "SELECT 1 FROM taxpayers AS pending_odb "
                "WHERE pending_odb.case_id = cases.id "
                "AND pending_odb.odb_result IS NULL"
                ")"
                ")"
            )
        rows = self.db.fetch_all(
            f"""
            SELECT
                cases.id AS case_id,
                cases.district_place,
                cases.recipient_position,
                cases.recipient_full_name,
                cases.recipient_display_name,
                cases.employee_name,
                cases.source_kind AS case_source_kind,
                cases.created_at AS case_created_at,
                uploads.original_filename,
                taxpayers.id AS taxpayer_id,
                taxpayers.name AS taxpayer_name,
                taxpayers.inn AS taxpayer_inn,
                taxpayers.name_source,
                taxpayers.inn_source,
                taxpayers.name_source_reference,
                taxpayers.inn_source_reference,
                taxpayers.display_order,
                taxpayers.abs_result,
                taxpayers.abs_account_result,
                taxpayers.abs_active_account_count,
                taxpayers.abs_closed_account_count,
                taxpayers.odb_result,
                (
                    SELECT pages.id
                    FROM pages
                    WHERE pages.case_id = cases.id
                    ORDER BY CASE WHEN pages.page_type = 'letter' THEN 0 ELSE 1 END,
                             pages.page_number
                    LIMIT 1
                ) AS source_page_id
                ,(
                    SELECT pages.page_number
                    FROM pages
                    WHERE pages.case_id = cases.id
                    ORDER BY CASE WHEN pages.page_type = 'letter' THEN 0 ELSE 1 END,
                             pages.page_number
                    LIMIT 1
                ) AS source_page_number
            FROM cases
            JOIN taxpayers ON taxpayers.case_id = cases.id
            JOIN uploads ON uploads.id = cases.upload_id
            WHERE {extra_where}
              AND NOT EXISTS (
                  SELECT 1 FROM case_match_reviews
                  WHERE case_match_reviews.status = 'pending'
                    AND (
                        case_match_reviews.left_case_id = cases.id
                        OR case_match_reviews.right_case_id = cases.id
                    )
              )
            ORDER BY cases.created_at, cases.id, taxpayers.display_order
            """
        )
        distinct_pairs = {
            tuple(sorted((item["left_case_id"], item["right_case_id"])))
            for item in self.db.fetch_all(
                """
                SELECT left_case_id, right_case_id
                FROM case_match_reviews
                WHERE status = 'resolved' AND resolution = 'distinct'
                """
            )
        }

        grouped: dict[str, dict[str, Any]] = {}
        for row in rows:
            district = clean_location(row.get("district_place"))
            office = self.canonical_gns_office(district)
            office_identity = office["office_key"] if office else district
            if office:
                district = office["district_place"]
            recipient = " ".join(
                (row.get("recipient_full_name") or "").split()
            )
            group_key = self._response_group_key(
                business_date,
                abs_bucket,
                office_identity,
                recipient,
            )
            group = grouped.setdefault(
                group_key,
                {
                    "group_key": group_key,
                    "business_date": business_date.isoformat(),
                    "abs_bucket": str(abs_bucket),
                    "district_place": district,
                    "gns_office_key": office["office_key"] if office else "",
                    "gns_email": self.office_delivery_email(office),
                    "recipient_position": row.get("recipient_position") or "",
                    "recipient_full_name": recipient,
                    "recipient_display_name": (
                        row.get("recipient_display_name") or ""
                    ),
                    "employee_name": row.get("employee_name") or "",
                    "case_ids": [],
                    "taxpayers": [],
                    "issues": [],
                    "warnings": [],
                    "_positions": set(),
                    "_display_names": set(),
                    "_employees": set(),
                    "_taxpayers_by_inn": {},
                    "_conflicts_by_inn": {},
                },
            )
            if row["case_id"] not in group["case_ids"]:
                group["case_ids"].append(row["case_id"])
            group["_positions"].add(
                self._normalize_group_value(row.get("recipient_position"))
            )
            group["_display_names"].add(
                self._normalize_group_value(
                    row.get("recipient_display_name")
                )
            )
            group["_employees"].add(
                self._normalize_group_value(row.get("employee_name"))
            )

            inn = re.sub(r"\D", "", row.get("taxpayer_inn") or "")
            if len(inn) != 14:
                warning = (
                    f"ИНН содержит {len(inn)} цифр вместо 14. "
                    "Сверьте его с исходным письмом."
                )
                if warning not in group["warnings"]:
                    group["warnings"].append(warning)
            raw_name = " ".join(
                (row.get("taxpayer_name") or "").split()
            )
            name = response_taxpayer_name(raw_name, inn)
            existing = group["_taxpayers_by_inn"].get(inn)
            source_record = {
                "taxpayer_id": row["taxpayer_id"],
                "source_case_id": row["case_id"],
                "source_page_id": row.get("source_page_id"),
                "source_page_number": row.get("source_page_number"),
                "original_filename": row.get("original_filename") or "",
                "case_source_kind": row.get("case_source_kind") or "",
                "name_source": row.get("name_source") or "",
                "inn_source": row.get("inn_source") or "",
                "name_source_reference": row.get("name_source_reference"),
                "inn_source_reference": row.get("inn_source_reference"),
                "name": name,
                "inn": inn,
            }
            if existing is None:
                taxpayer = {
                    "taxpayer_id": row["taxpayer_id"],
                    "source_case_id": row["case_id"],
                    "source_page_id": row.get("source_page_id"),
                    "source_page_number": row.get("source_page_number"),
                    "original_filename": row.get("original_filename") or "",
                    "case_source_kind": row.get("case_source_kind") or "",
                    "name_source": row.get("name_source") or "",
                    "inn_source": row.get("inn_source") or "",
                    "name_source_reference": row.get(
                        "name_source_reference"
                    ),
                    "inn_source_reference": row.get(
                        "inn_source_reference"
                    ),
                    "name": name,
                    "inn": inn,
                    "abs_result": row.get("abs_result"),
                    "abs_account_result": row.get("abs_account_result"),
                    "abs_active_account_count": row.get(
                        "abs_active_account_count"
                    ),
                    "abs_closed_account_count": row.get(
                        "abs_closed_account_count"
                    ),
                    "odb_result": row.get("odb_result"),
                }
                group["_taxpayers_by_inn"][inn] = taxpayer
                group["taxpayers"].append(taxpayer)
            elif self._normalize_group_value(existing["name"]) != (
                self._normalize_group_value(name)
            ):
                pair = tuple(
                    sorted((existing["source_case_id"], row["case_id"]))
                )
                if pair in distinct_pairs:
                    taxpayer = {
                        **source_record,
                        "abs_result": row.get("abs_result"),
                        "abs_account_result": row.get("abs_account_result"),
                        "abs_active_account_count": row.get(
                            "abs_active_account_count"
                        ),
                        "abs_closed_account_count": row.get(
                            "abs_closed_account_count"
                        ),
                        "odb_result": row.get("odb_result"),
                    }
                    group["_taxpayers_by_inn"][
                        f"{inn}:{row['case_id']}"
                    ] = taxpayer
                    group["taxpayers"].append(taxpayer)
                    continue
                issue = (
                    f"Для ИНН {inn} указаны разные наименования."
                )
                if issue not in group["issues"]:
                    group["issues"].append(issue)
                conflict_records = group["_conflicts_by_inn"].setdefault(
                    inn, []
                )
                for candidate in (existing, source_record):
                    if not any(
                        record["taxpayer_id"] == candidate["taxpayer_id"]
                        for record in conflict_records
                    ):
                        conflict_records.append(
                            {
                                "taxpayer_id": candidate["taxpayer_id"],
                                "source_case_id": candidate["source_case_id"],
                                "source_page_id": candidate.get(
                                    "source_page_id"
                                ),
                                "source_page_number": candidate.get(
                                    "source_page_number"
                                ),
                                "original_filename": candidate.get(
                                    "original_filename", ""
                                ),
                                "case_source_kind": candidate.get(
                                    "case_source_kind", ""
                                ),
                                "name_source": candidate.get(
                                    "name_source", ""
                                ),
                                "inn_source": candidate.get(
                                    "inn_source", ""
                                ),
                                "name_source_reference": candidate.get(
                                    "name_source_reference"
                                ),
                                "inn_source_reference": candidate.get(
                                    "inn_source_reference"
                                ),
                                "name": candidate["name"],
                                "inn": candidate["inn"],
                            }
                        )

        result: list[dict[str, Any]] = []
        for group in grouped.values():
            if len(group["_positions"]) != 1:
                group["issues"].append(
                    "У одного адресата отличаются должности в письмах."
                )
            if len(group["_display_names"]) != 1:
                group["issues"].append(
                    "У одного адресата отличаются формы обращения."
                )
            if len(group["_employees"]) != 1 or not group["employee_name"]:
                group["issues"].append(
                    "Для группы не определён единый исполнитель банка."
                )
            group["case_count"] = len(group["case_ids"])
            group["taxpayer_count"] = len(group["taxpayers"])
            group["conflicts"] = [
                {"inn": inn, "records": records}
                for inn, records in group["_conflicts_by_inn"].items()
            ]
            group["can_generate"] = bool(
                abs_bucket == AbsStatus.NOT_FOUND
                and group["taxpayers"]
                and not group["issues"]
            )
            for private_key in (
                "_positions",
                "_display_names",
                "_employees",
                "_taxpayers_by_inn",
                "_conflicts_by_inn",
            ):
                group.pop(private_key, None)
            result.append(group)

        return sorted(
            result,
            key=lambda item: (
                item["recipient_full_name"].casefold(),
                item["district_place"].casefold(),
            ),
        )

    def generate_all_ready_daily_responses(
        self,
        business_date: date | None = None,
    ) -> dict[str, Any]:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        ready_groups = [
            item
            for item in self._daily_response_groups(day, AbsStatus.NOT_FOUND)
            if item["can_generate"]
        ]
        created: list[str] = []
        errors: list[dict[str, str]] = []
        for group in ready_groups:
            try:
                group_id, _ = self.generate_daily_response(
                    group["group_key"], business_date=day
                )
                created.append(group_id)
            except (WorkflowValidationError, ValueError) as exc:
                errors.append(
                    {
                        "recipient": group["recipient_display_name"],
                        "message": str(exc),
                    }
                )
        return {"created": created, "errors": errors}

    def get_daily_response_group(
        self,
        group_key: str,
        business_date: date | None = None,
    ) -> dict[str, Any] | None:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        return next(
            (
                item
                for item in self._daily_response_groups(
                    day, AbsStatus.NOT_FOUND
                )
                if item["group_key"] == group_key
            ),
            None,
        )

    def generate_daily_response(
        self,
        group_key: str,
        business_date: date | None = None,
        taxpayers_per_page: int | None = None,
    ) -> tuple[str, Path]:
        with self._case_lock:
            return self._generate_daily_response_locked(
                group_key,
                business_date=business_date,
                taxpayers_per_page=taxpayers_per_page,
            )

    def _generate_daily_response_locked(
        self,
        group_key: str,
        business_date: date | None = None,
        taxpayers_per_page: int | None = None,
    ) -> tuple[str, Path]:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        group = self.get_daily_response_group(group_key, day)
        if not group:
            raise WorkflowValidationError(
                "Группа уже обработана или больше не готова к ответу"
            )
        if not group["can_generate"]:
            raise WorkflowValidationError(
                "Нельзя создать ответ: " + " ".join(group["issues"])
            )

        existing = self._mergeable_created_response_group(group)
        if existing:
            if taxpayers_per_page is not None:
                raise WorkflowValidationError(
                    "Позднее обращение можно добавить только в обычный "
                    "общий ответ без ручного разбиения."
                )
            return self._merge_ready_group_into_created(existing, group)

        group_id = uuid4().hex
        output = self.settings.responses_dir / self._response_word_filename(
            group["recipient_display_name"],
            day.isoformat(),
            group_id,
        )
        case_snapshot = {
            "district_place": group["district_place"],
            "recipient_position": group["recipient_position"],
            "recipient_display_name": group["recipient_display_name"],
            "employee_name": group["employee_name"],
        }
        per_page = taxpayers_per_page or group["taxpayer_count"]
        if per_page < 1 or per_page > group["taxpayer_count"]:
            raise WorkflowValidationError(
                "Количество лиц на странице должно быть от 1 до размера группы"
            )
        _, likely_overflow = self.word.render_pages(
            output,
            case_snapshot,
            group["taxpayers"],
            per_page,
        )
        letter_count = (
            group["taxpayer_count"] + per_page - 1
        ) // per_page

        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                INSERT INTO response_groups(
                    id, business_date, group_key, abs_bucket, status,
                    district_place, recipient_position, recipient_full_name,
                    recipient_display_name, employee_name, taxpayer_count,
                    taxpayers_per_letter, response_path,
                    response_page_overflow, created_at,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    group_id,
                    day.isoformat(),
                    group_key,
                    str(AbsStatus.NOT_FOUND),
                    "created",
                    group["district_place"],
                    group["recipient_position"],
                    group["recipient_full_name"],
                    group["recipient_display_name"],
                    group["employee_name"],
                    group["taxpayer_count"],
                    per_page,
                    str(output),
                    int(likely_overflow),
                    now,
                    now,
                ),
            )
            connection.executemany(
                """
                INSERT INTO response_group_cases(response_group_id, case_id)
                VALUES (?, ?)
                """,
                [(group_id, case_id) for case_id in group["case_ids"]],
            )
            connection.executemany(
                """
                INSERT INTO response_group_taxpayers(
                    response_group_id, display_order, source_case_id,
                    name, inn
                ) VALUES (?, ?, ?, ?, ?)
                """,
                [
                    (
                        group_id,
                        index,
                        taxpayer["source_case_id"],
                        taxpayer["name"],
                        taxpayer["inn"],
                    )
                    for index, taxpayer in enumerate(
                        group["taxpayers"], 1
                    )
                ],
            )
            connection.executemany(
                """
                INSERT INTO response_letters(
                    id, response_group_id, letter_order,
                    taxpayer_start_order, taxpayer_count,
                    created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    (
                        uuid4().hex,
                        group_id,
                        letter_order,
                        (letter_order - 1) * per_page + 1,
                        min(
                            per_page,
                            group["taxpayer_count"]
                            - (letter_order - 1) * per_page,
                        ),
                        now,
                        now,
                    )
                    for letter_order in range(1, letter_count + 1)
                ],
            )
            connection.executemany(
                """
                UPDATE cases
                SET status = ?, response_status = ?,
                    response_path = ?, updated_at = ?
                WHERE id = ? AND status = ?
                """,
                [
                    (
                        CaseStatus.RESPONSE_CREATED,
                        "grouped",
                        str(output),
                        now,
                        case_id,
                        CaseStatus.READY_FOR_RESPONSE,
                    )
                    for case_id in group["case_ids"]
                ],
            )

        self.db.audit(
            "response_group",
            group_id,
            "grouped_response_created",
            {
                "business_date": day.isoformat(),
                "case_count": group["case_count"],
                "taxpayer_count": group["taxpayer_count"],
                "filename": output.name,
                "taxpayers_per_page": per_page,
                "letter_count": letter_count,
            },
        )
        return group_id, output

    def _mergeable_created_response_group(
        self,
        ready_group: dict[str, Any],
    ) -> dict[str, Any] | None:
        candidates = self.db.fetch_all(
            """
            SELECT response_groups.*
            FROM response_groups
            WHERE response_groups.business_date = ?
              AND response_groups.group_key = ?
              AND response_groups.abs_bucket = ?
              AND response_groups.status = 'created'
              AND (
                  SELECT COUNT(*) FROM response_letters
                  WHERE response_letters.response_group_id = response_groups.id
              ) = 1
              AND NOT EXISTS (
                  SELECT 1
                  FROM signed_response_scans
                  JOIN response_letters
                    ON response_letters.id =
                       signed_response_scans.response_letter_id
                  WHERE response_letters.response_group_id = response_groups.id
                    AND signed_response_scans.status != 'superseded'
              )
              AND NOT EXISTS (
                  SELECT 1
                  FROM outlook_outgoing_messages
                  JOIN response_letters
                    ON response_letters.id =
                       outlook_outgoing_messages.response_letter_id
                  WHERE response_letters.response_group_id = response_groups.id
              )
              AND NOT EXISTS (
                  SELECT 1
                  FROM signed_scan_sessions
                  JOIN response_letters
                    ON response_letters.id =
                       signed_scan_sessions.response_letter_id
                  JOIN signed_scan_session_pages
                    ON signed_scan_session_pages.session_id =
                       signed_scan_sessions.id
                  WHERE response_letters.response_group_id = response_groups.id
                    AND signed_scan_sessions.status IN (
                        'collecting', 'technical_error'
                    )
                    AND signed_scan_session_pages.status = 'active'
              )
            ORDER BY response_groups.created_at, response_groups.id
            """,
            (
                ready_group["business_date"],
                ready_group["group_key"],
                str(AbsStatus.NOT_FOUND),
            ),
        )
        if not candidates:
            return None
        existing = candidates[0]
        for field in (
            "district_place",
            "recipient_position",
            "recipient_display_name",
            "employee_name",
        ):
            if self._normalize_group_value(existing.get(field)) != (
                self._normalize_group_value(ready_group.get(field))
            ):
                raise WorkflowValidationError(
                    "У уже созданного ответа этому адресату отличаются "
                    "реквизиты. Сначала проверьте их вручную."
                )
        return existing

    def _merge_ready_group_into_created(
        self,
        existing_group: dict[str, Any],
        ready_group: dict[str, Any],
    ) -> tuple[str, Path]:
        group_id = existing_group["id"]
        stored_group, stored_taxpayers, letters = (
            self._response_group_render_data(group_id)
        )
        if len(letters) != 1:
            raise WorkflowValidationError(
                "Позднее обращение нельзя безопасно добавить в разделённый Word."
            )
        output = ensure_within(
            Path(stored_group["response_path"]),
            self.settings.responses_dir,
        )
        if not output.exists():
            raise WorkflowValidationError(
                "Файл уже созданного ответа отсутствует."
            )

        combined = [dict(taxpayer) for taxpayer in stored_taxpayers]
        by_inn = {
            str(taxpayer["inn"]): self._normalize_group_value(taxpayer["name"])
            for taxpayer in combined
        }
        by_name = {
            self._normalize_group_value(taxpayer["name"]): str(taxpayer["inn"])
            for taxpayer in combined
        }
        additions: list[dict[str, Any]] = []
        for taxpayer in ready_group["taxpayers"]:
            inn = str(taxpayer["inn"])
            normalized_name = self._normalize_group_value(taxpayer["name"])
            if inn in by_inn:
                if by_inn[inn] != normalized_name:
                    raise WorkflowValidationError(
                        "У одинакового ИНН отличаются наименования. "
                        "Требуется ручная проверка."
                    )
                continue
            if normalized_name in by_name and by_name[normalized_name] != inn:
                raise WorkflowValidationError(
                    "У одинакового лица отличаются ИНН. "
                    "Требуется ручная проверка."
                )
            addition = {
                "source_case_id": taxpayer["source_case_id"],
                "name": taxpayer["name"],
                "inn": inn,
            }
            additions.append(addition)
            combined.append(addition)
            by_inn[inn] = normalized_name
            by_name[normalized_name] = inn

        temporary: Path | None = None
        backup: Path | None = None
        likely_overflow = bool(stored_group["response_page_overflow"])
        if additions:
            temporary = output.with_name(
                f".{output.stem}-late-merge-{uuid4().hex}.docx"
            )
            try:
                _, likely_overflow = self.word.render_pages(
                    temporary,
                    {
                        "district_place": stored_group["district_place"],
                        "recipient_position": stored_group["recipient_position"],
                        "recipient_display_name": stored_group[
                            "recipient_display_name"
                        ],
                        "employee_name": stored_group["employee_name"],
                    },
                    combined,
                    len(combined),
                    outgoing_numbers=[
                        str(letters[0]["outgoing_number"] or "")
                    ],
                )
                backup = output.with_name(
                    f".{output.stem}-before-late-merge-{uuid4().hex}.docx"
                )
                shutil.copy2(output, backup)
            except (OSError, WordTemplateError) as exc:
                temporary.unlink(missing_ok=True)
                raise WorkflowValidationError(
                    "Не удалось обновить общий Word."
                ) from exc

        case_ids = list(dict.fromkeys(ready_group["case_ids"]))
        placeholders = ",".join("?" for _ in case_ids)
        now = utc_now()
        replaced = False
        try:
            with self.db.connect() as connection:
                # Блокируем конкурирующую запись и повторно проверяем границу
                # сканирования/Outlook: между предварительным поиском группы и
                # этой точкой сотрудник мог начать отправку.
                connection.execute("BEGIN IMMEDIATE")
                merge_cutoff = connection.execute(
                    """
                    SELECT 1
                    FROM response_groups
                    WHERE response_groups.id = ?
                      AND response_groups.status = 'created'
                      AND (
                          EXISTS (
                              SELECT 1
                              FROM signed_response_scans
                              JOIN response_letters
                                ON response_letters.id =
                                   signed_response_scans.response_letter_id
                              WHERE response_letters.response_group_id =
                                    response_groups.id
                                AND signed_response_scans.status != 'superseded'
                          )
                          OR EXISTS (
                              SELECT 1
                              FROM outlook_outgoing_messages
                              JOIN response_letters
                                ON response_letters.id =
                                   outlook_outgoing_messages.response_letter_id
                              WHERE response_letters.response_group_id =
                                    response_groups.id
                          )
                          OR EXISTS (
                              SELECT 1
                              FROM signed_scan_sessions
                              JOIN response_letters
                                ON response_letters.id =
                                   signed_scan_sessions.response_letter_id
                              JOIN signed_scan_session_pages
                                ON signed_scan_session_pages.session_id =
                                   signed_scan_sessions.id
                              WHERE response_letters.response_group_id =
                                    response_groups.id
                                AND signed_scan_sessions.status IN (
                                    'collecting', 'technical_error'
                                )
                                AND signed_scan_session_pages.status = 'active'
                          )
                      )
                    """,
                    (group_id,),
                ).fetchone()
                current_group = connection.execute(
                    "SELECT status FROM response_groups WHERE id = ?",
                    (group_id,),
                ).fetchone()
                if not current_group or current_group["status"] != "created":
                    raise WorkflowValidationError(
                        "Существующий ответ уже изменился. Обновите страницу."
                    )
                if merge_cutoff:
                    raise WorkflowValidationError(
                        "Ответ уже начали сканировать или отправлять. "
                        "Новое обращение сохранено отдельно; повторите создание."
                    )
                current_cases = connection.execute(
                    "SELECT id, status FROM cases "
                    f"WHERE id IN ({placeholders})",
                    case_ids,
                ).fetchall()
                if len(current_cases) != len(case_ids) or any(
                    row["status"] != CaseStatus.READY_FOR_RESPONSE
                    for row in current_cases
                ):
                    raise WorkflowValidationError(
                        "Состав готовых обращений изменился. Обновите страницу."
                    )

                if additions and temporary is not None:
                    temporary.replace(output)
                    replaced = True

                connection.executemany(
                    "INSERT OR IGNORE INTO response_group_cases("
                    "response_group_id, case_id) VALUES (?, ?)",
                    [(group_id, case_id) for case_id in case_ids],
                )
                start_order = len(stored_taxpayers) + 1
                connection.executemany(
                    """
                    INSERT INTO response_group_taxpayers(
                        response_group_id, display_order, source_case_id,
                        name, inn
                    ) VALUES (?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            group_id,
                            start_order + index,
                            taxpayer["source_case_id"],
                            taxpayer["name"],
                            taxpayer["inn"],
                        )
                        for index, taxpayer in enumerate(additions)
                    ],
                )
                if additions:
                    connection.execute(
                        """
                        UPDATE response_groups
                        SET taxpayer_count = ?, taxpayers_per_letter = ?,
                            response_page_overflow = ?,
                            opened_for_print_at = NULL,
                            word_reopen_required = 1, updated_at = ?
                        WHERE id = ? AND status = 'created'
                        """,
                        (
                            len(combined),
                            len(combined),
                            int(likely_overflow),
                            now,
                            group_id,
                        ),
                    )
                    connection.execute(
                        """
                        UPDATE response_letters
                        SET taxpayer_start_order = 1, taxpayer_count = ?,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (len(combined), now, letters[0]["id"]),
                    )
                    connection.execute(
                        """
                        UPDATE signed_scan_sessions
                        SET status = 'cancelled', error_message = NULL,
                            updated_at = ?
                        WHERE response_letter_id = ?
                          AND status IN ('collecting', 'technical_error')
                          AND NOT EXISTS (
                              SELECT 1 FROM signed_scan_session_pages
                              WHERE signed_scan_session_pages.session_id =
                                    signed_scan_sessions.id
                                AND signed_scan_session_pages.status = 'active'
                          )
                        """,
                        (now, letters[0]["id"]),
                    )
                connection.executemany(
                    """
                    UPDATE cases
                    SET status = ?, response_status = ?,
                        response_path = ?, updated_at = ?
                    WHERE id = ? AND status = ?
                    """,
                    [
                        (
                            CaseStatus.RESPONSE_CREATED,
                            "grouped",
                            str(output),
                            now,
                            case_id,
                            CaseStatus.READY_FOR_RESPONSE,
                        )
                        for case_id in case_ids
                    ],
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('response_group', ?,
                              'late_cases_merged_into_response',
                              'Сотрудник', ?, ?)
                    """,
                    (
                        group_id,
                        json.dumps(
                            {
                                "added_case_count": len(case_ids),
                                "added_taxpayer_count": len(additions),
                                "duplicate_taxpayer_count": (
                                    len(ready_group["taxpayers"])
                                    - len(additions)
                                ),
                                "word_regenerated": bool(additions),
                            },
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
        except Exception:
            if replaced and backup is not None and backup.exists():
                try:
                    backup.replace(output)
                except OSError:
                    pass
            elif backup is not None:
                backup.unlink(missing_ok=True)
            raise
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

        return group_id, output

    def get_response_group(
        self, group_id: str
    ) -> dict[str, Any] | None:
        return self.db.fetch_one(
            "SELECT * FROM response_groups WHERE id = ?",
            (group_id,),
        )

    @staticmethod
    def _filename_part(value: str | None, fallback: str) -> str:
        normalized = re.sub(r"\s+", " ", str(value or "")).strip()
        cleaned = Path(sanitize_filename(normalized or fallback)).stem
        return cleaned[:72].strip(" ._") or fallback

    @classmethod
    def _response_word_filename(
        cls,
        recipient_display_name: str | None,
        date_label: str,
        unique_id: str,
    ) -> str:
        recipient = cls._filename_part(recipient_display_name, "адресат")
        return sanitize_filename(
            f"Ответ_{recipient}_{date_label}_{unique_id[:6]}.docx"
        )

    @classmethod
    def _signed_scan_filename(cls, letter: dict[str, Any]) -> str:
        recipient = cls._filename_part(
            letter.get("recipient_display_name"), "адресат"
        )
        outgoing_number = str(letter.get("outgoing_number") or "").strip()
        document_mark = (
            f"исх-{cls._filename_part(outgoing_number, 'номер')}"
            if outgoing_number
            else f"от-{letter.get('business_date') or 'без-даты'}"
        )
        return sanitize_filename(
            f"Скан_ответа_{recipient}_{document_mark}.pdf"
        )

    def response_group_download_filename(self, group_id: str) -> str:
        group = self.get_response_group(group_id)
        if not group:
            raise WorkflowValidationError("Общий ответ не найден")
        letters = self.db.fetch_all(
            """
            SELECT outgoing_number FROM response_letters
            WHERE response_group_id = ? ORDER BY letter_order
            """,
            (group_id,),
        )
        numbers = [
            str(letter["outgoing_number"]).strip()
            for letter in letters
            if str(letter.get("outgoing_number") or "").strip()
        ]
        suffix = (
            f"исх-{numbers[0]}"
            if len(letters) == 1 and len(numbers) == 1
            else str(group["business_date"])
        )
        recipient = self._filename_part(
            group.get("recipient_display_name"), "адресат"
        )
        return sanitize_filename(f"Ответ_{recipient}_{suffix}.docx")

    def signed_scan_download_filename(self, scan_id: str) -> str:
        scan = self.get_signed_response_scan(scan_id)
        if not scan:
            raise WorkflowValidationError("Подписанный скан не найден")
        letter = self.get_response_letter(scan["response_letter_id"])
        if not letter:
            raise WorkflowValidationError("Готовое письмо не найдено")
        return self._signed_scan_filename(letter)

    def mark_response_group_opened_for_print(
        self,
        group_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> None:
        if not self.get_response_group(group_id):
            raise WorkflowValidationError("Общий ответ не найден")
        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                UPDATE response_groups
                SET opened_for_print_at = ?, word_reopen_required = 0,
                    updated_at = ?
                WHERE id = ?
                """,
                (now, now, group_id),
            )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('response_group', ?, 'response_opened_for_print',
                          ?, '{}', ?)
                """,
                (group_id, actor, now),
            )

    @staticmethod
    def _canonical_outgoing_number(
        value: str | int | None,
        *,
        allow_empty: bool = False,
    ) -> str | None:
        text = str(value or "").strip()
        if not text:
            if allow_empty:
                return None
            raise WorkflowValidationError("Введите первый исходящий номер")
        if not text.isdigit():
            raise WorkflowValidationError(
                "Исходящий номер должен состоять только из цифр"
            )
        number = int(text)
        if number < 1 or number > 999_999_999:
            raise WorkflowValidationError(
                "Исходящий номер должен быть от 1 до 999999999"
            )
        return str(number)

    def list_outgoing_letters(
        self,
        business_date: date | None = None,
        *,
        only_unnumbered: bool = False,
    ) -> list[dict[str, Any]]:
        date_filter = ""
        parameters: tuple[Any, ...] = ()
        if business_date is not None:
            date_filter = "AND response_groups.business_date = ?"
            parameters = (business_date.isoformat(),)
        unnumbered_filter = (
            "AND response_letters.outgoing_number IS NULL"
            if only_unnumbered
            else ""
        )
        return self.db.fetch_all(
            f"""
            SELECT response_letters.*,
                   response_groups.business_date,
                   response_groups.recipient_display_name,
                   response_groups.district_place,
                   response_groups.response_path,
                   response_groups.created_at AS group_created_at
            FROM response_letters
            JOIN response_groups
              ON response_groups.id = response_letters.response_group_id
            WHERE response_groups.status = 'created'
              {date_filter}
              {unnumbered_filter}
            ORDER BY response_groups.created_at,
                     response_groups.rowid, response_letters.letter_order
            """,
            parameters,
        )

    def _propose_outgoing_numbers(
        self,
        first_number: str,
        letters: list[dict[str, Any]],
    ) -> tuple[dict[str, str], int]:
        """Assign only currently free numbers, preserving prior choices."""
        occupied_rows = self.db.fetch_all(
            """
            SELECT outgoing_number
            FROM response_letters
            WHERE outgoing_number IS NOT NULL AND outgoing_number != ''
            """
        )
        occupied: set[int] = set()
        for row in occupied_rows:
            number = self._canonical_outgoing_number(row["outgoing_number"])
            assert number is not None
            occupied.add(int(number))

        candidate = int(first_number)
        skipped_occupied_count = 0
        proposals: dict[str, str] = {}
        for letter in letters:
            while candidate in occupied:
                candidate += 1
                skipped_occupied_count += 1
            if candidate > 999_999_999:
                raise WorkflowValidationError(
                    "Диапазон исходящих номеров превышает 999999999"
                )
            proposals[letter["id"]] = str(candidate)
            occupied.add(candidate)
            candidate += 1
        return proposals, skipped_occupied_count

    def preview_outgoing_numbers(
        self,
        first_number: str | int,
        business_date: date | None = None,
    ) -> dict[str, Any]:
        day = business_date or datetime.now(self.BUSINESS_TIMEZONE).date()
        first = self._canonical_outgoing_number(first_number)
        letters = self.list_outgoing_letters(
            business_date,
            only_unnumbered=True,
        )
        updates, skipped_occupied_count = self._propose_outgoing_numbers(
            first, letters
        )
        proposals: list[dict[str, Any]] = []
        for letter in letters:
            proposal = dict(letter)
            proposal["proposed_number"] = updates[letter["id"]]
            proposals.append(proposal)

        proposed_numbers = [item["proposed_number"] for item in proposals]
        return {
            "business_date": day.isoformat(),
            "first_number": first,
            "last_number": (
                proposed_numbers[-1] if proposed_numbers else first
            ),
            "count": len(proposals),
            "letters": proposals,
            "skipped_occupied_count": skipped_occupied_count,
        }

    def _response_group_render_data(
        self,
        group_id: str,
    ) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
        group = self.get_response_group(group_id)
        if not group:
            raise WorkflowValidationError("Готовый ответ не найден")
        taxpayers = self.db.fetch_all(
            """
            SELECT name, inn
            FROM response_group_taxpayers
            WHERE response_group_id = ?
            ORDER BY display_order
            """,
            (group_id,),
        )
        letters = self.db.fetch_all(
            """
            SELECT * FROM response_letters
            WHERE response_group_id = ?
            ORDER BY letter_order
            """,
            (group_id,),
        )
        if not taxpayers or not letters:
            raise WorkflowValidationError(
                "Не удалось восстановить состав готового ответа"
            )
        return group, taxpayers, letters

    def _render_group_number_update(
        self,
        group_id: str,
        overrides: dict[str, str | None],
    ) -> tuple[Path, Path, bool]:
        group, taxpayers, letters = self._response_group_render_data(group_id)
        output = ensure_within(
            Path(group["response_path"]),
            self.settings.responses_dir,
        )
        if not output.exists():
            raise WorkflowValidationError("Файл готового ответа отсутствует")
        numbers = [
            str(overrides.get(letter["id"], letter["outgoing_number"]) or "")
            for letter in letters
        ]
        temporary = output.with_name(
            f".{output.stem}-numbering-{uuid4().hex}.docx"
        )
        case_snapshot = {
            "district_place": group["district_place"],
            "recipient_position": group["recipient_position"],
            "recipient_display_name": group["recipient_display_name"],
            "employee_name": group["employee_name"],
        }
        try:
            _, likely_overflow = self.word.render_pages(
                temporary,
                case_snapshot,
                taxpayers,
                int(group["taxpayers_per_letter"] or len(taxpayers)),
                outgoing_numbers=numbers,
            )
        except (OSError, WordTemplateError) as exc:
            temporary.unlink(missing_ok=True)
            raise WorkflowValidationError(
                "Не удалось обновить исходящий номер в Word"
            ) from exc
        return temporary, output, likely_overflow

    def _apply_outgoing_number_updates(
        self,
        updates: dict[str, str | None],
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        if not updates:
            raise WorkflowValidationError("Не выбраны письма для нумерации")
        letter_ids = list(updates)
        placeholders = ",".join("?" for _ in letter_ids)
        rows = self.db.fetch_all(
            """
            SELECT response_letters.*, response_groups.response_path
            FROM response_letters
            JOIN response_groups
              ON response_groups.id = response_letters.response_group_id
            """
            f"WHERE response_letters.id IN ({placeholders})",
            letter_ids,
        )
        if len(rows) != len(letter_ids):
            raise WorkflowValidationError("Одно из готовых писем не найдено")

        normalized = {
            letter_id: self._canonical_outgoing_number(
                value, allow_empty=True
            )
            for letter_id, value in updates.items()
        }
        requested_numbers = [value for value in normalized.values() if value]
        if len(requested_numbers) != len(set(requested_numbers)):
            raise WorkflowValidationError(
                "В выбранном диапазоне повторяется исходящий номер"
            )
        if requested_numbers:
            number_placeholders = ",".join("?" for _ in requested_numbers)
            id_placeholders = ",".join("?" for _ in letter_ids)
            occupied = self.db.fetch_one(
                "SELECT outgoing_number FROM response_letters "
                f"WHERE outgoing_number IN ({number_placeholders}) "
                f"AND id NOT IN ({id_placeholders}) LIMIT 1",
                [*requested_numbers, *letter_ids],
            )
            if occupied:
                raise WorkflowValidationError(
                    "Исходящий номер "
                    f"{occupied['outgoing_number']} уже используется"
                )

        before = {row["id"]: row["outgoing_number"] for row in rows}
        changed_rows = [
            row
            for row in rows
            if before[row["id"]] != normalized[row["id"]]
        ]
        if not changed_rows:
            return {
                "letter_count": 0,
                "group_count": 0,
                "superseded_scan_count": 0,
                "word_reopen_required": False,
            }

        changed_letter_ids = [row["id"] for row in changed_rows]
        changed_placeholders = ",".join("?" for _ in changed_letter_ids)
        # Нельзя менять номер во время получения листов: физическая копия
        # Word уже могла попасть на сканер.
        active_session = self.db.fetch_one(
            f"""
            SELECT id FROM signed_scan_sessions
            WHERE response_letter_id IN ({changed_placeholders})
              AND status IN ('collecting', 'technical_error')
            LIMIT 1
            """,
            changed_letter_ids,
        )
        if active_session:
            raise WorkflowValidationError(
                "Сначала завершите или отмените текущее сканирование."
            )
        # Черновик (включая уже отправленное письмо) — внешний объект. Его
        # нельзя оставить с подписанным PDF, который перестанет совпадать с
        # обновлённым Word.
        existing_outlook = self.db.fetch_one(
            f"""
            SELECT id FROM outlook_outgoing_messages
            WHERE response_letter_id IN ({changed_placeholders})
            LIMIT 1
            """,
            changed_letter_ids,
        )
        if existing_outlook:
            raise WorkflowValidationError(
                "Нельзя изменить исходящий номер: для письма уже создан "
                "черновик Outlook или подтверждена отправка."
            )

        group_ids = sorted(
            {row["response_group_id"] for row in changed_rows}
        )
        rendered: dict[str, tuple[Path, Path, bool]] = {}
        backups: dict[str, Path] = {}
        preserved_backups: set[str] = set()
        superseded_scan_count = 0
        try:
            for group_id in group_ids:
                rendered[group_id] = self._render_group_number_update(
                    group_id, normalized
                )
            for group_id, (_, output, _) in rendered.items():
                backup = output.with_name(
                    f".{output.stem}-before-numbering-{uuid4().hex}.docx"
                )
                shutil.copy2(output, backup)
                backups[group_id] = backup

            now = utc_now()
            with self.db.connect() as connection:
                current_rows = connection.execute(
                    "SELECT id, outgoing_number FROM response_letters "
                    f"WHERE id IN ({placeholders})",
                    letter_ids,
                ).fetchall()
                current = {
                    row["id"]: row["outgoing_number"] for row in current_rows
                }
                if current != before:
                    raise WorkflowValidationError(
                        "Список исходящих номеров уже изменился. Обновите страницу."
                    )
                active_session = connection.execute(
                    f"""
                    SELECT id FROM signed_scan_sessions
                    WHERE response_letter_id IN ({changed_placeholders})
                      AND status IN ('collecting', 'technical_error')
                    LIMIT 1
                    """,
                    changed_letter_ids,
                ).fetchone()
                if active_session:
                    raise WorkflowValidationError(
                        "Сначала завершите или отмените текущее сканирование."
                    )
                existing_outlook = connection.execute(
                    f"""
                    SELECT id FROM outlook_outgoing_messages
                    WHERE response_letter_id IN ({changed_placeholders})
                    LIMIT 1
                    """,
                    changed_letter_ids,
                ).fetchone()
                if existing_outlook:
                    raise WorkflowValidationError(
                        "Нельзя изменить исходящий номер: для письма уже создан "
                        "черновик Outlook или подтверждена отправка."
                    )
                if requested_numbers:
                    number_placeholders = ",".join(
                        "?" for _ in requested_numbers
                    )
                    id_placeholders = ",".join("?" for _ in letter_ids)
                    occupied = connection.execute(
                        "SELECT outgoing_number FROM response_letters "
                        f"WHERE outgoing_number IN ({number_placeholders}) "
                        f"AND id NOT IN ({id_placeholders}) LIMIT 1",
                        [*requested_numbers, *letter_ids],
                    ).fetchone()
                    if occupied:
                        raise WorkflowValidationError(
                            "Исходящий номер "
                            f"{occupied['outgoing_number']} уже используется"
                        )
                active_scans = connection.execute(
                    f"""
                    SELECT id, response_letter_id
                    FROM signed_response_scans
                    WHERE response_letter_id IN ({changed_placeholders})
                      AND status IN ('needs_confirmation', 'confirmed')
                    """,
                    changed_letter_ids,
                ).fetchall()
                superseded_scan_count = len(active_scans)
                for row in changed_rows:
                    letter_id = row["id"]
                    new_number = normalized[letter_id]
                    connection.execute(
                        """
                        UPDATE response_letters
                        SET outgoing_number = ?, assigned_at = ?, updated_at = ?
                        WHERE id = ?
                        """,
                        (
                            new_number,
                            now if new_number else None,
                            now,
                            letter_id,
                        ),
                    )
                    event_type = (
                        "outgoing_number_assigned"
                        if not before[letter_id] and new_number
                        else "outgoing_number_changed"
                    )
                    connection.execute(
                        """
                        INSERT INTO audit_events(
                            entity_type, entity_id, event_type,
                            actor, payload_json, created_at
                        ) VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            "response_letter",
                            letter_id,
                            event_type,
                            actor,
                            json.dumps(
                                {
                                    "response_group_id": row["response_group_id"],
                                    "before": before[letter_id],
                                    "after": new_number,
                                },
                                ensure_ascii=False,
                            ),
                            now,
                        ),
                    )
                if active_scans:
                    connection.execute(
                        f"""
                        UPDATE signed_response_scans
                        SET status = 'superseded', updated_at = ?
                        WHERE response_letter_id IN ({changed_placeholders})
                          AND status IN ('needs_confirmation', 'confirmed')
                        """,
                        (now, *changed_letter_ids),
                    )
                    for scan in active_scans:
                        connection.execute(
                            """
                            INSERT INTO audit_events(
                                entity_type, entity_id, event_type,
                                actor, payload_json, created_at
                            ) VALUES (?, ?, ?, ?, ?, ?)
                            """,
                            (
                                "signed_response_scan",
                                scan["id"],
                                "signed_scan_superseded_by_outgoing_number",
                                actor,
                                json.dumps(
                                    {
                                        "response_letter_id": scan[
                                            "response_letter_id"
                                        ],
                                        "reason": "outgoing_number_changed",
                                    },
                                    ensure_ascii=False,
                                ),
                                now,
                            ),
                        )
                for group_id, (temporary, output, overflow) in rendered.items():
                    temporary.replace(output)
                    connection.execute(
                        """
                        UPDATE response_groups
                        SET response_page_overflow = ?,
                            opened_for_print_at = NULL,
                            word_reopen_required = 1, updated_at = ?
                        WHERE id = ?
                        """,
                        (int(overflow), now, group_id),
                    )
        except sqlite3.IntegrityError as exc:
            preserved_backups.update(
                self._restore_numbering_backups(backups, rendered)
            )
            raise WorkflowValidationError(
                "Один из исходящих номеров уже используется"
            ) from exc
        except OSError as exc:
            preserved_backups.update(
                self._restore_numbering_backups(backups, rendered)
            )
            raise WorkflowValidationError(
                "Не удалось обновить Word. Закройте открытый файл и повторите."
            ) from exc
        except Exception:
            preserved_backups.update(
                self._restore_numbering_backups(backups, rendered)
            )
            raise
        finally:
            for temporary, _, _ in rendered.values():
                temporary.unlink(missing_ok=True)
            for group_id, backup in backups.items():
                if group_id not in preserved_backups:
                    backup.unlink(missing_ok=True)

        return {
            "letter_count": len(changed_letter_ids),
            "group_count": len(group_ids),
            "superseded_scan_count": superseded_scan_count,
            "word_reopen_required": True,
        }

    @staticmethod
    def _restore_numbering_backups(
        backups: dict[str, Path],
        rendered: dict[str, tuple[Path, Path, bool]],
    ) -> set[str]:
        failed: set[str] = set()
        for group_id, backup in backups.items():
            if not backup.exists():
                continue
            try:
                backup.replace(rendered[group_id][1])
            except OSError:
                # Не маскируем исходную ошибку и сохраняем резервную копию.
                failed.add(group_id)
        return failed

    def assign_outgoing_numbers(
        self,
        first_number: str | int,
        expected_letter_ids: list[str],
        business_date: date | None = None,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        first = self._canonical_outgoing_number(first_number)
        letters = self.list_outgoing_letters(
            business_date,
            only_unnumbered=True,
        )
        current_ids = [letter["id"] for letter in letters]
        if not current_ids:
            raise WorkflowValidationError(
                "Нет готовых писем без исходящего номера"
            )
        if not expected_letter_ids or len(expected_letter_ids) != len(
            set(expected_letter_ids)
        ):
            raise WorkflowValidationError(
                "Список готовых писем изменился. Обновите страницу."
            )
        if expected_letter_ids != current_ids:
            raise WorkflowValidationError(
                "Список готовых писем изменился. Обновите страницу."
            )
        updates, skipped_occupied_count = self._propose_outgoing_numbers(
            first, letters
        )
        summary = self._apply_outgoing_number_updates(updates, actor=actor)
        summary.update(
            {
                "first_number": first,
                "first_assigned_number": updates[current_ids[0]],
                "last_number": updates[current_ids[-1]],
                "skipped_occupied_count": skipped_occupied_count,
            }
        )
        return summary

    def set_outgoing_number(
        self,
        letter_id: str,
        outgoing_number: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        return self._apply_outgoing_number_updates(
            {letter_id: outgoing_number},
            actor=actor,
        )

    def get_response_letter(self, letter_id: str) -> dict[str, Any] | None:
        return self.db.fetch_one(
            """
            SELECT response_letters.*,
                   response_groups.recipient_display_name,
                   response_groups.recipient_full_name,
                   response_groups.district_place,
                   response_groups.response_path,
                   response_groups.business_date
            FROM response_letters
            JOIN response_groups
              ON response_groups.id = response_letters.response_group_id
            WHERE response_letters.id = ?
            """,
            (letter_id,),
        )

    def get_signed_response_scan(
        self, scan_id: str
    ) -> dict[str, Any] | None:
        return self.db.fetch_one(
            "SELECT * FROM signed_response_scans WHERE id = ?",
            (scan_id,),
        )

    def get_confirmed_signed_response_scan(
        self, letter_id: str
    ) -> dict[str, Any] | None:
        return self.db.fetch_one(
            """
            SELECT * FROM signed_response_scans
            WHERE response_letter_id = ? AND status = 'confirmed'
            ORDER BY confirmed_at DESC, created_at DESC
            LIMIT 1
            """,
            (letter_id,),
        )

    def get_signed_scan_session(
        self, session_id: str
    ) -> dict[str, Any] | None:
        session = self.db.fetch_one(
            "SELECT * FROM signed_scan_sessions WHERE id = ?",
            (session_id,),
        )
        if not session:
            return None
        session["pages"] = self.db.fetch_all(
            """
            SELECT * FROM signed_scan_session_pages
            WHERE session_id = ? AND status = 'active'
            ORDER BY page_order, created_at, id
            """,
            (session_id,),
        )
        return session

    def get_active_signed_scan_session(
        self, letter_id: str
    ) -> dict[str, Any] | None:
        row = self.db.fetch_one(
            """
            SELECT id FROM signed_scan_sessions
            WHERE response_letter_id = ?
              AND status IN ('collecting', 'technical_error')
            ORDER BY created_at DESC
            LIMIT 1
            """,
            (letter_id,),
        )
        return self.get_signed_scan_session(row["id"]) if row else None

    def start_signed_scan_session(
        self,
        letter_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        if not self.get_response_letter(letter_id):
            raise WorkflowValidationError("Готовое письмо не найдено")
        existing = self.get_active_signed_scan_session(letter_id)
        if existing:
            return existing

        session_id = uuid4().hex
        session_dir = ensure_within(
            self.settings.runtime_dir
            / "signed_scans"
            / letter_id
            / "sessions"
            / session_id,
            self.settings.runtime_dir,
        )
        session_dir.mkdir(parents=True, exist_ok=False)
        now = utc_now()
        try:
            with self.db.connect() as connection:
                connection.execute(
                    """
                    INSERT INTO signed_scan_sessions(
                        id, response_letter_id, status, page_count,
                        created_by, created_at, updated_at
                    ) VALUES (?, ?, 'collecting', 0, ?, ?, ?)
                    """,
                    (session_id, letter_id, actor, now, now),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('signed_scan_session', ?,
                              'signed_scan_session_started', ?, ?, ?)
                    """,
                    (
                        session_id,
                        actor,
                        json.dumps(
                            {"response_letter_id": letter_id},
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
        except Exception:
            shutil.rmtree(session_dir, ignore_errors=True)
            raise
        return self.get_signed_scan_session(session_id) or {}

    def _scan_session_page_destination(
        self,
        session: dict[str, Any],
        page_id: str,
    ) -> Path:
        return ensure_within(
            self.settings.runtime_dir
            / "signed_scans"
            / session["response_letter_id"]
            / "sessions"
            / session["id"]
            / "pages"
            / f"{page_id}.png",
            self.settings.runtime_dir,
        )

    def _set_signed_scan_session_error(
        self,
        session_id: str,
        message: str,
        *,
        actor: str,
    ) -> None:
        safe_message = message.strip() or "Не удалось получить лист со сканера"
        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                UPDATE signed_scan_sessions
                SET status = 'technical_error', error_message = ?, updated_at = ?
                WHERE id = ? AND status IN ('collecting', 'technical_error')
                """,
                (safe_message, now, session_id),
            )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('signed_scan_session', ?,
                          'signed_scan_page_failed', ?, ?, ?)
                """,
                (
                    session_id,
                    actor,
                    json.dumps({"message": safe_message}, ensure_ascii=False),
                    now,
                ),
            )

    def acquire_signed_scan_session_page(
        self,
        session_id: str,
        *,
        replace_page_id: str | None = None,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        session = self.get_signed_scan_session(session_id)
        if not session or session["status"] not in {
            "collecting",
            "technical_error",
        }:
            raise WorkflowValidationError("Сессия сканирования не найдена или завершена")
        replaced_page = None
        if replace_page_id:
            replaced_page = self.db.fetch_one(
                """
                SELECT * FROM signed_scan_session_pages
                WHERE id = ? AND session_id = ? AND status = 'active'
                """,
                (replace_page_id, session_id),
            )
            if not replaced_page:
                raise WorkflowValidationError("Лист для пересканирования не найден")

        page_id = uuid4().hex
        destination = self._scan_session_page_destination(session, page_id)
        destination.parent.mkdir(parents=True, exist_ok=True)
        try:
            scanner_settings = self.get_scanner_settings()
            self.scanner.acquire_a4(
                destination,
                dpi=int(scanner_settings["dpi"]),
                color_mode=str(scanner_settings["color_mode"]),
            )
            size_bytes = destination.stat().st_size
            if not size_bytes:
                raise WorkflowValidationError("Сканер вернул пустой лист")
            if size_bytes > self.settings.max_upload_bytes:
                raise WorkflowValidationError(
                    "Один лист превышает допустимый размер файла"
                )
            with Image.open(destination) as image:
                image.verify()
            digest = self._file_sha256(destination)
            if replaced_page:
                page_order = int(replaced_page["page_order"])
            else:
                page_order = max(
                    (
                        int(page["page_order"])
                        for page in session["pages"]
                    ),
                    default=0,
                ) + 1
            now = utc_now()
            with self.db.connect() as connection:
                if replaced_page:
                    connection.execute(
                        """
                        UPDATE signed_scan_session_pages
                        SET status = 'removed', updated_at = ?
                        WHERE id = ? AND session_id = ? AND status = 'active'
                        """,
                        (now, replace_page_id, session_id),
                    )
                connection.execute(
                    """
                    INSERT INTO signed_scan_session_pages(
                        id, session_id, status, page_order, original_path,
                        sha256, size_bytes, created_at, updated_at
                    ) VALUES (?, ?, 'active', ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        page_id,
                        session_id,
                        page_order,
                        str(destination),
                        digest,
                        size_bytes,
                        now,
                        now,
                    ),
                )
                page_count = connection.execute(
                    """
                    SELECT COUNT(*) AS count
                    FROM signed_scan_session_pages
                    WHERE session_id = ? AND status = 'active'
                    """,
                    (session_id,),
                ).fetchone()["count"]
                connection.execute(
                    """
                    UPDATE signed_scan_sessions
                    SET status = 'collecting', page_count = ?,
                        error_message = NULL, updated_at = ?
                    WHERE id = ?
                    """,
                    (page_count, now, session_id),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('signed_scan_session', ?, ?, ?, ?, ?)
                    """,
                    (
                        session_id,
                        (
                            "signed_scan_page_replaced"
                            if replaced_page
                            else "signed_scan_page_added"
                        ),
                        actor,
                        json.dumps(
                            {
                                "page_id": page_id,
                                "page_order": page_order,
                                "replaced_page_id": replace_page_id,
                            },
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
        except ScannerCancelled:
            destination.unlink(missing_ok=True)
            self.db.audit(
                "signed_scan_session",
                session_id,
                "signed_scan_page_cancelled",
                actor=actor,
            )
            raise
        except (ScannerError, WorkflowValidationError, OSError) as exc:
            destination.unlink(missing_ok=True)
            self._set_signed_scan_session_error(
                session_id,
                str(exc),
                actor=actor,
            )
            raise
        return self.get_signed_scan_session(session_id) or {}

    def remove_signed_scan_session_page(
        self,
        session_id: str,
        page_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        session = self.get_signed_scan_session(session_id)
        if not session or session["status"] not in {
            "collecting",
            "technical_error",
        }:
            raise WorkflowValidationError("Сессия сканирования не найдена или завершена")
        page = next(
            (item for item in session["pages"] if item["id"] == page_id),
            None,
        )
        if not page:
            raise WorkflowValidationError("Лист не найден")
        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                UPDATE signed_scan_session_pages
                SET status = 'removed', updated_at = ?
                WHERE id = ? AND session_id = ? AND status = 'active'
                """,
                (now, page_id, session_id),
            )
            remaining = connection.execute(
                """
                SELECT id FROM signed_scan_session_pages
                WHERE session_id = ? AND status = 'active'
                ORDER BY page_order, created_at, id
                """,
                (session_id,),
            ).fetchall()
            for order, row in enumerate(remaining, 1):
                connection.execute(
                    """
                    UPDATE signed_scan_session_pages
                    SET page_order = ?, updated_at = ? WHERE id = ?
                    """,
                    (order, now, row["id"]),
                )
            connection.execute(
                """
                UPDATE signed_scan_sessions
                SET page_count = ?, updated_at = ? WHERE id = ?
                """,
                (len(remaining), now, session_id),
            )
        self.db.audit(
            "signed_scan_session",
            session_id,
            "signed_scan_page_removed",
            {"page_id": page_id},
            actor=actor,
        )
        return self.get_signed_scan_session(session_id) or {}

    def move_signed_scan_session_page(
        self,
        session_id: str,
        page_id: str,
        direction: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        if direction not in {"up", "down"}:
            raise WorkflowValidationError("Неизвестное направление перемещения")
        session = self.get_signed_scan_session(session_id)
        if not session or session["status"] not in {
            "collecting",
            "technical_error",
        }:
            raise WorkflowValidationError("Сессия сканирования не найдена или завершена")
        pages = session["pages"]
        index = next(
            (position for position, page in enumerate(pages) if page["id"] == page_id),
            None,
        )
        if index is None:
            raise WorkflowValidationError("Лист не найден")
        target_index = index - 1 if direction == "up" else index + 1
        if target_index < 0 or target_index >= len(pages):
            return session
        current = pages[index]
        target = pages[target_index]
        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                "UPDATE signed_scan_session_pages SET page_order = ?, updated_at = ? WHERE id = ?",
                (target["page_order"], now, current["id"]),
            )
            connection.execute(
                "UPDATE signed_scan_session_pages SET page_order = ?, updated_at = ? WHERE id = ?",
                (current["page_order"], now, target["id"]),
            )
            connection.execute(
                "UPDATE signed_scan_sessions SET updated_at = ? WHERE id = ?",
                (now, session_id),
            )
        self.db.audit(
            "signed_scan_session",
            session_id,
            "signed_scan_pages_reordered",
            {"page_id": page_id, "direction": direction},
            actor=actor,
        )
        return self.get_signed_scan_session(session_id) or {}

    @staticmethod
    def _make_session_pages_pdf(
        pages: list[dict[str, Any]],
        destination: Path,
    ) -> int:
        rendered: list[Image.Image] = []
        try:
            for page in pages:
                with Image.open(Path(page["original_path"])) as image:
                    image.load()
                    if image.mode in {"RGBA", "LA"}:
                        background = Image.new("RGB", image.size, "white")
                        background.paste(image, mask=image.getchannel("A"))
                        rendered.append(background)
                    else:
                        rendered.append(image.convert("RGB"))
            if not rendered:
                raise WorkflowValidationError("Добавьте хотя бы один лист")
            first, *remaining = rendered
            first.save(
                destination,
                "PDF",
                save_all=True,
                append_images=remaining,
                resolution=300.0,
            )
            return len(rendered)
        except WorkflowValidationError:
            raise
        except (UnidentifiedImageError, OSError) as exc:
            raise WorkflowValidationError(
                "Не удалось собрать PDF из отсканированных листов"
            ) from exc
        finally:
            for image in rendered:
                image.close()

    def finalize_signed_scan_session(
        self,
        session_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        session = self.get_signed_scan_session(session_id)
        if not session or session["status"] not in {
            "collecting",
            "technical_error",
        }:
            raise WorkflowValidationError("Сессия сканирования не найдена или завершена")
        pages = session["pages"]
        if not pages:
            raise WorkflowValidationError("Добавьте хотя бы один лист")
        total_size = sum(int(page["size_bytes"]) for page in pages)
        if total_size > self.settings.max_upload_bytes:
            raise WorkflowValidationError(
                "Общий размер отсканированных листов превышает допустимый"
            )

        scan_id = uuid4().hex
        scan_dir = ensure_within(
            self.settings.runtime_dir
            / "signed_scans"
            / session["response_letter_id"]
            / scan_id,
            self.settings.runtime_dir,
        )
        scan_dir.mkdir(parents=True, exist_ok=False)
        letter = self.get_response_letter(session["response_letter_id"])
        if not letter:
            raise WorkflowValidationError("Готовое письмо не найдено")
        pdf_path = scan_dir / self._signed_scan_filename(letter)
        manifest_path = scan_dir / "source-pages.json"
        try:
            page_count = self._make_session_pages_pdf(pages, pdf_path)
            if pdf_path.stat().st_size > self.settings.max_upload_bytes:
                raise WorkflowValidationError(
                    "Созданный PDF превышает допустимый размер файла"
                )
            manifest_path.write_text(
                json.dumps(
                    {
                        "session_id": session_id,
                        "pages": [
                            {
                                "id": page["id"],
                                "order": page["page_order"],
                                "sha256": page["sha256"],
                            }
                            for page in pages
                        ],
                    },
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
            digest = self._file_sha256(pdf_path)
            now = utc_now()
            with self.db.connect() as connection:
                connection.execute(
                    """
                    UPDATE signed_response_scans
                    SET status = 'superseded', updated_at = ?
                    WHERE response_letter_id = ?
                      AND status IN ('needs_confirmation', 'confirmed')
                    """,
                    (now, session["response_letter_id"]),
                )
                connection.execute(
                    """
                    INSERT INTO signed_response_scans(
                        id, response_letter_id, status, source,
                        original_filename, original_path, pdf_path,
                        sha256, size_bytes, page_count,
                        correct_letter_confirmed, signature_confirmed,
                        bank_seal_confirmed, confirmed_by, confirmed_at,
                        created_at, updated_at
                    ) VALUES (?, ?, 'confirmed', 'wia', ?, ?, ?, ?, ?, ?,
                              1, 1, 1, ?, ?, ?, ?)
                    """,
                    (
                        scan_id,
                        session["response_letter_id"],
                        self._signed_scan_filename(letter),
                        str(manifest_path),
                        str(pdf_path),
                        digest,
                        pdf_path.stat().st_size,
                        page_count,
                        actor,
                        now,
                        now,
                        now,
                    ),
                )
                connection.execute(
                    """
                    UPDATE signed_scan_sessions
                    SET status = 'completed', page_count = ?,
                        result_scan_id = ?, error_message = NULL, updated_at = ?
                    WHERE id = ?
                    """,
                    (page_count, scan_id, now, session_id),
                )
                for entity_type, entity_id, event_type, payload in (
                    (
                        "signed_response_scan",
                        scan_id,
                        "signed_response_scan_registered",
                        {
                            "response_letter_id": session["response_letter_id"],
                            "source": "wia",
                            "sha256": digest,
                            "page_count": page_count,
                            "session_id": session_id,
                            "confirmed_during_scan": True,
                        },
                    ),
                    (
                        "signed_response_scan",
                        scan_id,
                        "signed_response_scan_confirmed",
                        {"confirmed_during_scan": True},
                    ),
                    (
                        "signed_scan_session",
                        session_id,
                        "signed_scan_session_completed",
                        {"result_scan_id": scan_id, "page_count": page_count},
                    ),
                ):
                    connection.execute(
                        """
                        INSERT INTO audit_events(
                            entity_type, entity_id, event_type,
                            actor, payload_json, created_at
                        ) VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            entity_type,
                            entity_id,
                            event_type,
                            actor,
                            json.dumps(payload, ensure_ascii=False),
                            now,
                        ),
                    )
        except Exception as exc:
            shutil.rmtree(scan_dir, ignore_errors=True)
            safe_message = (
                str(exc)
                if isinstance(exc, (WorkflowValidationError, OSError))
                else "Не удалось завершить сессию сканирования"
            )
            self._set_signed_scan_session_error(
                session_id,
                safe_message,
                actor=actor,
            )
            raise
        return self.get_signed_response_scan(scan_id) or {}

    def cancel_signed_scan_session(
        self,
        session_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> None:
        session = self.get_signed_scan_session(session_id)
        if not session or session["status"] not in {
            "collecting",
            "technical_error",
        }:
            raise WorkflowValidationError("Сессия сканирования не найдена или завершена")
        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                UPDATE signed_scan_sessions
                SET status = 'cancelled', error_message = NULL, updated_at = ?
                WHERE id = ?
                """,
                (now, session_id),
            )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('signed_scan_session', ?,
                          'signed_scan_session_cancelled', ?, ?, ?)
                """,
                (
                    session_id,
                    actor,
                    json.dumps(
                        {"page_count": len(session["pages"])},
                        ensure_ascii=False,
                    ),
                    now,
                ),
            )

    def register_signed_response_scan(
        self,
        letter_id: str,
        original_filename: str,
        stream: BinaryIO,
        *,
        source: str = "upload",
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        letter = self.get_response_letter(letter_id)
        if not letter:
            raise WorkflowValidationError("Готовое письмо не найдено")
        if source not in {"upload", "wia"}:
            raise WorkflowValidationError("Неизвестный источник скана")
        safe_name = sanitize_filename(original_filename or "scan.pdf")
        suffix = Path(safe_name).suffix.casefold()
        if suffix not in {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
            raise WorkflowValidationError(
                "Разрешены PDF, PNG, JPEG и TIFF"
            )

        scan_id = uuid4().hex
        scan_dir = ensure_within(
            self.settings.runtime_dir / "signed_scans" / letter_id / scan_id,
            self.settings.runtime_dir,
        )
        scan_dir.mkdir(parents=True, exist_ok=False)
        original = scan_dir / f"original{suffix}"
        temporary = original.with_suffix(original.suffix + ".part")
        digest = hashlib.sha256()
        size_bytes = 0
        try:
            with temporary.open("wb") as output:
                while chunk := stream.read(1024 * 1024):
                    size_bytes += len(chunk)
                    if size_bytes > self.settings.max_upload_bytes:
                        raise WorkflowValidationError(
                            "Скан превышает допустимый размер файла"
                        )
                    digest.update(chunk)
                    output.write(chunk)
            if not size_bytes:
                raise WorkflowValidationError("Получен пустой файл скана")
            temporary.replace(original)
            pdf_path = scan_dir / self._signed_scan_filename(letter)
            page_count = self._make_scan_pdf(original, pdf_path)
            now = utc_now()
            with self.db.connect() as connection:
                connection.execute(
                    """
                    UPDATE signed_response_scans
                    SET status = 'superseded', updated_at = ?
                    WHERE response_letter_id = ?
                      AND status IN ('needs_confirmation', 'confirmed')
                    """,
                    (now, letter_id),
                )
                connection.execute(
                    """
                    INSERT INTO signed_response_scans(
                        id, response_letter_id, status, source,
                        original_filename, original_path, pdf_path,
                        sha256, size_bytes, page_count,
                        correct_letter_confirmed, signature_confirmed,
                        bank_seal_confirmed, confirmed_by, confirmed_at,
                        created_at, updated_at
                    ) VALUES (?, ?, 'confirmed', ?, ?, ?, ?, ?, ?, ?,
                              1, 1, 1, ?, ?, ?, ?)
                    """,
                    (
                        scan_id,
                        letter_id,
                        source,
                        safe_name,
                        str(original),
                        str(pdf_path),
                        digest.hexdigest(),
                        size_bytes,
                        page_count,
                        actor,
                        now,
                        now,
                        now,
                    ),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('signed_response_scan', ?,
                              'signed_response_scan_registered', ?, ?, ?)
                    """,
                    (
                        scan_id,
                        actor,
                        json.dumps(
                            {
                                "response_letter_id": letter_id,
                                "source": source,
                                "sha256": digest.hexdigest(),
                                "page_count": page_count,
                                "confirmed_during_scan": True,
                            },
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('signed_response_scan', ?,
                              'signed_response_scan_confirmed', ?, ?, ?)
                    """,
                    (
                        scan_id,
                        actor,
                        json.dumps(
                            {"confirmed_during_scan": True},
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
                connection.execute(
                    """
                    UPDATE signed_scan_sessions
                    SET status = 'cancelled', error_message = NULL,
                        updated_at = ?
                    WHERE response_letter_id = ?
                      AND status IN ('collecting', 'technical_error')
                    """,
                    (now, letter_id),
                )
        except Exception:
            shutil.rmtree(scan_dir, ignore_errors=True)
            raise
        return self.get_signed_response_scan(scan_id) or {}

    @staticmethod
    def _make_scan_pdf(source: Path, destination: Path) -> int:
        if source.suffix.casefold() == ".pdf":
            try:
                reader = PdfReader(str(source))
                page_count = len(reader.pages)
                if page_count < 1:
                    raise WorkflowValidationError(
                        "PDF подписанного ответа не содержит страниц"
                    )
            except WorkflowValidationError:
                raise
            except Exception as exc:
                raise WorkflowValidationError("Не удалось открыть PDF скана") from exc
            shutil.copy2(source, destination)
            return page_count
        try:
            with Image.open(source) as image:
                pages: list[Image.Image] = []
                for frame in ImageSequence.Iterator(image):
                    frame.load()
                    if frame.mode in {"RGBA", "LA"}:
                        background = Image.new("RGB", frame.size, "white")
                        background.paste(frame, mask=frame.getchannel("A"))
                        pages.append(background)
                    else:
                        pages.append(frame.convert("RGB"))
                if not pages:
                    raise WorkflowValidationError(
                        "Изображение подписанного ответа не содержит страниц"
                    )
                first, *remaining = pages
                first.save(
                    destination,
                    "PDF",
                    save_all=True,
                    append_images=remaining,
                    resolution=300.0,
                )
                return len(pages)
        except WorkflowValidationError:
            raise
        except (UnidentifiedImageError, OSError) as exc:
            raise WorkflowValidationError("Не удалось открыть изображение скана") from exc

    def acquire_signed_response_scan(
        self,
        letter_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        session = self.start_signed_scan_session(letter_id, actor=actor)
        session = self.acquire_signed_scan_session_page(
            session["id"],
            actor=actor,
        )
        return self.finalize_signed_scan_session(
            session["id"],
            actor=actor,
        )

    def confirm_signed_response_scan(
        self,
        scan_id: str,
        *,
        correct_letter: bool,
        signature_present: bool,
        bank_seal_present: bool,
        actor: str = "Сотрудник",
    ) -> None:
        scan = self.get_signed_response_scan(scan_id)
        if not scan or scan["status"] != "needs_confirmation":
            raise WorkflowValidationError("Скан не найден или уже проверен")
        if not (correct_letter and signature_present and bank_seal_present):
            raise WorkflowValidationError(
                "Подтвердите правильное письмо, подписи и печать банка"
            )
        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                UPDATE signed_response_scans
                SET status = 'confirmed',
                    correct_letter_confirmed = 1,
                    signature_confirmed = 1,
                    bank_seal_confirmed = 1,
                    confirmed_by = ?, confirmed_at = ?, updated_at = ?
                WHERE id = ? AND status = 'needs_confirmation'
                """,
                (actor, now, now, scan_id),
            )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('signed_response_scan', ?,
                          'signed_response_scan_confirmed', ?, ?, ?)
                """,
                (
                    scan_id,
                    actor,
                    json.dumps(
                        {"response_letter_id": scan["response_letter_id"]},
                        ensure_ascii=False,
                    ),
                    now,
                ),
            )

    def repair_cleaned_responses(self) -> None:
        cases = self.db.fetch_all(
            """
            SELECT cases.*
            FROM cases
            WHERE cases.response_path IS NOT NULL
              AND EXISTS (
                  SELECT 1
                  FROM audit_events AS cleanup
                  WHERE cleanup.entity_type = 'case'
                    AND cleanup.entity_id = cases.id
                    AND cleanup.event_type =
                        'district_place_edge_noise_removed'
              )
              AND NOT EXISTS (
                  SELECT 1
                  FROM audit_events AS repaired
                  WHERE repaired.entity_type = 'case'
                    AND repaired.entity_id = cases.id
                    AND repaired.event_type =
                        'response_regenerated_after_cleanup'
              )
            """
        )
        for case in cases:
            path = Path(case["response_path"])
            try:
                _, likely_overflow = self.word.render(
                    path,
                    case,
                    self.get_taxpayers(case["id"]),
                )
            except (OSError, ValueError):
                continue
            self.db.execute(
                "UPDATE cases SET response_page_overflow = ? WHERE id = ?",
                (int(likely_overflow), case["id"]),
            )
            self.db.audit(
                "case",
                case["id"],
                "response_regenerated_after_cleanup",
                {"filename": path.name},
            )

    @classmethod
    def _business_day_bounds(
        cls, business_date: date
    ) -> tuple[str, str]:
        local_start = datetime.combine(
            business_date,
            time.min,
            tzinfo=cls.BUSINESS_TIMEZONE,
        )
        local_end = datetime.combine(
            business_date + timedelta(days=1),
            time.min,
            tzinfo=cls.BUSINESS_TIMEZONE,
        )
        return (
            local_start.astimezone(UTC).isoformat(timespec="seconds"),
            local_end.astimezone(UTC).isoformat(timespec="seconds"),
        )

    @staticmethod
    def _normalize_group_value(value: str | None) -> str:
        return " ".join((value or "").split()).casefold()

    @classmethod
    def _response_group_key(
        cls,
        business_date: date,
        abs_bucket: AbsStatus,
        district_place: str,
        recipient_full_name: str,
    ) -> str:
        identity = "\x1f".join(
            (
                business_date.isoformat(),
                str(abs_bucket),
                cls._normalize_group_value(district_place),
                cls._normalize_group_value(recipient_full_name),
            )
        )
        return hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24]

    def generate_response(self, case_id: str) -> Path:
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if case["status"] != CaseStatus.READY_FOR_RESPONSE:
            raise WorkflowValidationError(
                "Обращение ещё не готово к созданию ответа"
            )
        taxpayers = self.get_taxpayers(case_id)
        output = self.settings.responses_dir / self._response_word_filename(
            case.get("recipient_display_name"),
            datetime.now(self.BUSINESS_TIMEZONE).date().isoformat(),
            case_id,
        )
        _, likely_overflow = self.word.render(output, case, taxpayers)
        self.db.execute(
            """
            UPDATE cases
            SET status = ?, response_status = ?, response_path = ?,
                response_page_overflow = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                CaseStatus.RESPONSE_CREATED,
                "created",
                str(output),
                int(likely_overflow),
                utc_now(),
                case_id,
            ),
        )
        self.db.audit(
            "case",
            case_id,
            "response_created",
            {
                "filename": output.name,
                "taxpayer_count": len(taxpayers),
                "page_overflow": likely_overflow,
            },
        )
        return output

    def _refresh_upload_status(self, upload_id: str) -> None:
        page_counts = self.db.fetch_one(
            """
            SELECT
                COUNT(*) AS total,
                SUM(CASE WHEN status IN ('needs_review', 'technical_error')
                    THEN 1 ELSE 0 END) AS problem,
                SUM(CASE WHEN status IN ('registered', 'preview_ready', 'processing')
                    THEN 1 ELSE 0 END) AS pending
            FROM pages WHERE upload_id = ?
            """,
            (upload_id,),
        ) or {"total": 0, "problem": 0, "pending": 0}
        incomplete_cases = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM cases
            WHERE upload_id = ? AND status IN ('collecting', 'needs_review')
            """,
            (upload_id,),
        ) or {"count": 0}

        if page_counts["pending"]:
            status = UploadStatus.PROCESSING
        elif page_counts["problem"] or incomplete_cases["count"]:
            status = UploadStatus.NEEDS_REVIEW
        else:
            status = UploadStatus.READY
        self.db.execute(
            "UPDATE uploads SET status = ? WHERE id = ?",
            (status, upload_id),
        )

    def get_upload(self, upload_id: str) -> dict[str, Any] | None:
        return self.db.fetch_one("SELECT * FROM uploads WHERE id = ?", (upload_id,))

    def list_uploads(
        self,
        limit: int | None = None,
    ) -> list[dict[str, Any]]:
        if limit is None:
            return self.db.fetch_all(
                "SELECT * FROM uploads ORDER BY created_at DESC"
            )
        safe_limit = max(1, min(int(limit), 200))
        return self.db.fetch_all(
            "SELECT * FROM uploads ORDER BY created_at DESC LIMIT ?",
            (safe_limit,),
        )

    def list_incoming_work(
        self,
        limit: int | None = None,
        *,
        sort_order: str = "newest",
        offset: int = 0,
    ) -> list[dict[str, Any]]:
        direction = "ASC" if sort_order == "oldest" else "DESC"
        limit_clause = ""
        parameters: tuple[Any, ...] = ()
        if limit is not None:
            safe_limit = max(1, min(int(limit), 100))
            safe_offset = max(0, int(offset))
            limit_clause = "LIMIT ? OFFSET ?"
            parameters = (safe_limit, safe_offset)
        return self.db.fetch_all(
            f"""
            SELECT uploads.*,
                   CASE WHEN uploads.intake_source = 'outlook'
                        THEN MIN(outlook_messages.received_at)
                   END AS received_at,
                   CASE WHEN uploads.intake_source = 'outlook'
                        THEN GROUP_CONCAT(
                            DISTINCT outlook_messages.sender_smtp
                        )
                   END AS sender_smtp,
                   CASE WHEN uploads.intake_source = 'outlook'
                        THEN GROUP_CONCAT(
                            DISTINCT outlook_messages.original_sender_smtp
                        )
                   END AS original_sender_smtp,
                   CASE WHEN uploads.intake_source = 'outlook'
                        THEN COUNT(DISTINCT outlook_messages.source_key)
                        ELSE 0
                   END AS source_count
            FROM uploads
            LEFT JOIN outlook_attachments
              ON outlook_attachments.upload_id = uploads.id
            LEFT JOIN outlook_messages
              ON outlook_messages.source_key =
                 outlook_attachments.message_key
            GROUP BY uploads.id
            ORDER BY COALESCE(
                MIN(outlook_messages.received_at), uploads.created_at
            ) {direction}, uploads.created_at {direction}, uploads.id {direction}
            {limit_clause}
            """,
            parameters,
        )

    def count_incoming_work(self) -> int:
        row = self.db.fetch_one("SELECT COUNT(*) AS count FROM uploads")
        return int(row.get("count") or 0) if row else 0

    def get_upload_pages(self, upload_id: str) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            "SELECT * FROM pages WHERE upload_id = ? ORDER BY page_number",
            (upload_id,),
        )

    def get_page(self, page_id: str) -> dict[str, Any] | None:
        return self.db.fetch_one("SELECT * FROM pages WHERE id = ?", (page_id,))

    def get_page_pdf_path(self, page_id: str) -> Path:
        page = self.get_page(page_id)
        if not page:
            raise WorkflowValidationError("Страница не найдена")
        upload = self.get_upload(page["upload_id"])
        if not upload:
            raise WorkflowValidationError("Исходный PDF не найден")
        try:
            source = ensure_within(
                Path(upload["stored_path"]),
                self.settings.uploads_dir,
            )
            output = ensure_within(
                self.settings.previews_dir
                / upload["id"]
                / f"page-{int(page['page_number']):04d}.pdf",
                self.settings.previews_dir,
            )
            if not output.exists() or (
                source.stat().st_mtime_ns > output.stat().st_mtime_ns
            ):
                self.pdf.extract_page_pdf(
                    source,
                    int(page["page_number"]),
                    output,
                )
            return output
        except (OSError, PdfProcessingError, ValueError) as exc:
            raise WorkflowValidationError(str(exc)) from exc

    def list_review_pages(self) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            """
            SELECT pages.*, uploads.original_filename
            FROM pages
            JOIN uploads ON uploads.id = pages.upload_id
            WHERE pages.status IN ('needs_review', 'technical_error')
            ORDER BY uploads.created_at, pages.page_number
            """
        )

    def _case_match_common_key(self, case: dict[str, Any]) -> tuple[str, ...]:
        district = clean_location(case.get("district_place") or "")
        office = self.canonical_gns_office(district)
        return (
            str(office["office_key"] if office else district).casefold(),
            self._normalize_group_value(case.get("recipient_position")),
            self._normalize_group_value(case.get("recipient_full_name")),
            str(case.get("period_start") or ""),
            str(case.get("period_end") or ""),
            str(case.get("period_route") or ""),
            self._normalize_group_value(case.get("employee_name")),
        )

    def _case_match_name_key(self, taxpayer: dict[str, Any]) -> str:
        """Normalize only the presentation-level ИП prefix for comparison."""
        return self._normalize_group_value(
            response_taxpayer_name(
                taxpayer.get("name") or "",
                taxpayer.get("inn") or "",
            )
        )

    @staticmethod
    def _case_match_signature(
        left_case: dict[str, Any],
        right_case: dict[str, Any],
        left_taxpayer: dict[str, Any],
        right_taxpayer: dict[str, Any],
    ) -> str:
        payload = {
            "left": {
                "case_id": left_case["id"],
                "case_source": left_case.get("source_kind") or "",
                "official_parse_version": int(
                    left_case.get("official_parse_version") or 0
                ),
                "taxpayer_id": left_taxpayer["id"],
                "name": clean_taxpayer_name(left_taxpayer.get("name") or ""),
                "inn": re.sub(r"\D", "", left_taxpayer.get("inn") or ""),
                "name_source": left_taxpayer.get("name_source") or "",
                "inn_source": left_taxpayer.get("inn_source") or "",
            },
            "right": {
                "case_id": right_case["id"],
                "case_source": right_case.get("source_kind") or "",
                "official_parse_version": int(
                    right_case.get("official_parse_version") or 0
                ),
                "taxpayer_id": right_taxpayer["id"],
                "name": clean_taxpayer_name(right_taxpayer.get("name") or ""),
                "inn": re.sub(r"\D", "", right_taxpayer.get("inn") or ""),
                "name_source": right_taxpayer.get("name_source") or "",
                "inn_source": right_taxpayer.get("inn_source") or "",
            },
        }
        return hashlib.sha256(
            json.dumps(
                payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()

    def _case_match_candidates(self) -> list[dict[str, Any]]:
        cases = self.db.fetch_all(
            """
            SELECT cases.*
            FROM cases
            WHERE cases.fields_confirmed = 1
              AND cases.status NOT IN (
                  'response_created', 'completed', 'technical_error'
              )
              AND EXISTS (
                  SELECT 1 FROM taxpayers
                  WHERE taxpayers.case_id = cases.id
              )
            ORDER BY cases.created_at, cases.id
            """
        )
        if len(cases) < 2:
            return []
        taxpayers = self.db.fetch_all(
            """
            SELECT taxpayers.*
            FROM taxpayers
            JOIN cases ON cases.id = taxpayers.case_id
            WHERE cases.fields_confirmed = 1
              AND cases.status NOT IN (
                  'response_created', 'completed', 'technical_error'
              )
            ORDER BY taxpayers.case_id, taxpayers.display_order
            """
        )
        taxpayers_by_case: dict[str, list[dict[str, Any]]] = {}
        for taxpayer in taxpayers:
            taxpayers_by_case.setdefault(taxpayer["case_id"], []).append(
                taxpayer
            )
        groups: dict[tuple[str, ...], list[dict[str, Any]]] = {}
        for case in cases:
            groups.setdefault(self._case_match_common_key(case), []).append(case)

        candidates: list[dict[str, Any]] = []
        for matching_cases in groups.values():
            for left_index, first_case in enumerate(matching_cases):
                for second_case in matching_cases[left_index + 1 :]:
                    left_case, right_case = sorted(
                        (first_case, second_case), key=lambda item: item["id"]
                    )
                    left_taxpayers = taxpayers_by_case.get(left_case["id"], [])
                    right_taxpayers = taxpayers_by_case.get(right_case["id"], [])
                    if (
                        not left_taxpayers
                        or len(left_taxpayers) != len(right_taxpayers)
                    ):
                        continue
                    differences: list[
                        tuple[dict[str, Any], dict[str, Any], str]
                    ] = []
                    invalid_pair = False
                    for left_taxpayer, right_taxpayer in zip(
                        left_taxpayers, right_taxpayers, strict=True
                    ):
                        names_equal = self._case_match_name_key(
                            left_taxpayer
                        ) == self._case_match_name_key(right_taxpayer)
                        inns_equal = re.sub(
                            r"\D", "", left_taxpayer.get("inn") or ""
                        ) == re.sub(
                            r"\D", "", right_taxpayer.get("inn") or ""
                        )
                        if names_equal and inns_equal:
                            continue
                        if names_equal == inns_equal:
                            invalid_pair = True
                            break
                        differences.append(
                            (
                                left_taxpayer,
                                right_taxpayer,
                                "inn" if names_equal else "name",
                            )
                        )
                    if invalid_pair or len(differences) != 1:
                        continue
                    left_taxpayer, right_taxpayer, differing_field = (
                        differences[0]
                    )
                    candidates.append(
                        {
                            "left_case_id": left_case["id"],
                            "right_case_id": right_case["id"],
                            "left_taxpayer_id": left_taxpayer["id"],
                            "right_taxpayer_id": right_taxpayer["id"],
                            "differing_field": differing_field,
                            "signature_hash": self._case_match_signature(
                                left_case,
                                right_case,
                                left_taxpayer,
                                right_taxpayer,
                            ),
                        }
                    )
        return candidates

    def reconcile_case_match_reviews(self) -> int:
        candidates = self._case_match_candidates()
        candidate_keys = {
            (
                item["left_case_id"],
                item["right_case_id"],
                item["signature_hash"],
            )
            for item in candidates
        }
        now = utc_now()
        affected_case_ids: set[str] = set()
        created_ids: list[str] = []
        with self.db.connect() as connection:
            pending_rows = connection.execute(
                "SELECT * FROM case_match_reviews WHERE status = 'pending'"
            ).fetchall()
            for pending in pending_rows:
                key = (
                    pending["left_case_id"],
                    pending["right_case_id"],
                    pending["signature_hash"],
                )
                if key in candidate_keys:
                    continue
                connection.execute(
                    """
                    UPDATE case_match_reviews
                    SET status = 'superseded', updated_at = ?
                    WHERE id = ?
                    """,
                    (now, pending["id"]),
                )
                affected_case_ids.update(
                    (pending["left_case_id"], pending["right_case_id"])
                )

            for candidate in candidates:
                existing = connection.execute(
                    """
                    SELECT id, status
                    FROM case_match_reviews
                    WHERE left_case_id = ? AND right_case_id = ?
                      AND signature_hash = ?
                    LIMIT 1
                    """,
                    (
                        candidate["left_case_id"],
                        candidate["right_case_id"],
                        candidate["signature_hash"],
                    ),
                ).fetchone()
                if existing and existing["status"] != "pending":
                    continue
                if not existing:
                    review_id = uuid4().hex
                    connection.execute(
                        """
                        INSERT INTO case_match_reviews(
                            id, left_case_id, right_case_id,
                            left_taxpayer_id, right_taxpayer_id,
                            signature_hash, differing_field, status,
                            created_at, updated_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
                        """,
                        (
                            review_id,
                            candidate["left_case_id"],
                            candidate["right_case_id"],
                            candidate["left_taxpayer_id"],
                            candidate["right_taxpayer_id"],
                            candidate["signature_hash"],
                            candidate["differing_field"],
                            now,
                            now,
                        ),
                    )
                    created_ids.append(review_id)
                for case_id in (
                    candidate["left_case_id"],
                    candidate["right_case_id"],
                ):
                    connection.execute(
                        """
                        UPDATE cases
                        SET match_review_previous_status = CASE
                                WHEN status != 'needs_review' THEN status
                                ELSE COALESCE(
                                    match_review_previous_status, status
                                )
                            END,
                            status = 'needs_review', updated_at = ?
                        WHERE id = ?
                        """,
                        (now, case_id),
                    )
                    affected_case_ids.add(case_id)

            for case_id in affected_case_ids:
                pending_count = connection.execute(
                    """
                    SELECT COUNT(*) AS count
                    FROM case_match_reviews
                    WHERE status = 'pending'
                      AND (left_case_id = ? OR right_case_id = ?)
                    """,
                    (case_id, case_id),
                ).fetchone()["count"]
                if pending_count:
                    continue
                case = connection.execute(
                    """
                    SELECT status, match_review_previous_status
                    FROM cases WHERE id = ?
                    """,
                    (case_id,),
                ).fetchone()
                if not case or not case["match_review_previous_status"]:
                    continue
                if case["status"] == CaseStatus.NEEDS_REVIEW:
                    connection.execute(
                        """
                        UPDATE cases
                        SET status = ?, match_review_previous_status = NULL,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (case["match_review_previous_status"], now, case_id),
                    )
                else:
                    connection.execute(
                        """
                        UPDATE cases
                        SET match_review_previous_status = NULL,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (now, case_id),
                    )
            for review_id in created_ids:
                connection.execute(
                    """
                    INSERT INTO audit_events(
                        entity_type, entity_id, event_type,
                        actor, payload_json, created_at
                    ) VALUES ('case_match_review', ?,
                              'case_match_review_created', 'system', ?, ?)
                    """,
                    (
                        review_id,
                        json.dumps(
                            {"automatic_decision": False},
                            ensure_ascii=False,
                        ),
                        now,
                    ),
                )
        return len(created_ids)

    def has_pending_case_match_review(self, case_id: str) -> bool:
        row = self.db.fetch_one(
            """
            SELECT 1 AS found
            FROM case_match_reviews
            WHERE status = 'pending'
              AND (left_case_id = ? OR right_case_id = ?)
            LIMIT 1
            """,
            (case_id, case_id),
        )
        return bool(row)

    def list_case_match_reviews(self) -> list[dict[str, Any]]:
        self.reconcile_case_match_reviews()
        reviews = self.db.fetch_all(
            """
            SELECT case_match_reviews.*
            FROM case_match_reviews
            WHERE case_match_reviews.status = 'pending'
            ORDER BY case_match_reviews.created_at,
                     case_match_reviews.id
            """
        )
        return [
            self.get_case_match_review(item["id"]) or item
            for item in reviews
        ]

    def _case_match_source(
        self, case_id: str, taxpayer_id: str
    ) -> dict[str, Any] | None:
        case = self.db.fetch_one(
            """
            SELECT cases.*, uploads.original_filename,
                (
                    SELECT pages.id FROM pages
                    WHERE pages.case_id = cases.id
                    ORDER BY CASE WHEN pages.page_type = 'letter'
                                  THEN 0 ELSE 1 END,
                             pages.page_number
                    LIMIT 1
                ) AS source_page_id,
                (
                    SELECT pages.page_number FROM pages
                    WHERE pages.case_id = cases.id
                    ORDER BY CASE WHEN pages.page_type = 'letter'
                                  THEN 0 ELSE 1 END,
                             pages.page_number
                    LIMIT 1
                ) AS source_page_number
            FROM cases
            JOIN uploads ON uploads.id = cases.upload_id
            WHERE cases.id = ?
            """,
            (case_id,),
        )
        if not case:
            return None
        taxpayer = self.db.fetch_one(
            "SELECT * FROM taxpayers WHERE id = ? AND case_id = ?",
            (taxpayer_id, case_id),
        )
        if not taxpayer:
            return None
        case["taxpayer"] = taxpayer
        return case

    def get_case_match_review(
        self, review_id: str
    ) -> dict[str, Any] | None:
        review = self.db.fetch_one(
            "SELECT * FROM case_match_reviews WHERE id = ?", (review_id,)
        )
        if not review:
            return None
        review["left"] = self._case_match_source(
            review["left_case_id"], review["left_taxpayer_id"]
        )
        review["right"] = self._case_match_source(
            review["right_case_id"], review["right_taxpayer_id"]
        )
        review["issue_message"] = (
            "Совпадают письмо и наименование лица, но отличается ИНН."
            if review["differing_field"] == "inn"
            else "Совпадают письмо и ИНН, но отличается наименование лица."
        )
        return review

    def _release_case_match_block(
        self,
        connection: sqlite3.Connection,
        case_id: str,
        now: str,
        target_status: str | None = None,
    ) -> None:
        pending_count = connection.execute(
            """
            SELECT COUNT(*) AS count FROM case_match_reviews
            WHERE status = 'pending'
              AND (left_case_id = ? OR right_case_id = ?)
            """,
            (case_id, case_id),
        ).fetchone()["count"]
        case = connection.execute(
            """
            SELECT status, match_review_previous_status
            FROM cases WHERE id = ?
            """,
            (case_id,),
        ).fetchone()
        if not case:
            return
        restored = target_status or case["match_review_previous_status"]
        if pending_count:
            connection.execute(
                """
                UPDATE cases
                SET status = 'needs_review',
                    match_review_previous_status = COALESCE(?,
                        match_review_previous_status),
                    updated_at = ?
                WHERE id = ?
                """,
                (restored, now, case_id),
            )
            return
        connection.execute(
            """
            UPDATE cases
            SET status = ?, match_review_previous_status = NULL,
                updated_at = ?
            WHERE id = ?
            """,
            (restored or CaseStatus.NEEDS_REVIEW, now, case_id),
        )

    def resolve_case_match_review(
        self,
        review_id: str,
        *,
        decision: str,
        name_choice: str = "",
        inn_choice: str = "",
        manual_name: str = "",
        manual_inn: str = "",
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        self.reconcile_case_match_reviews()
        review = self.get_case_match_review(review_id)
        if not review or review.get("status") != "pending":
            raise WorkflowValidationError(
                "Проверка уже завершена или исходные данные изменились"
            )
        left = review.get("left")
        right = review.get("right")
        if not left or not right:
            raise WorkflowValidationError(
                "Одна из версий изменилась. Обновите очередь проверки"
            )
        if decision not in {"same", "distinct"}:
            raise WorkflowValidationError(
                "Укажите, относится ли это к одному обращению"
            )
        now = utc_now()
        case_ids = [review["left_case_id"], review["right_case_id"]]
        recheck_case_ids: set[str] = set()
        resolution_payload: dict[str, Any] = {"decision": decision}
        with self.db.connect() as connection:
            current = connection.execute(
                "SELECT status FROM case_match_reviews WHERE id = ?",
                (review_id,),
            ).fetchone()
            if not current or current["status"] != "pending":
                raise WorkflowValidationError("Проверка уже завершена")

            if decision == "same":
                records = [left["taxpayer"], right["taxpayer"]]
                records_by_id = {item["id"]: item for item in records}

                def selected_value(
                    field: str, choice: str, manual_value: str
                ) -> tuple[str, str, str | None]:
                    if choice == "manual":
                        return manual_value, str(ValueSource.MANUAL), None
                    record = records_by_id.get(choice)
                    if not record:
                        raise WorkflowValidationError(
                            f"Выберите источник для поля «{field}»"
                        )
                    is_name = field == "Наименование"
                    value_key = "name" if is_name else "inn"
                    source_key = "name_source" if is_name else "inn_source"
                    reference_key = (
                        "name_source_reference"
                        if is_name
                        else "inn_source_reference"
                    )
                    return (
                        str(record[value_key]),
                        str(record.get(source_key) or ValueSource.MANUAL),
                        record.get(reference_key) or f"taxpayer:{record['id']}",
                    )

                selected_name, name_source, name_reference = selected_value(
                    "Наименование", name_choice, manual_name
                )
                selected_inn, inn_source, inn_reference = selected_value(
                    "ИНН", inn_choice, manual_inn
                )
                clean_name = clean_taxpayer_name(selected_name)
                clean_inn = re.sub(r"\D", "", selected_inn)
                if not clean_name:
                    raise WorkflowValidationError(
                        "Укажите итоговое наименование"
                    )
                if not clean_inn:
                    raise WorkflowValidationError(
                        "Укажите итоговый ИНН"
                    )
                record_ids = [item["id"] for item in records]
                placeholders = ",".join("?" for _ in record_ids)
                for case_id in case_ids:
                    duplicate = connection.execute(
                        f"""
                        SELECT id FROM taxpayers
                        WHERE case_id = ? AND inn = ?
                          AND id NOT IN ({placeholders})
                        LIMIT 1
                        """,
                        (case_id, clean_inn, *record_ids),
                    ).fetchone()
                    if duplicate:
                        raise WorkflowValidationError(
                            "Итоговый ИНН уже указан для другого лица "
                            "в одном из обращений"
                        )
                for record in records:
                    inn_changed = str(record["inn"]) != clean_inn
                    if inn_changed:
                        recheck_case_ids.add(str(record["case_id"]))
                    connection.execute(
                        """
                        UPDATE taxpayers
                        SET name = ?, inn = ?,
                            name_source = ?, inn_source = ?,
                            name_source_reference = ?,
                            inn_source_reference = ?,
                            manually_confirmed = 1,
                            abs_result = CASE WHEN ? THEN NULL
                                              ELSE abs_result END,
                            abs_account_result = CASE WHEN ? THEN NULL
                                                      ELSE abs_account_result END,
                            abs_active_account_count = CASE WHEN ? THEN NULL
                                                            ELSE abs_active_account_count END,
                            abs_closed_account_count = CASE WHEN ? THEN NULL
                                                            ELSE abs_closed_account_count END,
                            odb_result = CASE WHEN ? THEN NULL
                                              ELSE odb_result END,
                            registry_status = NULL, registry_name = NULL,
                            registry_director = NULL,
                            registry_checked_at = NULL,
                            registry_provider = NULL, updated_at = ?
                        WHERE id = ?
                        """,
                        (
                            clean_name,
                            clean_inn,
                            name_source,
                            inn_source,
                            name_reference,
                            inn_reference,
                            int(inn_changed),
                            int(inn_changed),
                            int(inn_changed),
                            int(inn_changed),
                            int(inn_changed),
                            now,
                            record["id"],
                        ),
                    )
                resolution_payload.update(
                    {
                        "name_choice": name_choice,
                        "inn_choice": inn_choice,
                        "requires_abs_recheck": bool(recheck_case_ids),
                    }
                )

            connection.execute(
                """
                UPDATE case_match_reviews
                SET status = 'resolved', resolution = ?,
                    resolved_by = ?, resolved_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (decision, actor, now, now, review_id),
            )
            for case_id in case_ids:
                target_status = (
                    CaseStatus.READY_FOR_ABS
                    if case_id in recheck_case_ids
                    else None
                )
                if case_id in recheck_case_ids:
                    connection.execute(
                        """
                        UPDATE cases
                        SET abs_status = ?, response_status = NULL,
                            response_path = NULL,
                            response_page_overflow = NULL
                        WHERE id = ?
                        """,
                        (AbsStatus.NOT_CHECKED, case_id),
                    )
                self._release_case_match_block(
                    connection, case_id, now, target_status
                )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('case_match_review', ?,
                          'case_match_review_resolved', ?, ?, ?)
                """,
                (
                    review_id,
                    actor,
                    json.dumps(resolution_payload, ensure_ascii=False),
                    now,
                ),
            )

        self.reconcile_case_match_reviews()
        for case_id in case_ids:
            self.auto_check_abs(case_id)
        return {
            "decision": decision,
            "requires_abs_recheck": bool(recheck_case_ids),
            "recheck_case_ids": sorted(recheck_case_ids),
        }

    def get_case(self, case_id: str) -> dict[str, Any] | None:
        return self.db.fetch_one("SELECT * FROM cases WHERE id = ?", (case_id,))

    def get_case_source_page(self, case_id: str) -> dict[str, Any] | None:
        case = self.get_case(case_id)
        if not case:
            return None
        linked = self.db.fetch_one(
            """
            SELECT * FROM pages
            WHERE case_id = ?
            ORDER BY CASE WHEN page_type = 'letter' THEN 0 ELSE 1 END,
                     page_number
            LIMIT 1
            """,
            (case_id,),
        )
        if linked:
            return linked
        event = self.db.fetch_one(
            """
            SELECT payload_json
            FROM audit_events
            WHERE entity_type = 'case' AND entity_id = ?
              AND event_type = 'scan_case_created'
            ORDER BY id LIMIT 1
            """,
            (case_id,),
        )
        if not event:
            return None
        try:
            source_page_id = json.loads(event["payload_json"]).get(
                "source_page_id"
            )
        except (TypeError, ValueError, json.JSONDecodeError):
            return None
        if not source_page_id:
            return None
        source_page = self.get_page(str(source_page_id))
        if not source_page or source_page.get("upload_id") != case.get(
            "upload_id"
        ):
            return None
        return source_page

    def reopen_incomplete_case_review(
        self,
        case_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if (
            case.get("source_kind") != ValueSource.OCR_SCAN
            or case.get("fields_confirmed")
            or case.get("status")
            not in {CaseStatus.NEEDS_REVIEW, CaseStatus.TECHNICAL_ERROR}
        ):
            raise WorkflowValidationError(
                "Это обращение нельзя вернуть к ручной проверке"
            )
        source_page = self.get_case_source_page(case_id)
        if not source_page:
            raise WorkflowValidationError(
                "Не найден исходный лист для сверки данных"
            )
        if source_page.get("upload_id") != case.get("upload_id"):
            raise WorkflowValidationError(
                "Исходный лист относится к другому PDF"
            )
        linked_case_id = source_page.get("case_id")
        if linked_case_id not in {None, case_id}:
            raise WorkflowValidationError(
                "Исходный лист уже связан с другим обращением"
            )
        if (
            linked_case_id == case_id
            and source_page.get("status")
            in {PageStatus.NEEDS_REVIEW, PageStatus.TECHNICAL_ERROR}
        ):
            return source_page

        now = utc_now()
        before = {
            "case_id": linked_case_id,
            "page_type": source_page.get("page_type"),
            "status": source_page.get("status"),
            "manual_confirmed": bool(source_page.get("manual_confirmed")),
        }
        self.db.execute(
            """
            UPDATE pages
            SET case_id = ?, page_type = ?, type_confidence = 1,
                manual_confirmed = 0, status = ?,
                issue_code = ?, issue_message = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                case_id,
                PageType.LETTER,
                PageStatus.NEEDS_REVIEW,
                "incomplete_letter_data",
                "Заполните и подтвердите данные письма.",
                now,
                source_page["id"],
            ),
        )
        self.db.execute(
            "UPDATE cases SET status = ?, updated_at = ? WHERE id = ?",
            (CaseStatus.NEEDS_REVIEW, now, case_id),
        )
        payload = {
            "source_page_id": source_page["id"],
            "before": before,
            "after": {
                "case_id": case_id,
                "page_type": str(PageType.LETTER),
                "status": str(PageStatus.NEEDS_REVIEW),
                "manual_confirmed": False,
            },
        }
        self.db.audit(
            "case",
            case_id,
            "case_review_reopened",
            payload,
            actor=actor,
        )
        self.db.audit(
            "page",
            source_page["id"],
            "page_reopened_for_case_review",
            {"case_id": case_id, "before": before},
            actor=actor,
        )
        self._refresh_upload_status(case["upload_id"])
        return self.get_page(source_page["id"]) or source_page

    def restart_case_manual_review(
        self,
        case_id: str,
        *,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        """Return one unsent case to the manual form without erasing evidence."""
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        source_page = self.get_case_source_page(case_id)
        if not source_page or source_page.get("upload_id") != case.get("upload_id"):
            raise WorkflowValidationError(
                "Не найден исходный лист этого обращения"
            )
        if source_page.get("case_id") not in {None, case_id}:
            raise WorkflowValidationError(
                "Исходный лист уже связан с другим обращением"
            )

        outbound = self.db.fetch_one(
            """
            SELECT outlook_outgoing_messages.status,
                   outlook_outgoing_messages.resend_status
            FROM response_group_cases
            JOIN response_letters
              ON response_letters.response_group_id =
                 response_group_cases.response_group_id
            JOIN outlook_outgoing_messages
              ON outlook_outgoing_messages.response_letter_id =
                 response_letters.id
            WHERE response_group_cases.case_id = ?
              AND (
                   outlook_outgoing_messages.status IN ('sent', 'draft_created')
                   OR outlook_outgoing_messages.resend_status IN (
                       'sent', 'draft_created'
                   )
              )
            LIMIT 1
            """,
            (case_id,),
        )
        if outbound:
            if outbound["status"] == "sent" or outbound["resend_status"] == "sent":
                raise WorkflowValidationError(
                    "Письмо уже отправлено через Outlook; его нельзя вернуть "
                    "в обработку"
                )
            raise WorkflowValidationError(
                "Сначала удалите созданный черновик Outlook, затем запустите "
                "проверку заново"
            )

        group_rows = self.db.fetch_all(
            """
            SELECT response_group_id FROM response_group_cases
            WHERE case_id = ?
            """,
            (case_id,),
        )
        group_ids = [str(row["response_group_id"]) for row in group_rows]
        now = utc_now()
        before = {
            "status": case.get("status"),
            "fields_confirmed": bool(case.get("fields_confirmed")),
            "abs_status": case.get("abs_status"),
            "response_status": case.get("response_status"),
            "source_page_status": source_page.get("status"),
        }
        with self.db.connect() as connection:
            if group_ids:
                placeholders = ",".join("?" for _ in group_ids)
                connection.execute(
                    f"""
                    UPDATE signed_response_scans
                    SET status = 'superseded', updated_at = ?
                    WHERE response_letter_id IN (
                        SELECT id FROM response_letters
                        WHERE response_group_id IN ({placeholders})
                    ) AND status IN ('needs_confirmation', 'confirmed')
                    """,
                    (now, *group_ids),
                )
                connection.execute(
                    f"""
                    UPDATE signed_scan_sessions
                    SET status = 'cancelled', error_message = NULL, updated_at = ?
                    WHERE response_letter_id IN (
                        SELECT id FROM response_letters
                        WHERE response_group_id IN ({placeholders})
                    ) AND status IN ('collecting', 'technical_error')
                    """,
                    (now, *group_ids),
                )
                releasable_letters = connection.execute(
                    f"""
                    SELECT response_letters.id, response_letters.outgoing_number,
                           response_letters.response_group_id
                    FROM response_letters
                    WHERE response_letters.response_group_id IN ({placeholders})
                      AND TRIM(COALESCE(response_letters.outgoing_number, '')) != ''
                      AND NOT EXISTS (
                          SELECT 1 FROM outlook_outgoing_messages
                          WHERE outlook_outgoing_messages.response_letter_id =
                                response_letters.id
                      )
                    """,
                    group_ids,
                ).fetchall()
                for letter in releasable_letters:
                    connection.execute(
                        """
                        UPDATE response_letters
                        SET outgoing_number = NULL, assigned_at = NULL,
                            updated_at = ?
                        WHERE id = ? AND outgoing_number = ?
                        """,
                        (now, letter["id"], letter["outgoing_number"]),
                    )
                    connection.execute(
                        """
                        INSERT INTO audit_events(
                            entity_type, entity_id, event_type,
                            actor, payload_json, created_at
                        ) VALUES ('response_letter', ?, 'outgoing_number_released',
                                  ?, ?, ?)
                        """,
                        (
                            letter["id"],
                            actor,
                            json.dumps(
                                {
                                    "response_group_id": letter[
                                        "response_group_id"
                                    ],
                                    "before": letter["outgoing_number"],
                                    "reason": "manual_review_superseded",
                                },
                                ensure_ascii=False,
                            ),
                            now,
                        ),
                    )
                connection.execute(
                    f"""
                    UPDATE response_groups
                    SET status = 'superseded', updated_at = ?
                    WHERE id IN ({placeholders})
                    """,
                    (now, *group_ids),
                )
                connection.execute(
                    f"""
                    UPDATE cases
                    SET status = ?, response_status = NULL,
                        response_path = NULL, response_page_overflow = NULL,
                        updated_at = ?
                    WHERE id IN (
                        SELECT case_id FROM response_group_cases
                        WHERE response_group_id IN ({placeholders})
                    )
                    """,
                    (CaseStatus.READY_FOR_RESPONSE, now, *group_ids),
                )

            connection.execute(
                """
                UPDATE case_match_reviews
                SET status = 'superseded', updated_at = ?
                WHERE status = 'pending'
                  AND (left_case_id = ? OR right_case_id = ?)
                """,
                (now, case_id, case_id),
            )
            connection.execute(
                """
                UPDATE taxpayers
                SET manually_confirmed = 0,
                    abs_result = NULL, abs_account_result = NULL,
                    abs_active_account_count = NULL,
                    abs_closed_account_count = NULL, odb_result = NULL,
                    registry_status = NULL, registry_name = NULL,
                    registry_director = NULL, registry_checked_at = NULL,
                    registry_provider = NULL, updated_at = ?
                WHERE case_id = ?
                """,
                (now, case_id),
            )
            connection.execute(
                """
                UPDATE cases
                SET status = ?, fields_confirmed = 0,
                    abs_status = ?, response_status = NULL,
                    response_path = NULL, response_page_overflow = NULL,
                    match_review_previous_status = NULL, updated_at = ?
                WHERE id = ?
                """,
                (
                    CaseStatus.NEEDS_REVIEW,
                    AbsStatus.NOT_CHECKED,
                    now,
                    case_id,
                ),
            )
            connection.execute(
                """
                UPDATE pages
                SET case_id = ?, page_type = ?, type_confidence = 1,
                    manual_confirmed = 0, status = ?,
                    issue_code = ?, issue_message = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    case_id,
                    PageType.LETTER,
                    PageStatus.NEEDS_REVIEW,
                    "manual_restart_requested",
                    "Проверьте реквизиты письма заново. Результаты АБС и ОДБ сброшены.",
                    now,
                    source_page["id"],
                ),
            )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('case', ?, 'case_manual_review_restarted', ?, ?, ?)
                """,
                (
                    case_id,
                    actor,
                    json.dumps(
                        {
                            "source_page_id": source_page["id"],
                            "superseded_response_groups": group_ids,
                            "before": before,
                        },
                        ensure_ascii=False,
                    ),
                    now,
                ),
            )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('page', ?, 'page_manual_review_restarted', ?, ?, ?)
                """,
                (
                    source_page["id"],
                    actor,
                    json.dumps({"case_id": case_id}, ensure_ascii=False),
                    now,
                ),
            )

        self.reconcile_case_match_reviews()
        self._refresh_upload_status(case["upload_id"])
        return self.get_page(source_page["id"]) or source_page

    def get_taxpayers(self, case_id: str) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            """
            SELECT * FROM taxpayers
            WHERE case_id = ? ORDER BY display_order
            """,
            (case_id,),
        )

    def resolve_response_group_conflict(
        self,
        group_key: str,
        conflict_inn: str,
        *,
        name_choice: str,
        inn_choice: str,
        manual_name: str = "",
        manual_inn: str = "",
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        day = datetime.now(self.BUSINESS_TIMEZONE).date()
        group = next(
            (
                item
                for item in self._daily_response_groups(
                    day, AbsStatus.NOT_FOUND
                )
                if item["group_key"] == group_key
            ),
            None,
        )
        if not group:
            raise WorkflowValidationError(
                "Группа изменилась. Обновите страницу и проверьте данные снова"
            )
        conflict = next(
            (
                item
                for item in group["conflicts"]
                if item["inn"] == re.sub(r"\D", "", conflict_inn)
            ),
            None,
        )
        if not conflict:
            raise WorkflowValidationError(
                "Конфликт уже разрешён или больше не относится к этой группе"
            )
        records = conflict["records"]
        records_by_id = {record["taxpayer_id"]: record for record in records}

        def selected_value(
            field: str,
            choice: str,
            manual_value: str,
        ) -> tuple[str, str, str | None]:
            if choice == "manual":
                return manual_value, str(ValueSource.MANUAL), None
            record = records_by_id.get(choice)
            if not record:
                raise WorkflowValidationError(
                    f"Выберите источник для поля «{field}»"
                )
            source_column = "name_source" if field == "Наименование" else "inn_source"
            reference_column = (
                "name_source_reference"
                if field == "Наименование"
                else "inn_source_reference"
            )
            value_column = "name" if field == "Наименование" else "inn"
            reference = record.get(reference_column) or (
                f"taxpayer:{record['taxpayer_id']}"
            )
            return (
                str(record[value_column]),
                str(record.get(source_column) or ValueSource.MANUAL),
                reference,
            )

        selected_name, name_source, name_reference = selected_value(
            "Наименование",
            name_choice,
            manual_name,
        )
        selected_inn, inn_source, inn_reference = selected_value(
            "ИНН",
            inn_choice,
            manual_inn,
        )
        clean_name = clean_taxpayer_name(selected_name)
        clean_inn = re.sub(r"\D", "", selected_inn)
        if not clean_name:
            raise WorkflowValidationError("Укажите итоговое наименование")
        if not clean_inn:
            raise WorkflowValidationError(
                "Укажите итоговый ИНН"
            )

        record_ids = [record["taxpayer_id"] for record in records]
        case_ids = sorted({record["source_case_id"] for record in records})
        placeholders = ",".join("?" for _ in record_ids)
        for case_id in case_ids:
            duplicate = self.db.fetch_one(
                f"""
                SELECT id FROM taxpayers
                WHERE case_id = ? AND inn = ?
                  AND id NOT IN ({placeholders})
                LIMIT 1
                """,
                (case_id, clean_inn, *record_ids),
            )
            if duplicate:
                raise WorkflowValidationError(
                    "Итоговый ИНН уже указан для другого лица в одном из писем"
                )

        now = utc_now()
        recheck_case_ids: set[str] = set()
        with self.db.connect() as connection:
            current_rows = connection.execute(
                f"SELECT * FROM taxpayers WHERE id IN ({placeholders})",
                record_ids,
            ).fetchall()
            if len(current_rows) != len(record_ids):
                raise WorkflowValidationError(
                    "Одна из записей изменилась. Обновите страницу"
                )
            for row in current_rows:
                inn_changed = str(row["inn"]) != clean_inn
                if inn_changed:
                    recheck_case_ids.add(str(row["case_id"]))
                connection.execute(
                    """
                    UPDATE taxpayers
                    SET name = ?, inn = ?,
                        name_source = ?, inn_source = ?,
                        name_source_reference = ?, inn_source_reference = ?,
                        manually_confirmed = 1,
                        abs_result = CASE WHEN ? THEN NULL ELSE abs_result END,
                        abs_account_result = CASE WHEN ? THEN NULL ELSE abs_account_result END,
                        abs_active_account_count = CASE WHEN ? THEN NULL ELSE abs_active_account_count END,
                        abs_closed_account_count = CASE WHEN ? THEN NULL ELSE abs_closed_account_count END,
                        odb_result = CASE WHEN ? THEN NULL ELSE odb_result END,
                        registry_status = NULL, registry_name = NULL,
                        registry_director = NULL, registry_checked_at = NULL,
                        registry_provider = NULL, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        clean_name,
                        clean_inn,
                        name_source,
                        inn_source,
                        name_reference,
                        inn_reference,
                        int(inn_changed),
                        int(inn_changed),
                        int(inn_changed),
                        int(inn_changed),
                        int(inn_changed),
                        now,
                        row["id"],
                    ),
                )
            for case_id in recheck_case_ids:
                connection.execute(
                    """
                    UPDATE cases
                    SET status = ?, abs_status = ?, response_status = NULL,
                        response_path = NULL, response_page_overflow = NULL,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        CaseStatus.READY_FOR_ABS,
                        AbsStatus.NOT_CHECKED,
                        now,
                        case_id,
                    ),
                )
            connection.execute(
                """
                INSERT INTO audit_events(
                    entity_type, entity_id, event_type,
                    actor, payload_json, created_at
                ) VALUES ('response_group', ?,
                          'response_group_conflict_resolved', ?, ?, ?)
                """,
                (
                    group_key,
                    actor,
                    json.dumps(
                        {
                            "conflict_inn": conflict["inn"],
                            "record_ids": record_ids,
                            "name_choice": name_choice,
                            "inn_choice": inn_choice,
                            "recheck_case_ids": sorted(recheck_case_ids),
                        },
                        ensure_ascii=False,
                    ),
                    now,
                ),
            )
        for case_id in recheck_case_ids:
            taxpayers = self.get_taxpayers(case_id)
            if taxpayers and all(
                len(re.sub(r"\D", "", item.get("inn") or "")) != 14
                for item in taxpayers
            ):
                self.auto_check_abs(case_id)
        return {
            "updated_records": len(record_ids),
            "requires_abs_recheck": bool(recheck_case_ids),
            "recheck_case_ids": sorted(recheck_case_ids),
        }

    def correct_taxpayer(
        self,
        case_id: str,
        taxpayer_id: str,
        *,
        name: str,
        inn: str,
        actor: str = "Сотрудник",
    ) -> dict[str, Any]:
        case = self.get_case(case_id)
        if not case:
            raise WorkflowValidationError("Обращение не найдено")
        if case["status"] in {
            CaseStatus.RESPONSE_CREATED,
            CaseStatus.COMPLETED,
        }:
            raise WorkflowValidationError(
                "Ответ уже создан. Исправление исходных данных заблокировано"
            )
        if self.has_pending_case_match_review(case_id):
            raise WorkflowValidationError(
                "Эти версии уже находятся в ручном сравнении. "
                "Примените решение из очереди проверки"
            )
        taxpayer = self.db.fetch_one(
            "SELECT * FROM taxpayers WHERE id = ? AND case_id = ?",
            (taxpayer_id, case_id),
        )
        if not taxpayer:
            raise WorkflowValidationError("Налогоплательщик не найден")

        clean_name = clean_taxpayer_name(name)
        clean_inn = re.sub(r"\D", "", inn)
        if not clean_name:
            raise WorkflowValidationError("Укажите наименование")
        if not clean_inn:
            raise WorkflowValidationError("Укажите ИНН")
        duplicate = self.db.fetch_one(
            """
            SELECT id FROM taxpayers
            WHERE case_id = ? AND inn = ? AND id != ?
            """,
            (case_id, clean_inn, taxpayer_id),
        )
        if duplicate:
            raise WorkflowValidationError(
                "Такой ИНН уже есть в этом обращении. Исправьте существующую запись"
            )

        name_changed = clean_name != taxpayer["name"]
        inn_changed = clean_inn != taxpayer["inn"]
        if not name_changed and not inn_changed:
            raise WorkflowValidationError(
                "Данные не изменились. Исправьте неверное значение после сверки с письмом"
            )

        now = utc_now()
        with self.db.connect() as connection:
            connection.execute(
                """
                UPDATE taxpayers
                SET name = ?, inn = ?, name_source = ?, inn_source = ?,
                    name_source_reference = ?, inn_source_reference = ?,
                    manually_confirmed = 1,
                    abs_result = CASE WHEN ? THEN NULL ELSE abs_result END,
                    abs_account_result = CASE WHEN ? THEN NULL ELSE abs_account_result END,
                    abs_active_account_count = CASE WHEN ? THEN NULL ELSE abs_active_account_count END,
                    abs_closed_account_count = CASE WHEN ? THEN NULL ELSE abs_closed_account_count END,
                    odb_result = CASE WHEN ? THEN NULL ELSE odb_result END,
                    registry_status = NULL, registry_name = NULL,
                    registry_director = NULL, registry_checked_at = NULL,
                    registry_provider = NULL, updated_at = ?
                WHERE id = ? AND case_id = ?
                """,
                (
                    clean_name,
                    clean_inn,
                    ValueSource.MANUAL if name_changed else taxpayer["name_source"],
                    ValueSource.MANUAL if inn_changed else taxpayer["inn_source"],
                    None if name_changed else taxpayer["name_source_reference"],
                    None if inn_changed else taxpayer["inn_source_reference"],
                    int(inn_changed),
                    int(inn_changed),
                    int(inn_changed),
                    int(inn_changed),
                    int(inn_changed),
                    now,
                    taxpayer_id,
                    case_id,
                ),
            )
            if inn_changed:
                connection.execute(
                    """
                    UPDATE cases
                    SET status = ?, abs_status = ?, response_status = NULL,
                        response_path = NULL, response_page_overflow = NULL,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        CaseStatus.READY_FOR_ABS,
                        AbsStatus.NOT_CHECKED,
                        now,
                        case_id,
                    ),
                )

        changed_fields = [
            field
            for field, changed in (
                ("name", name_changed),
                ("inn", inn_changed),
            )
            if changed
        ]
        self.db.audit(
            "taxpayer",
            taxpayer_id,
            "taxpayer_manually_corrected",
            {
                "case_id": case_id,
                "changed_fields": changed_fields,
                "requires_abs_recheck": inn_changed,
            },
            actor=actor,
        )
        self.reconcile_case_match_reviews()
        invalid_inn = inn_changed and len(clean_inn) != 14
        if invalid_inn:
            self.auto_check_abs(case_id)
        result = {
            "changed_fields": changed_fields,
            "requires_abs_recheck": inn_changed and not invalid_inn,
        }
        if invalid_inn:
            result["invalid_inn"] = True
        return result

    def list_cases(self) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            """
            SELECT cases.*, uploads.original_filename,
                (SELECT COUNT(*) FROM taxpayers
                 WHERE taxpayers.case_id = cases.id) AS taxpayer_count
            FROM cases
            JOIN uploads ON uploads.id = cases.upload_id
            ORDER BY cases.created_at DESC
            """
        )

    def list_letter_history(
        self,
        query: str = "",
        *,
        page: int = 1,
        page_size: int = 40,
        sort_order: str = "date_desc",
    ) -> dict[str, Any]:
        search = " ".join(str(query or "").split())[:160]
        page_size = max(10, min(int(page_size), 100))
        requested_page = max(int(page), 1)
        if sort_order not in {
            "date_desc",
            "date_asc",
            "district_asc",
            "district_desc",
        }:
            sort_order = "date_desc"
        history_cte = """
            WITH letter_history AS (
                SELECT
                    cases.id,
                    cases.upload_id,
                    cases.status,
                    cases.created_at,
                    cases.district_place,
                    cases.recipient_display_name,
                    uploads.original_filename,
                    uploads.page_count,
                    uploads.intake_source,
                    uploads.created_at AS uploaded_at,
                    (SELECT COUNT(*) FROM taxpayers
                     WHERE taxpayers.case_id = cases.id)
                        AS taxpayer_count,
                    (SELECT GROUP_CONCAT(taxpayers.name, ' · ')
                     FROM taxpayers
                     WHERE taxpayers.case_id = cases.id)
                        AS taxpayer_names,
                    (SELECT GROUP_CONCAT(taxpayers.inn, ' · ')
                     FROM taxpayers
                     WHERE taxpayers.case_id = cases.id)
                        AS taxpayer_inns,
                    (SELECT GROUP_CONCAT(page_numbers.page_number, ', ')
                     FROM (
                        SELECT pages.page_number
                        FROM pages
                        WHERE pages.case_id = cases.id
                        ORDER BY pages.page_number
                     ) AS page_numbers)
                        AS source_page_numbers,
                    (SELECT MIN(pages.page_number)
                     FROM pages
                     WHERE pages.case_id = cases.id)
                        AS source_first_page,
                    CASE WHEN uploads.intake_source = 'outlook' THEN
                    (SELECT outlook_messages.sender_smtp
                     FROM outlook_attachments
                     JOIN outlook_messages
                       ON outlook_messages.source_key =
                          outlook_attachments.message_key
                     WHERE outlook_attachments.upload_id = cases.upload_id
                     ORDER BY outlook_messages.received_at
                     LIMIT 1) END AS sender_smtp,
                    CASE WHEN uploads.intake_source = 'outlook' THEN
                    (SELECT outlook_messages.original_sender_smtp
                     FROM outlook_attachments
                     JOIN outlook_messages
                       ON outlook_messages.source_key =
                          outlook_attachments.message_key
                     WHERE outlook_attachments.upload_id = cases.upload_id
                     ORDER BY outlook_messages.received_at
                     LIMIT 1) END AS original_sender_smtp,
                    CASE WHEN uploads.intake_source = 'outlook' THEN
                    (SELECT outlook_messages.received_at
                     FROM outlook_attachments
                     JOIN outlook_messages
                       ON outlook_messages.source_key =
                          outlook_attachments.message_key
                     WHERE outlook_attachments.upload_id = cases.upload_id
                     ORDER BY outlook_messages.received_at
                     LIMIT 1) END AS received_at,
                    (SELECT response_groups.id
                     FROM response_group_cases
                     JOIN response_groups
                       ON response_groups.id =
                          response_group_cases.response_group_id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY response_groups.created_at DESC
                     LIMIT 1) AS response_group_id,
                    (SELECT GROUP_CONCAT(response_letters.outgoing_number, ', ')
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     WHERE response_group_cases.case_id = cases.id
                       AND response_letters.outgoing_number IS NOT NULL
                       AND response_letters.outgoing_number != '')
                        AS outgoing_numbers,
                    (SELECT signed_response_scans.id
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN signed_response_scans
                       ON signed_response_scans.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                       AND signed_response_scans.status != 'superseded'
                     ORDER BY signed_response_scans.created_at DESC
                     LIMIT 1) AS signed_scan_id,
                    (SELECT outlook_outgoing_messages.id
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN outlook_outgoing_messages
                       ON outlook_outgoing_messages.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY outlook_outgoing_messages.created_at DESC
                     LIMIT 1) AS outlook_message_id,
                    (SELECT outlook_outgoing_messages.status
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN outlook_outgoing_messages
                       ON outlook_outgoing_messages.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY outlook_outgoing_messages.created_at DESC
                     LIMIT 1) AS mail_status,
                    (SELECT outlook_outgoing_messages.recipient_email
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN outlook_outgoing_messages
                       ON outlook_outgoing_messages.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY outlook_outgoing_messages.created_at DESC
                     LIMIT 1) AS recipient_email,
                    (SELECT outlook_outgoing_messages.sent_at
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN outlook_outgoing_messages
                       ON outlook_outgoing_messages.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY outlook_outgoing_messages.created_at DESC
                     LIMIT 1) AS sent_at,
                    (SELECT outlook_outgoing_messages.resend_status
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN outlook_outgoing_messages
                       ON outlook_outgoing_messages.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY outlook_outgoing_messages.created_at DESC
                     LIMIT 1) AS resend_status,
                    (SELECT outlook_outgoing_messages.resent_at
                     FROM response_group_cases
                     JOIN response_letters
                       ON response_letters.response_group_id =
                          response_group_cases.response_group_id
                     JOIN outlook_outgoing_messages
                       ON outlook_outgoing_messages.response_letter_id =
                          response_letters.id
                     WHERE response_group_cases.case_id = cases.id
                     ORDER BY outlook_outgoing_messages.created_at DESC
                     LIMIT 1) AS resent_at
                FROM cases
                JOIN uploads ON uploads.id = cases.upload_id
            )
        """
        where = ""
        search_parameters: tuple[Any, ...] = ()
        if search:
            where = """
                WHERE gns_fuzzy_match(
                    COALESCE(original_filename, '') || ' ' ||
                    COALESCE(district_place, '') || ' ' ||
                    COALESCE(recipient_display_name, '') || ' ' ||
                    COALESCE(taxpayer_names, '') || ' ' ||
                    COALESCE(taxpayer_inns, '') || ' ' ||
                    COALESCE(sender_smtp, '') || ' ' ||
                    COALESCE(original_sender_smtp, '') || ' ' ||
                    COALESCE(outgoing_numbers, '') || ' ' ||
                    COALESCE(status, ''),
                    ?
                ) = 1
            """
            search_parameters = (search,)
        order_by = {
            "date_desc": (
                "COALESCE(received_at, uploaded_at) DESC, uploaded_at DESC"
            ),
            "date_asc": (
                "COALESCE(received_at, uploaded_at) ASC, uploaded_at ASC"
            ),
            "district_asc": (
                "gns_search_normalize(COALESCE(district_place, '')) ASC, "
                "COALESCE(received_at, uploaded_at) DESC"
            ),
            "district_desc": (
                "gns_search_normalize(COALESCE(district_place, '')) DESC, "
                "COALESCE(received_at, uploaded_at) DESC"
            ),
        }[sort_order]
        total_row = self.db.fetch_one(
            history_cte
            + f"SELECT COUNT(*) AS count FROM letter_history {where}",
            search_parameters,
        ) or {"count": 0}
        total = int(total_row.get("count") or 0)
        page_count = max(1, (total + page_size - 1) // page_size)
        selected_page = min(requested_page, page_count)
        offset = (selected_page - 1) * page_size
        items = self.db.fetch_all(
            history_cte
            + f"""
                SELECT * FROM letter_history
                {where}
                ORDER BY {order_by},
                         original_filename,
                         COALESCE(source_first_page, 999999999),
                         id
                LIMIT ? OFFSET ?
            """,
            (*search_parameters, page_size, offset),
        )
        return {
            "items": items,
            "query": search,
            "sort_order": sort_order,
            "page": selected_page,
            "page_count": page_count,
            "page_size": page_size,
            "total": total,
        }

    def get_audit(self, limit: int = 200) -> list[dict[str, Any]]:
        rows = self.db.fetch_all(
            """
            SELECT * FROM audit_events
            ORDER BY id DESC LIMIT ?
            """,
            (limit,),
        )
        for row in rows:
            try:
                row["payload"] = json.loads(row["payload_json"])
            except json.JSONDecodeError:
                row["payload"] = {}
        return rows

    def dashboard_stats(self) -> dict[str, int]:
        self.reconcile_case_match_reviews()
        page_review_count = (
            self.db.fetch_one(
                """
                SELECT COUNT(*) AS n FROM pages
                WHERE status IN ('needs_review', 'technical_error')
                """
            )
            or {"n": 0}
        )["n"]
        match_review_count = (
            self.db.fetch_one(
                """
                SELECT COUNT(*) AS n FROM case_match_reviews
                WHERE status = 'pending'
                """
            )
            or {"n": 0}
        )["n"]
        return {
            "uploads": (
                self.db.fetch_one("SELECT COUNT(*) AS n FROM uploads") or {"n": 0}
            )["n"],
            "pages": (
                self.db.fetch_one("SELECT COUNT(*) AS n FROM pages") or {"n": 0}
            )["n"],
            "review": int(page_review_count) + int(match_review_count),
            "ready_abs": (
                self.db.fetch_one(
                    "SELECT COUNT(*) AS n FROM cases WHERE status = 'ready_for_abs'"
                )
                or {"n": 0}
            )["n"],
            "ready_response": (
                self.db.fetch_one(
                    """
                    SELECT COUNT(*) AS n FROM cases
                    WHERE status = 'ready_for_response'
                    """
                )
                or {"n": 0}
            )["n"],
        }
